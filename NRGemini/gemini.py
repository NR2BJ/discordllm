import discord
from discord import app_commands
from discord.ext import commands, tasks
import csv
import mafic
from google import genai
from google.genai import types
from bs4 import BeautifulSoup
import base64
import json
import os
import asyncio
import aiohttp
import hashlib
import html as html_module
import io
import warnings
import re
import tempfile
import traceback
import time as time_module
import math
import inspect
from importlib.metadata import PackageNotFoundError, version as package_version
from pathlib import Path
from PIL import Image
from datetime import datetime, timedelta, timezone
import glob as glob_module
import unicodedata
from urllib.parse import urlsplit

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

warnings.filterwarnings("ignore", category=DeprecationWarning, module="google.genai._api_client")

try:
    GENAI_SDK_VERSION = package_version("google-genai")
except PackageNotFoundError:
    GENAI_SDK_VERSION = "unknown"


def is_enabled(value):
    if isinstance(value, str):
        return value.strip().lower() in ("1", "true", "yes", "on")
    return bool(value)


ENABLE_UVLOOP = os.environ.get("GEMINI_ENABLE_UVLOOP", "0") == "1"
if ENABLE_UVLOOP:
    try:
        import uvloop
    except Exception:
        uvloop = None
    else:
        uvloop.install()
else:
    uvloop = None

if not discord.opus.is_loaded():
    try:
        discord.opus.load_opus('libopus.so.0')
    except OSError:
        discord.opus.load_opus('/opt/homebrew/lib/libopus.dylib')

# ================= 1. 설정 및 상수 =================
DISCORD_TOKEN = os.environ["GEMINI_DISCORD_TOKEN"]
GOOGLE_API_KEY = os.environ["GEMINI_API_KEY"].strip()
if not GOOGLE_API_KEY:
    raise RuntimeError("GEMINI_API_KEY가 비어 있습니다.")

try:
    CLIENT = genai.Client(api_key=GOOGLE_API_KEY)
except Exception as e:
    print(f"[ERROR] Google Client 초기화 실패: {e}")
    CLIENT = None

print(f"[GEMINI] google-genai SDK {GENAI_SDK_VERSION}")

FLASH_MODEL = "gemini-3.8-flash"
FLASH_LITE_MODEL = "gemini-3.5-flash-lite"
IMAGE_MODEL = "gemini-3.1-flash-image"
MAX_THINKING_LEVEL = "high"

AVAILABLE_MODELS = {
    "flash-lite": {
        "display_name": "Flash-Lite 최고 추론 (Gemini 3.5 Flash-Lite)",
        "text": FLASH_LITE_MODEL,
        "image": IMAGE_MODEL,
        "thinking_level": MAX_THINKING_LEVEL,
    },
    "flash": {
        "display_name": "Flash 최고 추론 (Gemini 3.8 Flash)",
        "text": FLASH_MODEL,
        "image": IMAGE_MODEL,
        "thinking_level": MAX_THINKING_LEVEL,
    },
}
TEXT_MODEL = FLASH_MODEL
# Deep Research는 모델 선택과 분리된 전용 명령어(/deepo, /deepm)로만 실행
DEEP_RESEARCH_AGENTS = {
    "deepo": "deep-research-preview-04-2026",
    "deepm": "deep-research-max-preview-04-2026",
}
DEEP_RESEARCH_POLL_INTERVAL_SECONDS = 10
DEEP_RESEARCH_STATUS_LOG_INTERVAL_SECONDS = 300
GEMINI_API_TIMEOUT_SECONDS = 120.0
GEMINI_IMAGE_TIMEOUT_SECONDS = 300.0
GEMINI_IMAGE_MAX_ATTEMPTS = 2
ROUTING_MODEL = os.environ.get(
    "GEMINI_ROUTING_MODEL",
    FLASH_LITE_MODEL,
)
DEFAULT_ACTION_ROUTER_ENABLED = is_enabled(os.environ.get("GEMINI_ENABLE_ACTION_ROUTER", "1"))
# VIDEO_GEN_MODEL = "veo-3.1-generate-preview"  # 비용 절감으로 임시 비활성화

DEFAULT_HISTORY_LIMIT = 20
MIN_HISTORY_LIMIT = 1
MAX_HISTORY_LIMIT = 100
DEFAULT_IMAGE_HISTORY_LIMIT = int(os.environ.get("GEMINI_IMAGE_HISTORY_LIMIT", "10"))
MIN_IMAGE_HISTORY_LIMIT = 0
MAX_IMAGE_HISTORY_LIMIT = 50
IMAGE_SEARCH_LIMIT = 2
MAX_DOC_CHARS = 12000
MEMORY_EXCERPT_MAX_CHARS = 5000
TEXT_ATTACHMENT_EXTENSIONS = (
    ".txt", ".py", ".json", ".md", ".c", ".cpp", ".js", ".html", ".css",
    ".log", ".xml", ".yaml", ".yml", ".toml", ".ini", ".sh", ".csv", ".tsv",
)
MODEL_HISTORY_LIMITS = {
    FLASH_MODEL: 20,
    FLASH_LITE_MODEL: 20,
}

IMAGE_ACTION_NOUN_RE = re.compile(
    r"(이미지|그림|사진|짤|일러스트|삽화|캐릭터|배경|썸네일|아이콘|로고|아바타|프로필|만화|웹툰|"
    r"\bimage\b|\bpicture\b|\bphoto\b|\billustration\b|\bdrawing\b|\blogo\b|\bicon\b|\bthumbnail\b|\bmeme\b|\bavatar\b|\bwallpaper\b)",
    re.IGNORECASE,
)
IMAGE_ACTION_VERB_RE = re.compile(
    r"(그려|그리|그린|그림으로|이미지로|사진처럼|짤로|생성|만들|제작|렌더|편집|수정|변환|합성|지워|넣어|추가|바꿔|"
    r"\bgenerate\b|\bdraw\b|\bcreate\b|\bmake\b|\brender\b|\bedit\b|\bmodify\b|\btransform\b|\bremove\b|\badd\b)",
    re.IGNORECASE,
)
IMAGE_DRAW_REQUEST_RE = re.compile(
    r"(그려\s*(줘|봐|라|달라|달라고|주라|주세요)?|그리(라고|라|세요|시오|기)|그리(?=\s|$|[.!?,])|"
    r"\bdraw\b|\bsketch\b)",
    re.IGNORECASE,
)
CURRENT_IMAGE_EDIT_RE = re.compile(
    r"(이거|이것|저거|저것|첨부|사진|이미지|그림|짤|reference|attached|this|that).{0,40}"
    r"(편집|수정|변환|합성|지워|넣어|추가|바꿔|스타일|풍으로|배경|edit|modify|transform|remove|add|style|background)",
    re.IGNORECASE,
)
IMAGE_ATTACHMENT_EXTENSIONS = (".png", ".jpg", ".jpeg", ".webp", ".gif")
FOLLOWUP_IMAGE_EDIT_RE = re.compile(
    r"(이거|이것|이게|저거|저것|그거|그것|아까|방금|전에|위에|최근|this|that|previous|recent|above).{0,80}"
    r"(만들|그려|그리|바꿔|변환|수정|편집|합성|처럼|같이|느낌|스타일|풍으로|make|draw|turn|change|edit|style)",
    re.IGNORECASE,
)
_COMMON_RULES = """## General Rules
- Respond in the user's language and mirror their tone (casual, formal, vulgar — match them).
- If the user sets up a specific roleplay scenario (character, world, setting), maintain it consistently within that conversation. Do NOT treat casual chat as "established context" that must persist.
- Analyze attached files (images/video/audio) when supported.
- Keep responses concise and Discord-appropriate."""

# search_types를 넣으면 Interactions API가 검색을 서버에서 실행하지 않고 클라이언트 function_call
# ('google:search', status=requires_action)로 넘겨 봇이 답을 못 받는다. 모델·store 설정과 무관하게 재현됨.
GOOGLE_SEARCH_TOOL = {"type": "google_search"}
RAW_CITATION_MARKER_RE = re.compile(
    r"(?:[ \t]*\[cite:\s*[^\]\r\n]+\])+",
    re.IGNORECASE,
)

SYSTEM_PROMPT = f"""
You are an AI member of this Discord server.

## Conversation Context
- Chat history is for CONTEXT ONLY — to understand the flow, not to recycle.
- When the topic changes, focus ONLY on the current topic. Do NOT drag in old topics.
- Do NOT reuse specific words, jokes, memes, or references from previous conversations to fake familiarity. Every new topic gets a fresh start.
- User's background info (job, interests, etc.) is for matching TONE only. Do NOT shoehorn it into answers unless the user explicitly brings it up or it has direct logical relevance to the current question.
- Answer new questions cleanly and on-point. Do NOT over-decorate answers by mixing in past context.
- You may briefly reference a previous topic AT MOST once if naturally relevant, then move on entirely.

{_COMMON_RULES}
"""

CHAT_SYSTEM_PROMPT = f"""
You are a helpful AI assistant.
{_COMMON_RULES}
- If an image is provided, analyze it carefully and respond based on both the image and the user's message.
"""

IMAGE_SYSTEM_PROMPT = f"""
You are an AI image generator assistant.
{_COMMON_RULES}
- Return a detailed English image prompt based on the user's request.
- The `message` field may contain a brief response to the user, or be empty.
- If a reference image is provided, describe the requested edit while preserving relevant source details.
"""

# 현재 파일(gemini.py)이 있는 폴더 경로를 구함
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

CONFIG_FILE = os.environ.get("CONFIG_FILE", os.path.join(BASE_DIR, "server.json"))
KST = timezone(timedelta(hours=9))
DOC_CACHE_DIR = os.environ.get("GEMINI_DOC_CACHE_DIR", os.path.join(os.path.dirname(CONFIG_FILE), "gemini_doc_cache"))
DOC_CACHE_AUTO_REPLAY_WINDOW = 2

SOUND_DIR = os.path.join(BASE_DIR, "sb")
SOUND_EXTENSIONS = ('.mp3', '.wav', '.flac', '.ogg', '.m4a', '.aac', '.opus', '.weba')
os.makedirs(SOUND_DIR, exist_ok=True)
os.makedirs(os.path.join(SOUND_DIR, "discord"), exist_ok=True)
LAVALINK_PAD_CACHE_DIR = Path(os.environ.get("GEMINI_LAVALINK_PAD_CACHE_DIR", os.path.join(SOUND_DIR, ".gemini_lavalink_pad")))
SOUND_VOLUME = 0.15  # 기본 볼륨 (0.0~1.0)
MIN_SOUND_VOLUME = 0.01
MAX_SOUND_VOLUME = 1.0
SIBO_LATENCY_COMP = 0.0  # Discord 전송 지연 보정 (초)
SIBO_PRELOAD_LEAD = 0.3  # 정각 전 ffmpeg 미리 띄워 파이프 버퍼 채우는 시간 (초)
LAVALINK_POST_CONNECT_GRACE = 0.50  # Lavalink 최초 연결 직후 첫 재생 앞머리 손실 완충
LAVALINK_SHORT_CLIP_PAD_MS = 220
LAVALINK_SIBO_PAD_MS = 120
VOICE_RECONNECT_DELAY = 3.0
VOICE_STATE_REJOIN_DELAY = 8.0
VOICE_RECONNECT_COOLDOWN = 20.0
VOICE_LIBRARY_RECOVERY_GRACE = 75.0
VOICE_HEALTH_GRACE = 30.0
VOICE_HEALTH_BAD_LIMIT = 3
VOICE_HEALTH_RECOVERY = os.environ.get("GEMINI_VOICE_HEALTH_RECOVERY", "0") == "1"
VOICE_SELF_DEAF = True
VOICE_SELF_MUTE = False
VOICE_FORCE_STATE = os.environ.get("GEMINI_VOICE_FORCE_STATE", "0") == "1"
VOICE_SILENCE_KEEPALIVE = os.environ.get("GEMINI_VOICE_SILENCE_KEEPALIVE", "0") == "1"
VOICE_SILENCE_RESUME_DELAY = 0.5
VOICE_DEBUG_LOGS = False
LOOP_MONITOR_INTERVAL = 2.0
LOOP_MONITOR_WARN = 2.5
LAVALINK_ENABLED = os.environ.get("GEMINI_USE_LAVALINK", "0") == "1"
LAVALINK_HOST = os.environ.get("LAVALINK_HOST", "lavalink")
LAVALINK_PORT = int(os.environ.get("LAVALINK_PORT", "2333"))
LAVALINK_PASSWORD = os.environ.get("LAVALINK_PASSWORD", "youshallnotpass")
LAVALINK_SECURE = os.environ.get("LAVALINK_SECURE", "0") == "1"
LAVALINK_LABEL = os.environ.get("LAVALINK_LABEL", "main")
LAVALINK_LOCAL_ROOT = Path(os.environ.get("LAVALINK_LOCAL_ROOT", "/opt/lavalink/sb"))
LAVALINK_TIMEOUT = float(os.environ.get("LAVALINK_TIMEOUT", "10"))

def format_channel_ref(channel):
    if channel is None:
        return "None"
    name = getattr(channel, "name", str(channel))
    channel_id = getattr(channel, "id", None)
    return f"{name}({channel_id})" if channel_id else name

def format_metric(value):
    if value is None:
        return "n/a"
    if isinstance(value, float):
        return f"{value:.3f}"
    return str(value)

def is_lavalink_player(vc):
    return LAVALINK_ENABLED and isinstance(vc, mafic.Player)

def voice_is_connected(vc):
    if vc is None:
        return False
    try:
        return vc.is_connected()
    except Exception:
        return bool(getattr(vc, "connected", False))

def voice_is_playing(vc):
    if vc is None:
        return False
    if is_lavalink_player(vc):
        return getattr(vc, "current", None) is not None and not getattr(vc, "paused", False)
    try:
        return vc.is_playing()
    except Exception:
        return False

def voice_is_paused(vc):
    if vc is None:
        return False
    if is_lavalink_player(vc):
        return bool(getattr(vc, "paused", False))
    try:
        return vc.is_paused()
    except Exception:
        return False

async def voice_stop(vc):
    if vc is None:
        return
    if is_lavalink_player(vc):
        await vc.stop()
        return
    vc.stop()

async def voice_disconnect(vc, *, force=True):
    if vc is None:
        return
    if is_lavalink_player(vc):
        await vc.disconnect(force=force)
        return
    try:
        await vc.disconnect(force=force)
    except TypeError:
        await vc.disconnect()

async def attempt_voice_state_rejoin(guild_id, channel_id):
    await asyncio.sleep(VOICE_STATE_REJOIN_DELAY)

    guild = bot.get_guild(guild_id)
    if guild is None or bot.is_closed():
        return

    if bot.get_voice_target(guild_id) != channel_id:
        return

    vc = guild.voice_client
    if vc and voice_is_connected(vc):
        return

    if not bot.can_retry_voice(guild_id):
        return

    try:
        log_voice_snapshot(
            "voice_state 빠른 재접속 시도",
            guild,
            vc,
            extra=f"target_channel_id={channel_id}",
        )
        vc = await ensure_voice_connection(
            guild,
            channel_id,
            move=True,
            reason="voice_state",
        )
        if vc:
            bot.reset_voice_retry(guild_id)
    except Exception as e:
        log_voice_debug(f"[VC] voice_state 빠른 재접속 실패 ({guild.name}): {e}")

def build_lavalink_local_identifier(filepath):
    relative = Path(filepath).resolve().relative_to(Path(SOUND_DIR).resolve())
    # Lavalink local source expects a filesystem path, not a file:// URI.
    return (LAVALINK_LOCAL_ROOT / relative).as_posix()

def get_voice_ws_closed(vc):
    ws = getattr(vc, "ws", None)
    if ws is None:
        return None
    closed = getattr(ws, "closed", None)
    if callable(closed):
        try:
            return closed()
        except Exception:
            return "error"
    return closed

def get_history_limit_for_text_model(text_model):
    return MODEL_HISTORY_LIMITS.get(text_model, DEFAULT_HISTORY_LIMIT)

def clamp_history_limit(value):
    try:
        limit = int(value)
    except (TypeError, ValueError):
        limit = DEFAULT_HISTORY_LIMIT
    return max(MIN_HISTORY_LIMIT, min(MAX_HISTORY_LIMIT, limit))

def get_history_limit(config):
    return clamp_history_limit((config or {}).get("history_limit", DEFAULT_HISTORY_LIMIT))


def clamp_image_history_limit(value):
    try:
        limit = int(value)
    except (TypeError, ValueError):
        limit = DEFAULT_IMAGE_HISTORY_LIMIT
    return max(MIN_IMAGE_HISTORY_LIMIT, min(MAX_IMAGE_HISTORY_LIMIT, limit))


def get_image_history_limit(config):
    return clamp_image_history_limit((config or {}).get("image_history_limit", DEFAULT_IMAGE_HISTORY_LIMIT))


def clamp_sound_volume(value):
    try:
        volume = float(value)
    except (TypeError, ValueError):
        return SOUND_VOLUME
    if volume > 1:
        volume /= 100.0
    return max(MIN_SOUND_VOLUME, min(MAX_SOUND_VOLUME, volume))


def get_voice_unhealthy_reason(vc):
    if vc is None:
        return None
    if is_lavalink_player(vc):
        return None
    reasons = []
    latency = getattr(vc, "latency", None)
    average_latency = getattr(vc, "average_latency", None)
    if isinstance(latency, float) and not math.isfinite(latency):
        reasons.append(f"latency={latency}")
    if isinstance(average_latency, float) and not math.isfinite(average_latency):
        reasons.append(f"avg_latency={average_latency}")
    ws_closed = get_voice_ws_closed(vc)
    if ws_closed is True:
        reasons.append("voice_ws_closed=True")
    endpoint = getattr(vc, "endpoint", None)
    if not endpoint:
        reasons.append("endpoint=missing")
    return ", ".join(reasons) if reasons else None

def describe_voice_client(vc):
    if vc is None:
        return "vc=None"
    if is_lavalink_player(vc):
        return (
            f"connected={voice_is_connected(vc)} "
            f"playing={voice_is_playing(vc)} "
            f"paused={voice_is_paused(vc)} "
            f"channel={format_channel_ref(getattr(vc, 'channel', None))} "
            f"node={getattr(getattr(vc, 'node', None), 'label', '-') } "
            f"ping={format_metric(getattr(vc, 'ping', None))} "
            f"current={getattr(getattr(vc, 'current', None), 'identifier', '-')}"
        )
    try:
        connected = vc.is_connected()
    except Exception as e:
        connected = f"err:{e.__class__.__name__}"
    try:
        playing = vc.is_playing()
    except Exception as e:
        playing = f"err:{e.__class__.__name__}"
    try:
        paused = vc.is_paused()
    except Exception as e:
        paused = f"err:{e.__class__.__name__}"
    endpoint = getattr(vc, "endpoint", None) or "-"
    session_id = getattr(vc, "session_id", None) or "-"
    ws_closed = get_voice_ws_closed(vc)
    return (
        f"connected={connected} playing={playing} paused={paused} "
        f"channel={format_channel_ref(getattr(vc, 'channel', None))} "
        f"endpoint={endpoint} session_id={session_id} "
        f"latency={format_metric(getattr(vc, 'latency', None))} "
        f"avg_latency={format_metric(getattr(vc, 'average_latency', None))} "
        f"ws_closed={ws_closed}"
    )

def describe_voice_state(state):
    if state is None:
        return "state=None"
    return (
        f"channel={format_channel_ref(getattr(state, 'channel', None))} "
        f"self_mute={getattr(state, 'self_mute', None)} "
        f"self_deaf={getattr(state, 'self_deaf', None)} "
        f"mute={getattr(state, 'mute', None)} "
        f"deaf={getattr(state, 'deaf', None)} "
        f"suppress={getattr(state, 'suppress', None)} "
        f"afk={getattr(state, 'afk', None)}"
    )

def log_voice_snapshot(prefix, guild=None, vc=None, extra=None):
    if not VOICE_DEBUG_LOGS:
        return
    parts = [f"[VC] {prefix}"]
    if guild is not None:
        parts.append(f"guild={guild.name}({guild.id})")
    if extra:
        parts.append(extra)
    parts.append(describe_voice_client(vc))
    print(" | ".join(parts))

def log_voice_debug(message):
    if VOICE_DEBUG_LOGS:
        print(message)

async def cleanup_stale_voice_client(vc, guild):
    """끊어진 VoiceClient 캐시를 정리해서 재접속 막힘 방지."""
    log_voice_snapshot("stale VC 정리 시작", guild, vc)
    try:
        await voice_disconnect(vc, force=True)
    except Exception as e:
        log_voice_debug(f"[VC] stale disconnect 실패 ({guild.name}): {e}")
    try:
        vc.cleanup()
    except Exception:
        pass
    log_voice_snapshot("stale VC 정리 완료", guild, guild.voice_client)

async def resolve_voice_channel(guild, channel_id):
    if not channel_id:
        return None
    channel = guild.get_channel(channel_id)
    if channel is None:
        try:
            channel = await bot.fetch_channel(channel_id)
        except Exception as e:
            log_voice_debug(f"[VC] 채널 조회 실패 ({guild.name}, {channel_id}): {e}")
            return None
    if not hasattr(channel, "connect"):
        log_voice_debug(f"[VC] 연결 불가 채널 타입 ({guild.name}, {channel_id})")
        return None
    return channel

async def apply_voice_preferences(guild, channel):
    if LAVALINK_ENABLED:
        return
    if not VOICE_FORCE_STATE:
        return
    me = guild.me or guild.get_member(bot.user.id)
    current_state = getattr(me, "voice", None) if me else None
    if (
        current_state
        and current_state.channel == channel
        and current_state.self_deaf == VOICE_SELF_DEAF
        and current_state.self_mute == VOICE_SELF_MUTE
    ):
        return
    try:
        await guild.change_voice_state(
            channel=channel,
            self_deaf=VOICE_SELF_DEAF,
            self_mute=VOICE_SELF_MUTE,
        )
    except Exception as e:
        log_voice_debug(f"[VC] voice state 선호값 적용 실패 ({guild.name}): {e}")

def make_silence_audio_source():
    return discord.FFmpegPCMAudio(
        "anullsrc=channel_layout=stereo:sample_rate=48000",
        before_options="-f lavfi",
        options="-vn",
    )

async def ensure_silence_keepalive(guild, vc=None, *, reason="unknown"):
    if LAVALINK_ENABLED:
        return False
    if not VOICE_SILENCE_KEEPALIVE:
        return False

    target_channel_id = bot.get_voice_target(guild.id)
    if not target_channel_id:
        return False

    if vc is None:
        vc = guild.voice_client

    if vc is None or not voice_is_connected(vc):
        return False

    try:
        if voice_is_playing(vc) or voice_is_paused(vc):
            return False
    except Exception:
        return False

    try:
        vc.play(make_silence_audio_source())
        log_voice_debug(f"[VC] 무음 keepalive 시작 ({guild.name}) [{reason}]")
        return True
    except Exception as e:
        log_voice_debug(f"[VC] 무음 keepalive 시작 실패 ({guild.name}) [{reason}]: {e}")
        return False

async def resume_silence_keepalive(guild_id, delay=VOICE_SILENCE_RESUME_DELAY, *, reason="unknown"):
    if not VOICE_SILENCE_KEEPALIVE:
        return

    if delay > 0:
        await asyncio.sleep(delay)

    guild = bot.get_guild(guild_id)
    if guild is None or bot.is_closed():
        return

    target_channel_id = bot.get_voice_target(guild.id)
    if not target_channel_id:
        return

    vc = guild.voice_client
    if vc is None or not vc.is_connected():
        vc = await ensure_voice_connection(
            guild,
            target_channel_id,
            move=True,
            reason=f"{reason}_resume",
        )
        if vc is None:
            return

    await ensure_silence_keepalive(guild, vc, reason=reason)

def schedule_silence_keepalive_resume(guild_id, delay=VOICE_SILENCE_RESUME_DELAY, *, reason="unknown"):
    if not VOICE_SILENCE_KEEPALIVE or bot.is_closed():
        return

    try:
        asyncio.run_coroutine_threadsafe(
            resume_silence_keepalive(guild_id, delay=delay, reason=reason),
            bot.loop,
        )
    except Exception as e:
        guild = bot.get_guild(guild_id)
        guild_name = guild.name if guild else guild_id
        log_voice_debug(f"[VC] 무음 keepalive 재개 예약 실패 ({guild_name}) [{reason}]: {e}")

async def ensure_voice_connection(guild, channel_id, *, move=True, reason="unknown"):
    channel = await resolve_voice_channel(guild, channel_id)
    if channel is None:
        log_voice_debug(f"[VC] 연결 대상 채널 조회 실패 ({guild.name}, {channel_id}) [{reason}]")
        return None

    async with bot.get_voice_lock(guild.id):
        vc = guild.voice_client
        log_voice_snapshot(
            f"ensure_voice_connection 시작 [{reason}]",
            guild,
            vc,
            extra=f"target={format_channel_ref(channel)} move={move}",
        )
        if vc and voice_is_connected(vc):
            if move and vc.channel != channel:
                old_channel = vc.channel
                if is_lavalink_player(vc):
                    await vc.disconnect(force=True)
                    vc = await channel.connect(
                        cls=mafic.Player,
                        timeout=LAVALINK_TIMEOUT,
                        reconnect=True,
                        self_deaf=VOICE_SELF_DEAF,
                        self_mute=VOICE_SELF_MUTE,
                    )
                else:
                    await vc.move_to(channel)
                log_voice_snapshot(
                    f"이동 완료 [{reason}]",
                    guild,
                    vc,
                    extra=f"from={format_channel_ref(old_channel)} to={format_channel_ref(channel)}",
                )
            await apply_voice_preferences(guild, channel)
            bot.mark_voice_connected(guild.id)
            await ensure_silence_keepalive(guild, vc, reason=reason)
            return vc

        if vc:
            await cleanup_stale_voice_client(vc, guild)

        if LAVALINK_ENABLED:
            await bot.ensure_lavalink_node()
            vc = await channel.connect(
                cls=mafic.Player,
                timeout=LAVALINK_TIMEOUT,
                reconnect=True,
                self_deaf=VOICE_SELF_DEAF,
                self_mute=VOICE_SELF_MUTE,
            )
        else:
            try:
                connect_kwargs = {"reconnect": True}
                if VOICE_FORCE_STATE:
                    connect_kwargs.update(
                        self_deaf=VOICE_SELF_DEAF,
                        self_mute=VOICE_SELF_MUTE,
                    )
                vc = await channel.connect(**connect_kwargs)
            except TypeError:
                vc = await channel.connect()

        await apply_voice_preferences(guild, channel)
        log_voice_snapshot(
            f"연결 완료 [{reason}]",
            guild,
            vc,
            extra=f"target={format_channel_ref(channel)}",
        )
        bot.mark_voice_connected(guild.id)
        await ensure_silence_keepalive(guild, vc, reason=reason)
        return vc

# ================= 사운드보드 유틸리티 =================

def scan_sound_files():
    """재귀 스캔. 키: 상대경로(확장자 제외, 소문자). 값: 절대경로."""
    sounds = {}
    for root, dirs, files in os.walk(SOUND_DIR):
        dirs[:] = [d for d in dirs if not d.startswith(".") and d != "__pycache__"]
        for f in files:
            if f.startswith("."):
                continue
            ext = os.path.splitext(f)[1].lower()
            if ext in SOUND_EXTENSIONS:
                fullpath = os.path.join(root, f)
                relpath = os.path.relpath(fullpath, SOUND_DIR)
                key = os.path.splitext(relpath)[0].lower()
                key = key.replace("\\", "/")
                # macOS NFD → NFC 정규화 (한글 호환)
                key = unicodedata.normalize("NFC", key)
                sounds[key] = fullpath
    return sounds

async def temp_reply(interaction, content, delay=5):
    """메시지 보내고 delay초 후 자동 삭제."""
    msg = await interaction.response.send_message(content)
    actual = await interaction.original_response()
    await asyncio.sleep(delay)
    try:
        await actual.delete()
    except Exception:
        pass

async def temp_followup(interaction, content, delay=5):
    """followup 메시지 보내고 delay초 후 자동 삭제."""
    msg = await interaction.followup.send(content)
    await asyncio.sleep(delay)
    try:
        await msg.delete()
    except Exception:
        pass


async def ensure_lavalink_lead_in(filepath, lead_in_ms):
    if lead_in_ms <= 0:
        return filepath
    src = Path(filepath)
    cache_dir = LAVALINK_PAD_CACHE_DIR
    try:
        cache_dir.mkdir(parents=True, exist_ok=True)
    except Exception as e:
        print(f"[LAVALINK] 선행 무음 패딩 캐시 생성 실패: {cache_dir}: {e}")
        return filepath
    file_stamp = f"{src.resolve()}::{src.stat().st_mtime_ns}::{lead_in_ms}"
    digest = hashlib.sha1(file_stamp.encode("utf-8")).hexdigest()[:16]
    safe_stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", unicodedata.normalize("NFC", src.stem)).strip("._") or "sound"
    output_path = cache_dir / f"{safe_stem}_lead_{lead_in_ms}_{digest}.wav"
    if output_path.exists():
        return str(output_path)

    process = await asyncio.create_subprocess_exec(
        "ffmpeg",
        "-y",
        "-i",
        str(src),
        "-af",
        f"adelay={lead_in_ms}:all=1",
        str(output_path),
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )
    _, stderr = await process.communicate()
    if process.returncode != 0:
        err = stderr.decode("utf-8", errors="ignore").strip()
        print(f"[LAVALINK] 선행 무음 패딩 생성 실패: {err}")
        return filepath
    return str(output_path)


async def play_lavalink_local_sound(player, filepath, volume, *, replace=True, lead_in_ms=0):
    guild = getattr(player, "guild", None)
    if guild is not None:
        connected_age = bot.get_voice_connected_age(guild.id)
        if connected_age is not None and connected_age < LAVALINK_POST_CONNECT_GRACE:
            await asyncio.sleep(LAVALINK_POST_CONNECT_GRACE - connected_age)
    playback_path = await ensure_lavalink_lead_in(filepath, lead_in_ms)
    identifier = build_lavalink_local_identifier(playback_path)
    await player.play(
        identifier,
        volume=max(0, min(1000, int(volume * 100))),
        replace=replace,
    )

async def play_sound_in_vc(guild, member, sound_name, bot_instance):
    if not member.voice or not member.voice.channel:
        return False, "음성 채널에 먼저 접속해주세요."

    key = unicodedata.normalize("NFC", sound_name.lower())
    if key not in bot_instance.sound_files:
        available = ', '.join(sorted(bot_instance.sound_files.keys()))
        return False, f"'{sound_name}' 사운드를 찾을 수 없습니다. 사용 가능: {available}"

    filepath = bot_instance.sound_files[key]
    bot_instance.set_voice_target(guild.id, member.voice.channel.id)
    log_voice_snapshot(
        "사운드 재생 요청",
        guild,
        guild.voice_client,
        extra=f"user={member.display_name} sound={key} target={format_channel_ref(member.voice.channel)}",
    )
    vc = await ensure_voice_connection(
        guild,
        member.voice.channel.id,
        move=True,
        reason="soundboard",
    )
    if vc is None:
        return False, "음성 채널 연결에 실패했습니다."

    if voice_is_playing(vc):
        await voice_stop(vc)

    if is_lavalink_player(vc):
        await play_lavalink_local_sound(
            vc,
            filepath,
            bot_instance.sound_volume,
            replace=True,
            lead_in_ms=LAVALINK_SHORT_CLIP_PAD_MS,
        )
        return True, f"**{sound_name}** 재생 중 (Lavalink, vol {int(bot_instance.sound_volume * 100)}%)"

    audio = discord.PCMVolumeTransformer(
        discord.FFmpegPCMAudio(filepath, options="-af loudnorm=I=-16:TP=-1.5:LRA=11"),
        volume=bot_instance.sound_volume
    )

    def after_play(error):
        status = f"error={error}" if error else "ok"
        log_voice_snapshot(
            f"사운드 재생 종료 [{key}]",
            guild,
            vc,
            extra=f"status={status}",
        )
        schedule_silence_keepalive_resume(
            guild.id,
            reason=f"soundboard:{key}",
        )

    vc.play(audio, after=after_play)
    return True, f"**{sound_name}** 재생 중 (vol {int(bot_instance.sound_volume * 100)}%)"

# ================= 2. 관리자 클래스 =================
class DataManager:
    @staticmethod
    def load_json(filename, default_data):
        if os.path.exists(filename):
            try:
                with open(filename, 'r', encoding='utf-8') as f: return json.load(f)
            except Exception as e:
                print(f"[WARN] JSON 로드 실패: {filename}: {e}")
        return default_data

    @staticmethod
    def save_json(filename, data):
        try:
            os.makedirs(os.path.dirname(os.path.abspath(filename)), exist_ok=True)
            tmp_filename = f"{filename}.tmp"
            with open(tmp_filename, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            os.replace(tmp_filename, filename)
        except Exception as e:
            print(f"[WARN] JSON 저장 실패: {filename}: {e}")


IMAGE_ACTION_SCHEMA = {
    "type": "object",
    "properties": {
        "action": {
            "type": "string",
            "enum": ["respond_text", "generate_image"],
        },
        "refined_prompt": {"type": "string"},
        "use_recent_images": {"type": "boolean"},
        "reason": {"type": "string"},
    },
    "required": ["action", "refined_prompt", "use_recent_images", "reason"],
}

IMAGE_PROMPT_SCHEMA = {
    "type": "object",
    "properties": {
        "prompt": {"type": "string"},
        "message": {"type": "string"},
    },
    "required": ["prompt"],
}


class DocumentCache:
    def __init__(self, cache_dir):
        self.cache_dir = cache_dir
        self.docs_dir = os.path.join(cache_dir, "docs")
        self.index_file = os.path.join(cache_dir, "index.json")
        os.makedirs(self.docs_dir, exist_ok=True)
        self.index = self._load_index()

    def _default_index(self):
        return {"version": 1, "next_id": 1, "docs": {}, "message_to_docs": {}}

    def _load_index(self):
        if os.path.exists(self.index_file):
            try:
                with open(self.index_file, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                if isinstance(data, dict):
                    data.setdefault("version", 1)
                    data.setdefault("next_id", 1)
                    data.setdefault("docs", {})
                    data.setdefault("message_to_docs", {})
                    return data
            except Exception:
                pass
        return self._default_index()

    def _save_index(self):
        with open(self.index_file, 'w', encoding='utf-8') as f:
            json.dump(self.index, f, indent=2, ensure_ascii=False)

    def _make_doc_id(self):
        doc_id = f"doc_{self.index['next_id']:06d}"
        self.index["next_id"] += 1
        return doc_id

    def _doc_path(self, doc_id):
        return os.path.join(self.docs_dir, f"{doc_id}.txt")

    def _write_doc(self, doc_id, content):
        with open(self._doc_path(doc_id), 'w', encoding='utf-8') as f:
            f.write(content)

    def store_document(self, doc_type, source_label, content, message_id, channel_id):
        message_key = str(message_id)
        docs_for_message = self.index["message_to_docs"].setdefault(message_key, [])
        now = datetime.now(timezone.utc).isoformat()

        for doc_id in docs_for_message:
            doc = self.index["docs"].get(doc_id)
            if doc and doc.get("type") == doc_type and doc.get("source_label") == source_label:
                self._write_doc(doc_id, content)
                doc.update({
                    "channel_id": str(channel_id),
                    "updated_at": now,
                    "size": len(content),
                })
                self._save_index()
                return doc_id

        doc_id = self._make_doc_id()
        self._write_doc(doc_id, content)
        self.index["docs"][doc_id] = {
            "id": doc_id,
            "type": doc_type,
            "source_label": source_label,
            "message_id": message_key,
            "channel_id": str(channel_id),
            "created_at": now,
            "updated_at": now,
            "size": len(content),
        }
        docs_for_message.append(doc_id)
        self._save_index()
        return doc_id

    def get_docs_for_message(self, message_id):
        message_key = str(message_id)
        result = []
        for doc_id in self.index["message_to_docs"].get(message_key, []):
            doc = self.index["docs"].get(doc_id)
            if doc:
                result.append(doc)
        return result

    def has_source(self, message_id, doc_type, source_label):
        for doc in self.get_docs_for_message(message_id):
            if doc.get("type") == doc_type and doc.get("source_label") == source_label:
                return True
        return False

    def get_content(self, doc_id):
        path = self._doc_path(doc_id)
        if not os.path.exists(path):
            return None
        try:
            with open(path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception:
            return None

    def render_blocks(self, message_id, full_content=False):
        blocks = []
        if not full_content:
            return blocks
        for doc in self.get_docs_for_message(message_id):
            doc_id = doc["id"]
            label = doc.get("source_label", "")
            content = self.get_content(doc_id)
            if content is None:
                continue
            source = "Webpage" if doc.get("type") == "webpage" else "Attached file"
            blocks.append(f'{source} "{label}":\n\n{content}')
        return blocks


def trim_document_text(text, limit=MAX_DOC_CHARS):
    cleaned = (text or "").strip()
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[:limit].rstrip()


def decode_text_bytes(raw_bytes):
    for encoding in ("utf-8", "cp949", "latin1"):
        try:
            return raw_bytes.decode(encoding)
        except Exception:
            continue
    return None


def extract_pdf_text(raw_bytes):
    if PdfReader is None:
        return None
    try:
        reader = PdfReader(io.BytesIO(raw_bytes))
        parts = []
        for page in reader.pages[:20]:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return trim_document_text("\n\n".join(part for part in parts if part.strip()))
    except Exception:
        return None


def extract_docx_text(raw_bytes):
    if DocxDocument is None:
        return None
    try:
        document = DocxDocument(io.BytesIO(raw_bytes))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
        return trim_document_text("\n".join(paragraphs))
    except Exception:
        return None


def extract_xlsx_text(raw_bytes):
    if load_workbook is None:
        return None
    try:
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        blocks = []
        for sheet in workbook.worksheets[:5]:
            rows = []
            for row in sheet.iter_rows(max_row=50, values_only=True):
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(f'Sheet "{sheet.title}":\n' + "\n".join(rows))
        workbook.close()
        return trim_document_text("\n\n".join(blocks))
    except Exception:
        return None


def extract_csv_text(raw_bytes, delimiter=","):
    decoded = decode_text_bytes(raw_bytes)
    if not decoded:
        return None
    try:
        reader = csv.reader(io.StringIO(decoded), delimiter=delimiter)
        rows = []
        for idx, row in enumerate(reader):
            if idx >= 100:
                break
            cells = [cell.strip() for cell in row if cell and cell.strip()]
            if cells:
                rows.append(" | ".join(cells))
        return trim_document_text("\n".join(rows))
    except Exception:
        return trim_document_text(decoded)


async def extract_attachment_text(att):
    filename = (att.filename or "").lower()
    raw_bytes = await att.read()

    if filename.endswith(".pdf"):
        return extract_pdf_text(raw_bytes), "pdf"
    if filename.endswith(".docx"):
        return extract_docx_text(raw_bytes), "docx"
    if filename.endswith(".xlsx"):
        return extract_xlsx_text(raw_bytes), "xlsx"
    if filename.endswith(".csv"):
        return extract_csv_text(raw_bytes, delimiter=","), "csv"
    if filename.endswith(".tsv"):
        return extract_csv_text(raw_bytes, delimiter="\t"), "tsv"
    if filename.endswith(TEXT_ATTACHMENT_EXTENSIONS):
        decoded = decode_text_bytes(raw_bytes)
        if decoded:
            return trim_document_text(decoded), "text"
    return None, None


def get_server_display_name(guild, user):
    user_id = getattr(user, "id", None)
    member = guild.get_member(user_id) if guild and user_id is not None else None
    if member is not None:
        return getattr(member, "display_name", None) or getattr(member, "name", None) or str(user_id)
    if guild and isinstance(user, discord.Member) and getattr(user.guild, "id", None) == guild.id:
        return getattr(user, "display_name", None) or getattr(user, "name", None) or str(user_id)
    if guild:
        return str(user_id or "unknown")
    return getattr(user, "display_name", None) or getattr(user, "name", None) or str(user_id or "unknown")


def coerce_json_object(text):
    if isinstance(text, dict):
        return text
    if not isinstance(text, str):
        return None

    candidate = text.strip()
    if not candidate:
        return None
    if "```" in candidate:
        match = re.search(r"```(?:json)?\s*(\{.*\})\s*```", candidate, re.DOTALL)
        if match:
            candidate = match.group(1).strip()

    start = candidate.find("{")
    end = candidate.rfind("}")
    if start >= 0 and end > start:
        candidate = candidate[start:end + 1]

    try:
        return json.loads(candidate)
    except Exception:
        return None


def build_recent_conversation_excerpt(history_texts, limit=MEMORY_EXCERPT_MAX_CHARS, window=8):
    lines = []
    for text in (history_texts or [])[-window:]:
        cleaned = (text or "").strip()
        if cleaned:
            lines.append(cleaned)
    excerpt = "\n\n".join(lines)
    if len(excerpt) > limit:
        excerpt = excerpt[-limit:].lstrip()
    return excerpt


def build_contextual_image_prompt(refined_prompt, current_request, history_texts):
    recent_context = build_recent_conversation_excerpt(history_texts, limit=4000)
    parts = [
        "Create an image for the CURRENT Discord request.",
        "Use recent conversation only to resolve references such as 'this', 'that', 'as above', mentioned characters, styles, or scenes.",
        "Do not create a generic educational diagram unless the current request explicitly asks for one.",
        "If the preferred visual prompt is generic or incomplete, make it specific using the current request and recent context.",
    ]
    if recent_context:
        parts.append(f"Recent conversation context:\n{recent_context}")
    parts.append(f"Current user request:\n{(current_request or '').strip()}")
    if refined_prompt:
        parts.append(f"Preferred visual prompt:\n{refined_prompt.strip()}")
    return "\n\n".join(parts)


def compact_prompt_for_status(prompt, limit=500):
    prompt = (prompt or "").strip()
    if len(prompt) <= limit:
        return prompt
    return prompt[:limit].rstrip() + "..."


def is_image_attachment(attachment):
    content_type = (getattr(attachment, "content_type", None) or "").lower()
    filename = (getattr(attachment, "filename", None) or "").lower()
    return content_type.startswith("image/") or filename.endswith(IMAGE_ATTACHMENT_EXTENSIONS)


def is_audio_or_video_attachment(attachment):
    content_type = (getattr(attachment, "content_type", None) or "").lower()
    return content_type.startswith("video/") or content_type.startswith("audio/")


def attachment_media_type(attachment):
    content_type = (getattr(attachment, "content_type", None) or "").lower()
    if "/" in content_type:
        return content_type.split("/", 1)[0]
    if is_image_attachment(attachment):
        return "image"
    return "file"


def infer_attachment_mime(attachment, detected_mime=None):
    if detected_mime:
        return detected_mime.split(";", 1)[0].strip()
    content_type = getattr(attachment, "content_type", None)
    if content_type:
        return content_type.split(";", 1)[0].strip()
    filename = (getattr(attachment, "filename", None) or "").lower()
    if filename.endswith((".jpg", ".jpeg")):
        return "image/jpeg"
    if filename.endswith(".webp"):
        return "image/webp"
    if filename.endswith(".gif"):
        return "image/gif"
    if filename.endswith(".png"):
        return "image/png"
    return "application/octet-stream"


def get_structured_helper_model(text_model=None):
    return ROUTING_MODEL


def has_image_generation_signal(text, *, has_current_images=False):
    text = (text or "").strip()
    if not text:
        return False
    if IMAGE_DRAW_REQUEST_RE.search(text):
        return True
    if IMAGE_ACTION_NOUN_RE.search(text) and IMAGE_ACTION_VERB_RE.search(text):
        return True
    if has_current_images and CURRENT_IMAGE_EDIT_RE.search(text):
        return True
    return False


def should_call_action_router(
    text,
    *,
    has_current_images=False,
    has_recent_images=False,
    allow_image_generation=True,
    allow_recent_images=True,
):
    if allow_image_generation and has_image_generation_signal(text, has_current_images=has_current_images):
        return True
    if allow_recent_images and has_recent_images and FOLLOWUP_IMAGE_EDIT_RE.search(text or ""):
        return True
    return False


def build_gemini_system_prompt():
    return SYSTEM_PROMPT


async def maybe_await(value):
    if inspect.isawaitable(value):
        return await value
    return value


def describe_gemini_exception(error, timeout_seconds=GEMINI_API_TIMEOUT_SECONDS):
    if isinstance(error, (TimeoutError, asyncio.TimeoutError)):
        return f"TimeoutError: Gemini API가 {timeout_seconds:.0f}초 안에 응답하지 않았습니다."
    message = str(error).strip()
    return f"{type(error).__name__}: {message}" if message else type(error).__name__


def get_interactions_api():
    if CLIENT is None:
        return None
    return getattr(getattr(CLIENT, "aio", None), "interactions", None) or getattr(CLIENT, "interactions", None)


def interaction_media_type(mime_type):
    mime_type = (mime_type or "").lower()
    if mime_type.startswith("image/"):
        return "image"
    if mime_type.startswith("audio/"):
        return "audio"
    if mime_type.startswith("video/"):
        return "video"
    if mime_type in {"application/pdf", "text/csv"}:
        return "document"
    return None


def encode_interaction_data(data):
    if isinstance(data, str):
        return data
    if isinstance(data, (bytes, bytearray, memoryview)):
        return base64.b64encode(bytes(data)).decode("ascii")
    return None


def gemini_part_to_interaction_content(part):
    text = getattr(part, "text", None)
    if text:
        return {"type": "text", "text": text}

    inline_data = getattr(part, "inline_data", None)
    if inline_data is not None:
        mime_type = getattr(inline_data, "mime_type", None)
        content_type = interaction_media_type(mime_type)
        data = encode_interaction_data(getattr(inline_data, "data", None))
        if content_type and data:
            return {"type": content_type, "mime_type": mime_type, "data": data}

    file_data = getattr(part, "file_data", None)
    if file_data is not None:
        mime_type = getattr(file_data, "mime_type", None)
        content_type = interaction_media_type(mime_type)
        uri = getattr(file_data, "file_uri", None)
        if content_type and uri:
            return {"type": content_type, "mime_type": mime_type, "uri": uri}
    return None


def gemini_contents_to_interaction_steps(contents):
    steps = []
    for content in contents or []:
        role = getattr(content, "role", None)
        step_type = "model_output" if role == "model" else "user_input"
        blocks = []
        for part in getattr(content, "parts", None) or []:
            block = gemini_part_to_interaction_content(part)
            if block:
                blocks.append(block)
        if blocks:
            steps.append({"type": step_type, "content": blocks})
    return steps


def merge_adjacent_contents(contents):
    merged = []
    for content in contents or []:
        if merged and merged[-1].role == content.role:
            merged[-1] = types.Content(
                role=content.role,
                parts=[*(merged[-1].parts or []), *(content.parts or [])],
            )
        else:
            merged.append(content)
    return merged


def label_discord_user_contents(entries):
    """Add natural speaker labels while keeping Discord IDs out of model input."""
    ordered_user_ids = []
    display_names = {}
    for content, user_id, display_name in entries:
        if content.role != "user" or user_id is None:
            continue
        if user_id not in display_names:
            ordered_user_ids.append(user_id)
        display_names[user_id] = (display_name or str(user_id)).strip()

    users_by_name = {}
    for user_id in ordered_user_ids:
        users_by_name.setdefault(display_names[user_id], []).append(user_id)

    speaker_labels = {}
    for display_name, user_ids in users_by_name.items():
        for index, user_id in enumerate(user_ids, start=1):
            speaker_labels[user_id] = (
                display_name if len(user_ids) == 1 else f"{display_name} ({index})"
            )

    labeled = []
    for content, user_id, _display_name in entries:
        if content.role == "user" and user_id in speaker_labels:
            content = types.Content(
                role=content.role,
                parts=[types.Part(text=f"{speaker_labels[user_id]}:\n"), *(content.parts or [])],
            )
        labeled.append(content)
    return labeled


def image_generation_inputs_to_blocks(contents):
    blocks = []
    for item in contents or []:
        if isinstance(item, str):
            blocks.append({"type": "text", "text": item})
            continue
        if isinstance(item, Image.Image):
            image_buffer = io.BytesIO()
            item.save(image_buffer, format="PNG")
            blocks.append(
                {
                    "type": "image",
                    "mime_type": "image/png",
                    "data": base64.b64encode(image_buffer.getvalue()).decode("ascii"),
                }
            )
    return blocks


async def create_model_interaction(
    *,
    model,
    input_data,
    system_instruction=None,
    tools=None,
    thinking_level=None,
    response_schema=None,
    response_format=None,
    store=False,
    timeout_seconds=GEMINI_API_TIMEOUT_SECONDS,
):
    interactions_api = get_interactions_api()
    if interactions_api is None:
        raise RuntimeError("현재 google-genai SDK에서 Interactions API를 찾을 수 없습니다. 패키지 업데이트가 필요합니다.")

    kwargs = {"model": model, "input": input_data, "store": store}
    if system_instruction:
        kwargs["system_instruction"] = system_instruction
    if tools:
        kwargs["tools"] = tools
    if thinking_level is not None:
        kwargs["generation_config"] = {"thinking_level": thinking_level}
    if response_schema is not None:
        kwargs["response_format"] = [
            {"type": "text", "mime_type": "application/json", "schema": response_schema}
        ]
    elif response_format is not None:
        kwargs["response_format"] = response_format
    return await asyncio.wait_for(
        maybe_await(interactions_api.create(timeout=timeout_seconds, **kwargs)),
        timeout=timeout_seconds,
    )


async def call_interaction_get(interactions_api, interaction_id):
    try:
        return await maybe_await(interactions_api.get(interaction_id))
    except TypeError:
        return await maybe_await(interactions_api.get(id=interaction_id))


def get_final_interaction_text_blocks(interaction):
    output_steps = []
    collecting = False
    steps = get_interaction_output_value(interaction, "steps", None) or []
    for step in reversed(steps):
        step_type = get_interaction_output_value(step, "type", "")
        if step_type == "user_input":
            break
        if step_type == "model_output":
            collecting = True
            output_steps.append(step)
        elif collecting:
            break

    blocks = []
    for step in reversed(output_steps):
        for block in get_interaction_output_value(step, "content", None) or []:
            if get_interaction_output_value(block, "type", "") == "text":
                blocks.append(block)
    return blocks


def collect_interaction_url_citations(text_blocks):
    citations = []
    seen_urls = set()
    for block in text_blocks:
        for annotation in get_interaction_output_value(block, "annotations", None) or []:
            annotation_type = str(get_interaction_output_value(annotation, "type", "")).lower()
            if not annotation_type.endswith("url_citation"):
                continue
            url = str(get_interaction_output_value(annotation, "url", "") or "").strip()
            parsed = urlsplit(url)
            if parsed.scheme not in {"http", "https"} or not parsed.netloc or url in seen_urls:
                continue
            title = re.sub(
                r"\s+",
                " ",
                str(get_interaction_output_value(annotation, "title", "") or ""),
            ).strip()
            citations.append({"title": title or parsed.netloc, "url": url})
            seen_urls.add(url)
    return citations


def format_interaction_citations(citations):
    lines = []
    for index, citation in enumerate(citations, start=1):
        title = (
            citation["title"]
            .replace("\\", "\\\\")
            .replace("[", "\\[")
            .replace("]", "\\]")
        )
        url = (
            citation["url"]
            .replace("\\", "%5C")
            .replace(" ", "%20")
            .replace("(", "%28")
            .replace(")", "%29")
        )
        lines.append(f"{index}. [{title}]({url})")
    return "**출처:**\n" + "\n".join(lines)


def extract_interaction_text(interaction):
    text_blocks = get_final_interaction_text_blocks(interaction)
    if text_blocks:
        raw_text = "".join(
            get_interaction_output_value(block, "text", "") or ""
            for block in text_blocks
        )
    else:
        output_text = get_interaction_output_value(interaction, "output_text", "")
        raw_text = output_text if isinstance(output_text, str) else ""

    text, marker_count = RAW_CITATION_MARKER_RE.subn("", raw_text)
    text = text.strip()
    citations = [
        citation
        for citation in collect_interaction_url_citations(text_blocks)
        if citation["url"] not in text
    ]
    if marker_count and not citations:
        print(f"[WARN] Gemini citation markers had no URL annotations: count={marker_count}")
    if citations:
        citation_text = format_interaction_citations(citations)
        return f"{text}\n\n{citation_text}" if text else citation_text
    return text


def get_interaction_usage(interaction):
    usage = get_interaction_output_value(interaction, "usage")
    return int(get_interaction_output_value(usage, "total_tokens", 0) or 0)


def get_interaction_status(interaction):
    return get_interaction_output_value(interaction, "status", "completed") or "completed"


def get_interaction_output_value(output, key, default=None):
    if isinstance(output, dict):
        return output.get(key, default)
    return getattr(output, key, default)


def summarize_interaction_request(
    input_data,
    *,
    model,
    thinking_level,
    system_instruction,
    tools,
    store,
):
    """Summarize request shape without logging message text or media payloads."""
    step_counts = {}
    block_counts = {}
    text_chars = 0
    inline_media_bytes = 0
    uri_media = 0

    for step in input_data or []:
        step_type = str(get_interaction_output_value(step, "type", "unknown"))
        step_counts[step_type] = step_counts.get(step_type, 0) + 1
        for block in get_interaction_output_value(step, "content", None) or []:
            block_type = str(get_interaction_output_value(block, "type", "unknown"))
            block_counts[block_type] = block_counts.get(block_type, 0) + 1
            text_chars += len(get_interaction_output_value(block, "text", "") or "")

            data = get_interaction_output_value(block, "data")
            if isinstance(data, str):
                payload = data.split(",", 1)[1] if data.startswith("data:") and "," in data else data
                inline_media_bytes += max(0, (len(payload) * 3) // 4 - payload[-2:].count("="))
            elif isinstance(data, (bytes, bytearray, memoryview)):
                inline_media_bytes += len(data)
            if get_interaction_output_value(block, "uri"):
                uri_media += 1

    tool_types = [
        str(get_interaction_output_value(tool, "type", type(tool).__name__))
        for tool in (tools or [])
    ]
    return {
        "model": model,
        "thinking_level": thinking_level or "model_default",
        "store": bool(store),
        "system_chars": len(system_instruction or ""),
        "steps": step_counts,
        "blocks": block_counts,
        "stateless_model_history_without_thought_steps": bool(
            step_counts.get("model_output") and not step_counts.get("thought")
        ),
        "text_chars": text_chars,
        "inline_media_bytes": inline_media_bytes,
        "uri_media": uri_media,
        "tools": tool_types,
    }


async def create_traced_interaction(
    call_trace,
    *,
    phase,
    model,
    input_data,
    system_instruction,
    thinking_level,
    tools=None,
):
    request_summary = summarize_interaction_request(
        input_data,
        model=model,
        thinking_level=thinking_level,
        system_instruction=system_instruction,
        tools=tools,
        store=False,
    )
    started = time_module.perf_counter()
    print(
        f"[GEMINI] {phase} request start "
        f"request={json.dumps(request_summary, ensure_ascii=False, default=str, separators=(',', ':'))}"
    )
    try:
        interaction = await create_model_interaction(
            model=model,
            input_data=input_data,
            system_instruction=system_instruction,
            tools=tools,
            thinking_level=thinking_level,
        )
    except Exception as e:
        elapsed_ms = round((time_module.perf_counter() - started) * 1000)
        call_trace["calls"].append(
            {
                "phase": phase,
                "elapsed_ms": elapsed_ms,
                "request": request_summary,
            }
        )
        print(
            f"[ERROR] Gemini {phase} interaction failed "
            f"trace={json.dumps(call_trace, ensure_ascii=False, default=str, separators=(',', ':'))} "
            f"exception={describe_gemini_exception(e)}"
        )
        raise

    elapsed_ms = round((time_module.perf_counter() - started) * 1000)
    call_trace["calls"].append(
        {
            "phase": phase,
            "elapsed_ms": elapsed_ms,
            "request": request_summary,
        }
    )
    print(
        f"[GEMINI] {phase} request complete elapsed_ms={elapsed_ms} "
        f"status={get_interaction_status(interaction)} id={get_interaction_output_value(interaction, 'id', '-') or '-'}"
    )
    return interaction


def get_interaction_present_fields(interaction):
    if isinstance(interaction, dict):
        return sorted(str(key) for key in interaction)

    fields = getattr(interaction, "model_fields_set", None)
    if fields is None:
        fields = getattr(interaction, "__fields_set__", None)
    return sorted(str(key) for key in (fields or []))


def classify_content_free_interaction(interaction):
    steps = get_interaction_output_value(interaction, "steps", None) or []
    usage = get_interaction_output_value(interaction, "usage")
    output_tokens = int(get_interaction_output_value(usage, "total_output_tokens", 0) or 0)
    thought_tokens = int(get_interaction_output_value(usage, "total_thought_tokens", 0) or 0)

    if not steps:
        if thought_tokens and not output_tokens:
            return "thought_tokens_without_output_steps"
        return "no_output_steps"

    model_outputs = [
        step for step in steps
        if get_interaction_output_value(step, "type", "") == "model_output"
    ]
    if not model_outputs:
        return "no_model_output_step"
    if any(get_interaction_output_value(step, "error") for step in model_outputs):
        return "model_output_error"
    return "model_output_without_text"


def get_interaction_error_message(interaction):
    error = get_interaction_output_value(interaction, "error")
    steps = get_interaction_output_value(interaction, "steps", None) or []
    if not error:
        for step in reversed(steps):
            error = get_interaction_output_value(step, "error")
            if error:
                break
    if not error:
        return ""
    message = get_interaction_output_value(error, "message")
    code = get_interaction_output_value(error, "code")
    if message and code:
        return f"{code}: {message}"
    return str(message or code or error)


def interaction_step_to_dict(step):
    if isinstance(step, dict):
        return step
    model_dump = getattr(step, "model_dump", None)
    if callable(model_dump):
        return model_dump(exclude_none=True, mode="json")
    to_dict = getattr(step, "dict", None)
    if callable(to_dict):
        return to_dict(exclude_none=True)
    return None


def collect_search_continuation_steps(interaction):
    """Collect a completed search turn that is missing its final model output."""
    continuation_steps = []
    has_successful_result = False
    for step in get_interaction_output_value(interaction, "steps", None) or []:
        step_type = get_interaction_output_value(step, "type", "")
        if step_type not in {"thought", "google_search_call", "google_search_result"}:
            continue
        serialized = interaction_step_to_dict(step)
        if serialized:
            continuation_steps.append(serialized)
        if step_type == "google_search_result" and not get_interaction_output_value(step, "is_error", False):
            has_successful_result = True
    return continuation_steps if has_successful_result else []


def summarize_interaction_response(interaction):
    """Return a content-free diagnostic summary for failed/empty interactions."""
    steps = get_interaction_output_value(interaction, "steps", None) or []
    step_summaries = []
    for step in steps:
        step_type = str(get_interaction_output_value(step, "type", "unknown"))
        summary = {"type": step_type}

        if step_type == "model_output":
            content = get_interaction_output_value(step, "content", None) or []
            summary["content"] = [
                {
                    "type": str(get_interaction_output_value(item, "type", "unknown")),
                    "text_chars": len(get_interaction_output_value(item, "text", "") or ""),
                }
                for item in content
            ]
            error = get_interaction_output_value(step, "error")
            if error:
                summary["error"] = {
                    "code": get_interaction_output_value(error, "code"),
                    "message": get_interaction_output_value(error, "message"),
                    "details": get_interaction_output_value(error, "details"),
                }
        elif step_type == "thought":
            thought_summary = get_interaction_output_value(step, "summary", None) or []
            summary["summary_chars"] = sum(
                len(get_interaction_output_value(item, "text", "") or "")
                for item in thought_summary
            )
        elif step_type.endswith("_result"):
            is_error = get_interaction_output_value(step, "is_error")
            if is_error is not None:
                summary["is_error"] = bool(is_error)

        step_summaries.append(summary)

    usage = get_interaction_output_value(interaction, "usage")
    usage_summary = {
        key: get_interaction_output_value(usage, key)
        for key in (
            "total_input_tokens",
            "total_output_tokens",
            "total_thought_tokens",
            "total_tool_use_tokens",
            "total_tokens",
        )
        if get_interaction_output_value(usage, key) is not None
    }
    diagnostic = {
        "id": get_interaction_output_value(interaction, "id"),
        "status": get_interaction_status(interaction),
        "model": get_interaction_output_value(interaction, "model"),
        "empty_kind": classify_content_free_interaction(interaction),
        "present_fields": get_interaction_present_fields(interaction),
        "created": get_interaction_output_value(interaction, "created"),
        "updated": get_interaction_output_value(interaction, "updated"),
        "service_tier": get_interaction_output_value(interaction, "service_tier"),
        "response_modalities": get_interaction_output_value(interaction, "response_modalities"),
        "steps": step_summaries,
        "usage": usage_summary,
    }
    return json.dumps(diagnostic, ensure_ascii=False, default=str, separators=(",", ":"))


def decode_interaction_image(data):
    if not data:
        return None
    if isinstance(data, bytes):
        return data
    if not isinstance(data, str):
        return None
    payload = data.split(",", 1)[1] if data.startswith("data:") and "," in data else data
    try:
        return base64.b64decode(payload)
    except Exception:
        return None


def collect_interaction_outputs(interaction):
    """Interaction 응답에서 텍스트/이미지 블록을 순서대로 수집.

    google-genai 2.x의 Interactions API는 결과를 `outputs`가 아니라
    `steps[].content[]`(type=model_output)로 반환한다. thought/tool 스텝은 건너뛰고
    model_output 스텝의 content 블록만 모은다."""
    steps = (
        interaction.get("steps")
        if isinstance(interaction, dict)
        else getattr(interaction, "steps", None)
    ) or []
    blocks = []
    for step in steps:
        if get_interaction_output_value(step, "type", "") != "model_output":
            continue
        blocks.extend(get_interaction_output_value(step, "content", None) or [])
    collected = []
    for output in blocks:
        output_type = get_interaction_output_value(output, "type", "")
        text = get_interaction_output_value(output, "text")
        data = get_interaction_output_value(output, "data")
        mime_type = get_interaction_output_value(output, "mime_type") or get_interaction_output_value(output, "mimeType")
        if output_type == "text" and text:
            collected.append({"type": "text", "text": text})
        elif output_type == "image" and data:
            image_bytes = decode_interaction_image(data)
            if image_bytes:
                collected.append({"type": "image", "data": image_bytes, "mime_type": mime_type or "image/png"})
        elif text:
            collected.append({"type": "text", "text": text})
    return collected


DEEP_RESEARCH_HTML_ALLOWED_TAGS = {
    "a", "blockquote", "br", "code", "dd", "del", "div", "dl", "dt", "em",
    "h1", "h2", "h3", "h4", "h5", "h6", "hr", "i", "img", "li", "ol",
    "p", "pre", "s", "span", "strong", "sub", "sup", "table", "tbody", "td",
    "tfoot", "th", "thead", "tr", "ul",
}
DEEP_RESEARCH_HTML_DANGEROUS_TAGS = {
    "base", "button", "embed", "form", "frame", "frameset", "iframe", "input",
    "link", "meta", "object", "script", "select", "style", "textarea",
}


def prepare_deep_research_image(raw, mime_type):
    """해상도 손실 없이 PNG를 더 작은 lossless WebP로 바꿀 수 있을 때만 변환한다."""
    detected_mime = mime_type or "application/octet-stream"
    try:
        with Image.open(io.BytesIO(raw)) as image:
            format_mime = {
                "GIF": "image/gif",
                "JPEG": "image/jpeg",
                "PNG": "image/png",
                "WEBP": "image/webp",
            }
            detected_mime = format_mime.get(image.format, detected_mime)
            if image.format != "PNG" or getattr(image, "is_animated", False):
                return raw, detected_mime

            image.load()
            converted = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            webp_buffer = io.BytesIO()
            save_kwargs = {"format": "WEBP", "lossless": True, "method": 6}
            if image.info.get("icc_profile"):
                save_kwargs["icc_profile"] = image.info["icc_profile"]
            converted.save(webp_buffer, **save_kwargs)
            webp = webp_buffer.getvalue()
            if len(webp) < len(raw):
                print(f"[DeepResearch HTML] lossless WebP: {len(raw):,} -> {len(webp):,} bytes")
                return webp, "image/webp"
    except Exception as e:
        print(f"[DeepResearch HTML] image optimization skipped: {e}")
    return raw, detected_mime


def sanitize_deep_research_html(fragment):
    """Markdown 결과에서 실행 가능 콘텐츠와 외부 이미지 요청을 제거한다."""
    soup = BeautifulSoup(fragment, "html.parser")
    for tag in list(soup.find_all(True)):
        name = tag.name.lower()
        if name in DEEP_RESEARCH_HTML_DANGEROUS_TAGS:
            tag.decompose()
            continue
        if name not in DEEP_RESEARCH_HTML_ALLOWED_TAGS:
            tag.unwrap()
            continue
        if name == "img":
            src = (tag.get("src") or "").strip()
            alt = (tag.get("alt") or "외부 이미지").strip()
            note = soup.new_tag("span")
            note["class"] = ["external-image-note"]
            if urlsplit(src).scheme in ("http", "https"):
                link = soup.new_tag("a", href=src, target="_blank", rel="noopener noreferrer")
                link.string = f"외부 이미지 링크: {alt}"
                note.append(link)
            else:
                note.string = f"외부 이미지 제외: {alt}"
            tag.replace_with(note)
            continue

        clean_attrs = {}
        raw_id = (tag.get("id") or "").strip()
        if raw_id and name in ("div", "li", "p", "span", "sup"):
            clean_attrs["id"] = re.sub(r"[^\w:.-]+", "-", raw_id)[:160]
        if name == "a":
            href = (tag.get("href") or "").strip()
            scheme = urlsplit(href).scheme.lower()
            if href.startswith("#") or scheme in ("http", "https", "mailto"):
                clean_attrs["href"] = href
            if scheme in ("http", "https", "mailto"):
                clean_attrs["target"] = "_blank"
                clean_attrs["rel"] = "noopener noreferrer"
            allowed_classes = {
                value for value in (tag.get("class") or [])
                if value in ("footnote-ref", "footnote-backref")
            }
            if allowed_classes:
                clean_attrs["class"] = sorted(allowed_classes)
        elif name == "code":
            language_classes = [
                value for value in (tag.get("class") or [])
                if re.fullmatch(r"language-[A-Za-z0-9_+-]+", value)
            ]
            if language_classes:
                clean_attrs["class"] = language_classes
        elif name == "div" and "footnote" in (tag.get("class") or []):
            clean_attrs["class"] = ["footnote"]
        elif name in ("th", "td"):
            alignment = (tag.get("align") or "").lower()
            style = tag.get("style") or ""
            style_match = re.fullmatch(r"\s*text-align\s*:\s*(left|center|right)\s*;?\s*", style)
            if not alignment and style_match:
                alignment = style_match.group(1)
            if alignment in ("left", "center", "right"):
                clean_attrs["class"] = [f"align-{alignment}"]
        elif name == "ol" and str(tag.get("start") or "").isdigit():
            clean_attrs["start"] = str(tag["start"])
        tag.attrs = clean_attrs
    return soup


def build_deep_research_html(outputs, *, agent_id, title):
    try:
        import markdown as markdown_module
    except Exception as e:
        raise RuntimeError("Markdown 패키지가 설치되어 있지 않습니다.") from e

    markdown_parts = []
    embedded_images = []
    text_char_count = 0
    for output in outputs:
        if output["type"] == "text":
            text = output.get("text", "")
            markdown_parts.append(text)
            text_char_count += len(text)
        elif output["type"] == "image" and output.get("data"):
            index = len(embedded_images) + 1
            raw = output["data"]
            token = f"DEEPRESEARCHIMAGE{index:04d}{hashlib.sha1(raw).hexdigest()[:12]}PLACEHOLDER"
            markdown_parts.append(f"\n\n{token}\n\n")
            embedded_images.append((token, raw, output.get("mime_type")))

    renderer = markdown_module.Markdown(
        extensions=["extra", "sane_lists"],
        output_format="html5",
    )
    soup = sanitize_deep_research_html(renderer.convert("\n\n".join(markdown_parts)))

    for index, (token, raw, mime_type) in enumerate(embedded_images, start=1):
        placeholder = soup.find(string=lambda value: isinstance(value, str) and value.strip() == token)
        if placeholder is None:
            continue
        prepared, prepared_mime = prepare_deep_research_image(raw, mime_type)
        encoded = base64.b64encode(prepared).decode("ascii")
        figure = soup.new_tag("figure")
        figure["class"] = ["report-figure"]
        image_tag = soup.new_tag(
            "img",
            src=f"data:{prepared_mime};base64,{encoded}",
            alt=f"Deep Research 시각 자료 {index}",
            loading="lazy",
        )
        image_tag["tabindex"] = "0"
        image_tag["role"] = "button"
        image_tag["aria-label"] = f"시각 자료 {index} 크게 보기"
        caption = soup.new_tag("figcaption")
        caption.string = f"시각 자료 {index}"
        figure.extend([image_tag, caption])
        parent = placeholder.parent
        if parent is not None and parent.name == "p" and parent.get_text(strip=True) == token:
            parent.replace_with(figure)
        else:
            placeholder.replace_with(figure)

    table_count = 0
    for table in list(soup.find_all("table")):
        table_count += 1
        first_row = table.find("tr")
        column_count = len(first_row.find_all(("th", "td"), recursive=False)) if first_row else 0
        table["class"] = ["data-table"] + (["wide-table"] if column_count >= 7 else [])
        table["data-sortable"] = "true"
        wrapper = soup.new_tag("div")
        wrapper["class"] = ["table-wrap"]
        wrapper["tabindex"] = "0"
        table.wrap(wrapper)

    used_ids = set()
    toc_items = []
    for index, heading in enumerate(soup.find_all(("h1", "h2", "h3")), start=1):
        label = heading.get_text(" ", strip=True) or f"Section {index}"
        base_id = re.sub(r"[^\w-]+", "-", label.lower()).strip("-") or f"section-{index}"
        heading_id = base_id
        suffix = 2
        while heading_id in used_ids:
            heading_id = f"{base_id}-{suffix}"
            suffix += 1
        used_ids.add(heading_id)
        heading["id"] = heading_id
        level = int(heading.name[1])
        toc_items.append(
            f'<li class="toc-level-{level}"><a href="#{html_module.escape(heading_id, quote=True)}">'
            f'{html_module.escape(label)}</a></li>'
        )

    content = soup.new_tag("div")
    content["class"] = ["report-content"]
    current_section = None
    section_index = 0
    for child in list(soup.contents):
        child_name = getattr(child, "name", None)
        if current_section is None or child_name in ("h1", "h2"):
            section_index += 1
            current_section = soup.new_tag("section")
            current_section["class"] = ["report-section"]
            current_section["data-section"] = str(section_index)
            content.append(current_section)
        current_section.append(child.extract())

    safe_title = html_module.escape((title or "Deep Research Report").strip())
    safe_agent = html_module.escape(agent_id)
    generated_at = html_module.escape(datetime.now(KST).strftime("%Y-%m-%d %H:%M JST"))
    language = "ko" if re.search(r"[가-힣]", title or "") else "en"
    toc_html = "".join(toc_items) or '<li class="toc-empty">목차가 없습니다.</li>'
    body_html = str(content)
    image_count = len(embedded_images)

    return (
        '<!doctype html>\n'
        f'<html lang="{language}">\n<head>\n'
        '<meta charset="utf-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; img-src data:; '
        'style-src \'unsafe-inline\'; script-src \'unsafe-inline\'; connect-src \'none\'; '
        'font-src data:; object-src \'none\'; base-uri \'none\'; form-action \'none\'">\n'
        f'<title>{safe_title} · Deep Research</title>\n'
        """<style>
:root {
  color-scheme: light;
  --canvas: #e9eee9;
  --paper: #fffdf8;
  --paper-2: #f4f2e9;
  --ink: #172522;
  --muted: #65736e;
  --line: #d8ded8;
  --accent: #0d6b5d;
  --accent-strong: #084b42;
  --accent-soft: #d9ece6;
  --warm: #c56a3a;
  --code: #13201e;
  --shadow: 0 24px 70px rgba(27, 48, 42, .16);
}
:root[data-theme="dark"] {
  color-scheme: dark;
  --canvas: #111816;
  --paper: #17201e;
  --paper-2: #1f2a27;
  --ink: #eef4f0;
  --muted: #a9b8b2;
  --line: #33433e;
  --accent: #63c5b3;
  --accent-strong: #8ed9cb;
  --accent-soft: #203d37;
  --warm: #f09a68;
  --code: #0c1211;
  --shadow: 0 24px 70px rgba(0, 0, 0, .38);
}
* { box-sizing: border-box; }
html { scroll-behavior: smooth; }
body {
  margin: 0;
  background:
    radial-gradient(circle at 8% 4%, color-mix(in srgb, var(--accent) 18%, transparent), transparent 28rem),
    linear-gradient(145deg, var(--canvas), color-mix(in srgb, var(--canvas) 84%, #fff));
  color: var(--ink);
  font-family: "Noto Sans CJK KR", "Noto Sans KR", "Apple SD Gothic Neo", "Malgun Gothic", sans-serif;
  line-height: 1.72;
  overflow-wrap: anywhere;
}
button, input { font: inherit; }
a { color: var(--accent); text-underline-offset: .18em; }
a:hover { color: var(--warm); }
.toolbar {
  position: sticky;
  top: 0;
  z-index: 30;
  display: flex;
  gap: .55rem;
  align-items: center;
  padding: .7rem max(1rem, calc((100vw - 1180px) / 2));
  border-bottom: 1px solid color-mix(in srgb, var(--line) 78%, transparent);
  background: color-mix(in srgb, var(--paper) 88%, transparent);
  backdrop-filter: blur(16px);
}
.toolbar input {
  min-width: 12rem;
  flex: 1;
  padding: .58rem .8rem;
  color: var(--ink);
  border: 1px solid var(--line);
  border-radius: .65rem;
  background: var(--paper);
}
.toolbar button, .section-toggle {
  padding: .5rem .72rem;
  color: var(--ink);
  border: 1px solid var(--line);
  border-radius: .65rem;
  background: var(--paper);
  cursor: pointer;
}
.toolbar button:hover, .section-toggle:hover { border-color: var(--accent); color: var(--accent); }
.document {
  width: min(1180px, calc(100% - 2rem));
  margin: 2rem auto 5rem;
  overflow: hidden;
  border: 1px solid color-mix(in srgb, var(--line) 82%, transparent);
  border-radius: 1.35rem;
  background: var(--paper);
  box-shadow: var(--shadow);
}
.cover {
  position: relative;
  isolation: isolate;
  min-height: 25rem;
  padding: clamp(2.2rem, 6vw, 5.5rem);
  display: flex;
  flex-direction: column;
  justify-content: flex-end;
  overflow: hidden;
  border-bottom: 1px solid var(--line);
  background:
    linear-gradient(120deg, color-mix(in srgb, var(--accent-soft) 86%, var(--paper)), var(--paper-2));
}
.cover::before {
  content: "";
  position: absolute;
  z-index: -1;
  width: 26rem;
  height: 26rem;
  right: -7rem;
  top: -10rem;
  border: 5rem solid color-mix(in srgb, var(--accent) 20%, transparent);
  border-radius: 50%;
}
.eyebrow { margin: 0 0 1rem; color: var(--accent); font-size: .76rem; font-weight: 800; letter-spacing: .18em; text-transform: uppercase; }
.cover h1 { max-width: 18ch; margin: 0; font-size: clamp(2.1rem, 5.4vw, 4.8rem); line-height: 1.08; letter-spacing: -.045em; }
.cover-meta { margin-top: 2rem; color: var(--muted); font-size: .88rem; }
.stat-row { display: flex; flex-wrap: wrap; gap: .65rem; margin-top: 1.2rem; }
.stat { padding: .42rem .68rem; border: 1px solid color-mix(in srgb, var(--accent) 25%, var(--line)); border-radius: 999px; background: color-mix(in srgb, var(--paper) 72%, transparent); font-size: .78rem; }
.report-grid { display: grid; grid-template-columns: 15.5rem minmax(0, 1fr); gap: clamp(1.5rem, 4vw, 4rem); padding: clamp(1.5rem, 5vw, 4.5rem); }
.toc { align-self: start; position: sticky; top: 5rem; max-height: calc(100vh - 7rem); overflow: auto; padding-right: .75rem; }
.toc h2 { margin: 0 0 .8rem; font-size: .78rem; letter-spacing: .13em; text-transform: uppercase; color: var(--muted); }
.toc ol { list-style: none; padding: 0; margin: 0; }
.toc li { margin: .18rem 0; }
.toc a { display: block; padding: .34rem .5rem; border-left: 2px solid transparent; color: var(--muted); text-decoration: none; font-size: .82rem; line-height: 1.35; }
.toc a:hover, .toc a.active { color: var(--accent); border-left-color: var(--accent); background: var(--accent-soft); }
.toc-level-2 a { padding-left: 1rem; }
.toc-level-3 a { padding-left: 1.55rem; font-size: .76rem; }
.report-grid > article { min-width: 0; }
.report-content { min-width: 0; }
.report-section { position: relative; margin: 0 0 2.8rem; scroll-margin-top: 5rem; }
.report-section[hidden] { display: none; }
.report-section.collapsed > :not(h1):not(h2) { display: none; }
.section-toggle { float: right; margin-left: .75rem; padding: .24rem .5rem; font-size: .7rem; font-weight: 500; letter-spacing: 0; }
.report-content h1, .report-content h2, .report-content h3 { color: var(--ink); text-wrap: balance; }
.report-content h1 { margin: 0 0 1.4rem; padding-bottom: .65rem; border-bottom: 3px solid var(--accent); font-size: clamp(1.8rem, 3.5vw, 2.7rem); line-height: 1.2; letter-spacing: -.035em; }
.report-content h2 { margin: 0 0 1.2rem; font-size: clamp(1.35rem, 2.3vw, 1.85rem); line-height: 1.28; letter-spacing: -.02em; }
.report-content h3 { margin: 2rem 0 .75rem; color: var(--accent-strong); font-size: 1.12rem; }
.report-content h4 { margin: 1.5rem 0 .5rem; }
.report-content p { margin: .65rem 0 1rem; }
.report-content strong { color: var(--accent-strong); }
.report-content blockquote { margin: 1.4rem 0; padding: .9rem 1.15rem; border-left: .28rem solid var(--warm); background: var(--paper-2); color: color-mix(in srgb, var(--ink) 83%, var(--muted)); }
.report-content ul, .report-content ol { padding-left: 1.35rem; }
.report-content li { margin: .35rem 0; }
pre { overflow: auto; padding: 1rem 1.1rem; border-radius: .8rem; background: var(--code); color: #e8f1ed; font: .83rem/1.55 ui-monospace, SFMono-Regular, Consolas, monospace; }
code { padding: .12em .32em; border-radius: .28rem; background: var(--paper-2); font-family: ui-monospace, SFMono-Regular, Consolas, monospace; font-size: .9em; }
pre code { padding: 0; background: transparent; }
.table-wrap { max-width: 100%; overflow: auto; margin: 1.35rem 0 1.8rem; border: 1px solid var(--line); border-radius: .85rem; background: var(--paper); }
.data-table { width: 100%; border-collapse: collapse; font-size: .85rem; line-height: 1.45; }
.data-table th, .data-table td { min-width: 8rem; padding: .68rem .75rem; border-bottom: 1px solid var(--line); border-right: 1px solid color-mix(in srgb, var(--line) 62%, transparent); vertical-align: top; }
.data-table th { position: sticky; top: 0; z-index: 1; color: var(--paper); background: var(--accent-strong); text-align: left; cursor: pointer; }
.data-table tbody tr:nth-child(even) { background: var(--paper-2); }
.data-table tbody tr:hover { background: var(--accent-soft); }
.data-table .align-center { text-align: center; }
.data-table .align-right { text-align: right; font-variant-numeric: tabular-nums; }
.wide-table { min-width: 64rem; font-size: .78rem; }
.report-figure { margin: 1.8rem 0 2.2rem; text-align: center; }
.report-figure img { display: block; width: auto; max-width: 100%; max-height: 78vh; margin: 0 auto; border: 1px solid var(--line); border-radius: .85rem; background: #fff; box-shadow: 0 12px 32px rgba(20, 45, 38, .14); cursor: zoom-in; }
.report-figure figcaption { margin-top: .6rem; color: var(--muted); font-size: .78rem; }
.external-image-note { display: inline-block; padding: .35rem .55rem; border: 1px dashed var(--line); border-radius: .45rem; color: var(--muted); font-size: .82rem; }
.footnote { margin-top: 3rem; padding-top: 1rem; border-top: 1px solid var(--line); color: var(--muted); font-size: .8rem; }
.no-results { display: none; padding: 3rem; text-align: center; color: var(--muted); }
.no-results.visible { display: block; }
dialog { width: min(96vw, 1400px); max-height: 94vh; padding: 0; border: 0; border-radius: 1rem; background: transparent; }
dialog::backdrop { background: rgba(4, 12, 10, .86); backdrop-filter: blur(6px); }
.lightbox-inner { position: relative; padding: 2.5rem 1rem 1rem; text-align: center; }
.lightbox-inner img { max-width: 100%; max-height: 86vh; border-radius: .7rem; background: #fff; }
.lightbox-close { position: absolute; top: 0; right: .8rem; border: 0; color: #fff; background: transparent; font-size: 1.8rem; cursor: pointer; }
@media (max-width: 820px) {
  .toolbar { flex-wrap: wrap; }
  .toolbar input { flex-basis: 100%; order: 2; }
  .document { width: min(100% - 1rem, 1180px); margin-top: .5rem; border-radius: .85rem; }
  .cover { min-height: 21rem; }
  .report-grid { grid-template-columns: minmax(0, 1fr); }
  .toc { position: static; max-height: none; padding: 1rem; border: 1px solid var(--line); border-radius: .8rem; background: var(--paper-2); }
}
@media print {
  @page { size: A4; margin: 16mm; }
  body { background: #fff; }
  .toolbar, .section-toggle, dialog { display: none !important; }
  .document { width: 100%; margin: 0; border: 0; border-radius: 0; box-shadow: none; }
  .cover { min-height: 230mm; break-after: page; }
  .report-grid { display: block; padding: 0; }
  .toc { position: static; max-height: none; margin-bottom: 2rem; break-after: page; }
  .report-section { break-inside: auto; }
  .report-section.collapsed > :not(h1):not(h2) { display: block; }
  .table-wrap { overflow: visible; break-inside: avoid; }
  .wide-table { min-width: 0; font-size: 7pt; }
  .data-table th, .data-table td { min-width: 0; padding: 4pt; }
  .report-figure { break-inside: avoid; }
  a { color: inherit; }
}
</style>
"""
        '</head>\n<body>\n'
        '<header class="toolbar" aria-label="보고서 도구">\n'
        '<button id="theme-toggle" type="button" aria-label="테마 전환">테마</button>\n'
        '<button id="expand-all" type="button">모두 펼치기</button>\n'
        '<button id="collapse-all" type="button">모두 접기</button>\n'
        '<input id="report-search" type="search" placeholder="보고서에서 검색" aria-label="보고서 검색">\n'
        '<span id="search-count" aria-live="polite"></span>\n'
        '</header>\n'
        '<main class="document">\n'
        '<section class="cover">\n'
        '<p class="eyebrow">Gemini Deep Research</p>\n'
        f'<h1>{safe_title}</h1>\n'
        f'<div class="cover-meta">Agent {safe_agent}<br>{generated_at}</div>\n'
        '<div class="stat-row">\n'
        f'<span class="stat">본문 {text_char_count:,}자</span>\n'
        f'<span class="stat">표 {table_count:,}개</span>\n'
        f'<span class="stat">시각 자료 {image_count:,}개</span>\n'
        '</div>\n</section>\n'
        '<div class="report-grid">\n'
        f'<nav class="toc" aria-label="목차"><h2>Contents</h2><ol>{toc_html}</ol></nav>\n'
        f'<article>{body_html}<div id="no-results" class="no-results">검색 결과가 없습니다.</div></article>\n'
        '</div>\n</main>\n'
        '<dialog id="lightbox"><div class="lightbox-inner"><button class="lightbox-close" '
        'type="button" aria-label="닫기">×</button><img alt="확대 이미지"></div></dialog>\n'
        """<script>
(() => {
  const root = document.documentElement;
  const sections = [...document.querySelectorAll('.report-section')];
  const search = document.getElementById('report-search');
  const count = document.getElementById('search-count');
  const noResults = document.getElementById('no-results');

  try {
    const saved = localStorage.getItem('deep-research-theme');
    if (saved) root.dataset.theme = saved;
  } catch (_) {}
  document.getElementById('theme-toggle').addEventListener('click', () => {
    root.dataset.theme = root.dataset.theme === 'dark' ? 'light' : 'dark';
    try { localStorage.setItem('deep-research-theme', root.dataset.theme); } catch (_) {}
  });

  sections.forEach((section) => {
    const heading = section.querySelector(':scope > h1, :scope > h2');
    if (!heading) return;
    const button = document.createElement('button');
    button.type = 'button';
    button.className = 'section-toggle';
    button.textContent = '접기';
    button.setAttribute('aria-expanded', 'true');
    button.addEventListener('click', () => {
      const collapsed = section.classList.toggle('collapsed');
      button.textContent = collapsed ? '펼치기' : '접기';
      button.setAttribute('aria-expanded', String(!collapsed));
    });
    heading.prepend(button);
  });

  const setCollapsed = (collapsed) => sections.forEach((section) => {
    const button = section.querySelector(':scope > h1 .section-toggle, :scope > h2 .section-toggle');
    if (!button) return;
    section.classList.toggle('collapsed', collapsed);
    button.textContent = collapsed ? '펼치기' : '접기';
    button.setAttribute('aria-expanded', String(!collapsed));
  });
  document.getElementById('expand-all').addEventListener('click', () => setCollapsed(false));
  document.getElementById('collapse-all').addEventListener('click', () => setCollapsed(true));

  search.addEventListener('input', () => {
    const query = search.value.trim().toLocaleLowerCase();
    let visible = 0;
    sections.forEach((section) => {
      const matches = !query || section.textContent.toLocaleLowerCase().includes(query);
      section.hidden = !matches;
      if (matches) {
        visible += 1;
        if (query) section.classList.remove('collapsed');
      }
    });
    count.textContent = query ? `${visible}개 섹션` : '';
    noResults.classList.toggle('visible', Boolean(query) && visible === 0);
  });

  document.querySelectorAll('table[data-sortable="true"]').forEach((table) => {
    table.querySelectorAll('thead th').forEach((header, column) => {
      header.tabIndex = 0;
      header.title = '클릭해서 정렬';
      const sort = () => {
        const body = table.tBodies[0];
        if (!body) return;
        const ascending = header.getAttribute('aria-sort') !== 'ascending';
        table.querySelectorAll('th[aria-sort]').forEach((item) => item.removeAttribute('aria-sort'));
        header.setAttribute('aria-sort', ascending ? 'ascending' : 'descending');
        const rows = [...body.rows];
        rows.sort((a, b) => {
          const left = a.cells[column]?.textContent.trim() || '';
          const right = b.cells[column]?.textContent.trim() || '';
          return left.localeCompare(right, undefined, { numeric: true, sensitivity: 'base' }) * (ascending ? 1 : -1);
        });
        rows.forEach((row) => body.appendChild(row));
      };
      header.addEventListener('click', sort);
      header.addEventListener('keydown', (event) => {
        if (event.key === 'Enter' || event.key === ' ') { event.preventDefault(); sort(); }
      });
    });
  });

  const dialog = document.getElementById('lightbox');
  const dialogImage = dialog.querySelector('img');
  const openImage = (image) => {
    dialogImage.src = image.src;
    dialogImage.alt = image.alt;
    if (dialog.showModal) dialog.showModal(); else dialog.setAttribute('open', '');
  };
  document.querySelectorAll('.report-figure img').forEach((image) => {
    image.addEventListener('click', () => openImage(image));
    image.addEventListener('keydown', (event) => {
      if (event.key === 'Enter' || event.key === ' ') { event.preventDefault(); openImage(image); }
    });
  });
  dialog.querySelector('.lightbox-close').addEventListener('click', () => dialog.close());
  dialog.addEventListener('click', (event) => { if (event.target === dialog) dialog.close(); });

  if ('IntersectionObserver' in window) {
    const links = new Map([...document.querySelectorAll('.toc a')].map((link) => [decodeURIComponent(link.hash.slice(1)), link]));
    const observer = new IntersectionObserver((entries) => {
      entries.filter((entry) => entry.isIntersecting).forEach((entry) => {
        links.forEach((link) => link.classList.remove('active'));
        links.get(entry.target.id)?.classList.add('active');
      });
    }, { rootMargin: '-15% 0px -75% 0px' });
    document.querySelectorAll('.report-content h1, .report-content h2, .report-content h3').forEach((heading) => observer.observe(heading));
  }
})();
</script>
"""
        '</body>\n</html>\n'
    ).encode("utf-8")


def build_deep_research_report(interaction, *, agent_id, title):
    outputs = collect_interaction_outputs(interaction)
    if not outputs:
        raise RuntimeError("Deep Research outputs가 비어 있습니다.")
    text_preview = "\n\n".join(output["text"] for output in outputs if output["type"] == "text").strip()
    image_count = sum(1 for output in outputs if output["type"] == "image")
    html_bytes = build_deep_research_html(outputs, agent_id=agent_id, title=title)
    return {
        "kind": "deep_research_report",
        "html_bytes": html_bytes,
        "filename": f"deep_research_{datetime.now(KST).strftime('%Y%m%d_%H%M%S')}.html",
        "summary": (
            f"[Deep Research] 완료했습니다. 오프라인 HTML 보고서를 첨부합니다."
            f"\n텍스트 출력 {len(text_preview):,}자, 이미지 {image_count}개"
        ),
    }


async def call_deep_research(research_topic, agent_id):
    """대화 컨텍스트 없이 입력된 주제만 단독으로 리서치한다."""
    if not CLIENT:
        return "Error: API Client not initialized"
    research_topic = (research_topic or "").strip()
    if not research_topic:
        return "[ERROR] Deep Research에 전달할 텍스트가 없습니다."
    prompt = (
        "Research the topic below and produce a thorough, well-structured report. "
        "Write the report in the same language as the topic (e.g. Korean topic → Korean report).\n\n"
        "## Research Topic\n"
        f"{research_topic}"
    )

    interactions_api = get_interactions_api()
    if interactions_api is None:
        return "[ERROR] 현재 google-genai SDK에서 Interactions API를 찾을 수 없습니다. 패키지 업데이트가 필요합니다."

    try:
        interaction = await maybe_await(
            interactions_api.create(
                input=prompt,
                agent=agent_id,
                background=True,
                agent_config={"type": "deep-research", "thinking_summaries": "auto", "visualization": "auto"},
            )
        )
        interaction_id = getattr(interaction, "id", None) or (interaction.get("id") if isinstance(interaction, dict) else None)
        if not interaction_id:
            return "[ERROR] Deep Research 작업 ID를 받지 못했습니다."

        print(f"[DEEP] started agent={agent_id} interaction_id={interaction_id}")
        started_at = time_module.monotonic()
        last_status = None
        last_status_log_at = started_at
        while True:
            await asyncio.sleep(DEEP_RESEARCH_POLL_INTERVAL_SECONDS)
            result = await call_interaction_get(interactions_api, interaction_id)
            status = getattr(result, "status", None) or (result.get("status") if isinstance(result, dict) else None)
            now = time_module.monotonic()
            if status != last_status or now - last_status_log_at >= DEEP_RESEARCH_STATUS_LOG_INTERVAL_SECONDS:
                elapsed = int(now - started_at)
                print(
                    f"[DEEP] interaction_id={interaction_id} status={status or 'unknown'} "
                    f"elapsed={elapsed}s"
                )
                last_status = status
                last_status_log_at = now
            if status == "completed":
                try:
                    return build_deep_research_report(result, agent_id=agent_id, title=research_topic)
                except Exception as e:
                    text = extract_interaction_text(result)
                    if text:
                        return f"[ERROR] Deep Research HTML 생성 실패: {e}\n\n{text}"
                    return f"[ERROR] Deep Research HTML 생성 실패: {e}"
            if status in ("failed", "cancelled", "incomplete", "budget_exceeded", "requires_action"):
                error = getattr(result, "error", None) or (result.get("error") if isinstance(result, dict) else None)
                detail = f": {error}" if error else ""
                return f"[ERROR] Deep Research 중단 (status={status}){detail}"
    except Exception as e:
        return f"[ERROR] Deep Research 호출 실패: {e}"


async def route_message_action(
    last_user_msg,
    *,
    has_current_images=False,
    has_recent_images=False,
    text_model=None,
    recent_context=None,
):
    if not CLIENT:
        return {
            "action": "respond_text",
            "refined_prompt": "",
            "use_recent_images": False,
            "reason": "client unavailable",
        }
    if text_model is None:
        text_model = TEXT_MODEL
    text_model = get_structured_helper_model(text_model)
    recent_context = (recent_context or "").strip()
    prompt = f"""
You are a routing planner for a Discord AI bot.
Decide the next action for the assistant.

Available actions:
- respond_text: normal text reply.
- generate_image: the user wants a new image or wants to edit/transform an image. In this case, write a refined English prompt for the image model.

Rules:
- Do not infer that product names, drinks, people, or objects refer to stored images by themselves.
- Do not choose generate_image unless the user clearly wants image creation or visual transformation/editing.
- If the current message is only a complaint, acknowledgement, or confusion about a previous answer/image, choose respond_text.
- If the current message is a follow-up image request or edit request, use the recent conversation excerpt to resolve the subject/style and write a specific refined prompt.
- `use_recent_images` should be true only for editing or transforming a recent image.

Current message has image attachments: {has_current_images}
Recent image references are available: {has_recent_images}

Recent conversation excerpt:
{recent_context or '(empty)'}

User message:
{last_user_msg}
""".strip()

    try:
        interaction = await create_model_interaction(
            model=text_model,
            input_data=prompt,
            response_schema=IMAGE_ACTION_SCHEMA,
        )
        payload = coerce_json_object(extract_interaction_text(interaction))
        if not payload:
            raise ValueError("empty or invalid JSON route response")
        action = payload.get("action", "respond_text")
        if action not in {"respond_text", "generate_image"}:
            action = "respond_text"
        if (
            action == "generate_image"
            and not payload.get("use_recent_images")
            and not has_current_images
            and has_recent_images
            and FOLLOWUP_IMAGE_EDIT_RE.search(last_user_msg or "")
        ):
            payload["use_recent_images"] = True
        return {
            "action": action,
            "refined_prompt": (payload.get("refined_prompt") or "").strip(),
            "use_recent_images": bool(payload.get("use_recent_images", False)),
            "reason": (payload.get("reason") or "").strip(),
        }
    except Exception as e:
        print(f"[IMG] 라우팅 판정 실패: {e}")
        is_followup_edit = bool(FOLLOWUP_IMAGE_EDIT_RE.search(last_user_msg or ""))
        if has_image_generation_signal(last_user_msg, has_current_images=has_current_images) or is_followup_edit:
            return {
                "action": "generate_image",
                "refined_prompt": "",
                "use_recent_images": bool(is_followup_edit and has_recent_images and not has_current_images),
                "reason": "routing failure direct image signal",
            }
        return {
            "action": "respond_text",
            "refined_prompt": "",
            "use_recent_images": False,
            "reason": "routing failure",
        }

# ================= 3. 유틸리티 & API =================
async def download_media(url):
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(url) as resp:
                if resp.status == 200:
                    data = await resp.read()
                    return io.BytesIO(data), resp.headers.get('Content-Type')
                else:
                    print(f"[WARN] download_media failed: HTTP {resp.status} for {url[:80]}")
    except Exception as e:
        print(f"[WARN] download_media error: {e}")
    return None, None

async def upload_file_api(file_data, mime_type):
    """File API로 대용량 미디어 업로드 후 ACTIVE 상태까지 polling"""
    if not CLIENT: return None
    tmp = None
    try:
        suffix = '.' + mime_type.split('/')[-1] if '/' in mime_type else '.bin'
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
        tmp.write(file_data)
        tmp.close()
        uploaded = await CLIENT.aio.files.upload(
            file=Path(tmp.name),
            config=types.UploadFileConfig(mime_type=mime_type)
        )
        for _ in range(30):  # 최대 60초
            if uploaded.state and uploaded.state.name == "ACTIVE":
                return uploaded
            await asyncio.sleep(2)
            uploaded = await CLIENT.aio.files.get(name=uploaded.name)
        print(f"[WARN] File API 업로드 타임아웃: {mime_type}")
        return None
    except Exception as e:
        print(f"[ERROR] File API 업로드 실패: {e}")
        return None
    finally:
        if tmp:
            try: os.unlink(tmp.name)
            except: pass

async def scrap_webpage(url):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"
    }
    try:
        async with aiohttp.ClientSession(headers=headers) as session:
            async with session.get(url, timeout=5) as resp:
                if resp.status == 200:
                    html = await resp.text()
                    soup = BeautifulSoup(html, 'html.parser')
                    for script in soup(["script", "style", "nav", "footer", "iframe"]):
                        script.extract()
                    text = soup.get_text(separator='\n')
                    clean_text = '\n'.join([line.strip() for line in text.splitlines() if line.strip()])
                    return clean_text
    except Exception as e:
        print(f"[WARN] URL 읽기 실패: {e}")
    return None


def split_discord_message(content, limit=1900):
    chunks = []
    remaining = content or ""
    while len(remaining) > limit:
        split_at = remaining.rfind("\n\n", 0, limit + 1)
        if split_at < limit // 2:
            split_at = remaining.rfind("\n", 0, limit + 1)
        if split_at < limit // 2:
            split_at = remaining.rfind(" ", 0, limit + 1)
        if split_at <= 0:
            split_at = limit
        chunks.append(remaining[:split_at].rstrip())
        remaining = remaining[split_at:].lstrip("\n")
    if remaining:
        chunks.append(remaining)
    return chunks


async def send_long_message(channel, content):
    for chunk in split_discord_message(content):
        await channel.send(chunk)


async def call_nano_banana(contents, image_model=None):
    if not CLIENT: return None
    if image_model is None:
        image_model = IMAGE_MODEL

    for attempt in range(1, GEMINI_IMAGE_MAX_ATTEMPTS + 1):
        try:
            input_blocks = image_generation_inputs_to_blocks(contents)
            if not input_blocks:
                raise ValueError("이미지 생성 입력이 비어 있습니다.")
            interaction = await create_model_interaction(
                model=image_model,
                input_data=input_blocks,
                response_format={"type": "image"},
                timeout_seconds=GEMINI_IMAGE_TIMEOUT_SECONDS,
            )
            status = get_interaction_status(interaction)
            if status not in {"completed", "incomplete"}:
                raise Exception(f"Image interaction status: {status}")

            for output in collect_interaction_outputs(interaction):
                if output["type"] == "image":
                    return output["data"]

            output_image = get_interaction_output_value(interaction, "output_image")
            image_data = decode_interaction_image(get_interaction_output_value(output_image, "data"))
            if image_data:
                return image_data

            raise Exception("No image data found in response")

        except Exception as e:
            error_msg = describe_gemini_exception(e, GEMINI_IMAGE_TIMEOUT_SECONDS)
            retryable = isinstance(e, (TimeoutError, asyncio.TimeoutError)) or any(
                marker in error_msg.lower()
                for marker in ("connection", "429", "500", "502", "503", "504", "unavailable")
            )
            if not retryable or attempt == GEMINI_IMAGE_MAX_ATTEMPTS:
                raise
            print(
                f"[WARN] 이미지 생성 실패 (시도 {attempt}/{GEMINI_IMAGE_MAX_ATTEMPTS}): "
                f"{error_msg} -> 5초 후 재시도..."
            )
            await asyncio.sleep(5)

async def call_gemini(history_contents, text_model=None, thinking_level=None):
    if not CLIENT:
        raise RuntimeError("API Client not initialized")
    if text_model is None:
        text_model = TEXT_MODEL
    if thinking_level is None:
        thinking_level = MAX_THINKING_LEVEL

    input_data = gemini_contents_to_interaction_steps(history_contents)
    system_instruction = build_gemini_system_prompt()
    call_trace = {"sdk_version": GENAI_SDK_VERSION, "calls": []}
    interaction = await create_traced_interaction(
        call_trace,
        phase="initial",
        model=text_model,
        input_data=input_data,
        system_instruction=system_instruction,
        tools=[GOOGLE_SEARCH_TOOL],
        thinking_level=thinking_level,
    )
    token_usage = get_interaction_usage(interaction)

    if not extract_interaction_text(interaction):
        continuation_steps = collect_search_continuation_steps(interaction)
        if continuation_steps:
            print(
                "[WARN] Google Search 결과 뒤 최종 출력이 없어 같은 검색 결과로 생성을 이어갑니다. "
                f"response={summarize_interaction_response(interaction)} "
                f"call={json.dumps(call_trace, ensure_ascii=False, default=str, separators=(',', ':'))}"
            )
            continuation_input = [*input_data, *continuation_steps]
            interaction = await create_traced_interaction(
                call_trace,
                phase="search_continuation",
                model=text_model,
                input_data=continuation_input,
                system_instruction=system_instruction,
                thinking_level=thinking_level,
            )
            token_usage += get_interaction_usage(interaction)

    return interaction, token_usage, call_trace


# ================= 4. 봇 메인 로직 =================
class GeminiBot(commands.Bot):
    def __init__(self):
        super().__init__(command_prefix="!", intents=discord.Intents.all())
        self.server_configs = DataManager.load_json(CONFIG_FILE, {})
        model_config_changed = False
        supported_text_models = {info["text"] for info in AVAILABLE_MODELS.values()}
        for guild_id, guild_config in self.server_configs.items():
            if guild_id.startswith("_") or not isinstance(guild_config, dict):
                continue
            configured_model = guild_config.get("text_model")
            if configured_model and configured_model not in supported_text_models:
                guild_config["text_model"] = TEXT_MODEL
                guild_config["image_model"] = IMAGE_MODEL
                guild_config["thinking_level"] = MAX_THINKING_LEVEL
                model_config_changed = True
                print(
                    f"[CONFIG] guild={guild_id} unsupported text model reset: "
                    f"{configured_model} -> {TEXT_MODEL}"
                )
            elif configured_model and guild_config.get("thinking_level") != MAX_THINKING_LEVEL:
                guild_config["thinking_level"] = MAX_THINKING_LEVEL
                model_config_changed = True
                print(f"[CONFIG] guild={guild_id} thinking level fixed: {MAX_THINKING_LEVEL}")
        if model_config_changed:
            DataManager.save_json(CONFIG_FILE, self.server_configs)
        self.channel_locks = {}
        self.voice_locks = {}
        self.voice_retry_after = {}
        self.voice_connected_at = {}
        self.voice_disconnected_at = {}
        self.voice_bad_state_counts = {}
        self.loop_monitor_expected = None
        self.doc_cache = DocumentCache(DOC_CACHE_DIR)
        self.sound_files = scan_sound_files()
        global_cfg = self.get_global_config()
        self.sound_volume = clamp_sound_volume(global_cfg.get("sound_volume", SOUND_VOLUME))
        if global_cfg.get("sound_volume") != self.sound_volume:
            global_cfg["sound_volume"] = self.sound_volume
            self.save_config()
        print(f"[SOUND] 볼륨 로드: {int(self.sound_volume * 100)}% ({CONFIG_FILE})")
        self._sound_dir_mtime = self._get_sound_dir_mtime()
        global SIBO_LATENCY_COMP
        SIBO_LATENCY_COMP = self.server_configs.get("_global", {}).get("sibo_comp", SIBO_LATENCY_COMP)
        self.active_views = []  # View 참조 유지용
        self.mafic_pool = mafic.NodePool(self) if LAVALINK_ENABLED else None
        self.mafic_node = None
        self.mafic_connect_lock = asyncio.Lock()

    def _get_sound_dir_mtime(self):
        """SOUND_DIR 전체의 최신 mtime 반환."""
        latest = 0
        for root, dirs, files in os.walk(SOUND_DIR):
            dirs[:] = [d for d in dirs if not d.startswith(".") and d != "__pycache__"]
            for name in files + dirs:
                if name.startswith("."):
                    continue
                try:
                    mt = os.path.getmtime(os.path.join(root, name))
                    if mt > latest:
                        latest = mt
                except OSError:
                    pass
            # root 폴더 자체의 mtime도 체크 (파일 추가/삭제 감지)
            try:
                mt = os.path.getmtime(root)
                if mt > latest:
                    latest = mt
            except OSError:
                pass
        return latest

    def get_config(self, guild_id):
        gid = str(guild_id)
        if gid not in self.server_configs:
            self.server_configs[gid] = {
                "channel_id": None,
                "reset_time": None,
                "text_model": TEXT_MODEL,
                "image_model": IMAGE_MODEL,
                "thinking_level": MAX_THINKING_LEVEL,
                "history_limit": DEFAULT_HISTORY_LIMIT,
                "image_history_limit": DEFAULT_IMAGE_HISTORY_LIMIT,
                "action_router": DEFAULT_ACTION_ROUTER_ENABLED,
            }
        self.server_configs[gid].setdefault("text_model", TEXT_MODEL)
        self.server_configs[gid].setdefault("image_model", IMAGE_MODEL)
        self.server_configs[gid]["thinking_level"] = MAX_THINKING_LEVEL
        self.server_configs[gid].setdefault("history_limit", DEFAULT_HISTORY_LIMIT)
        self.server_configs[gid].setdefault("image_history_limit", DEFAULT_IMAGE_HISTORY_LIMIT)
        self.server_configs[gid].setdefault("action_router", DEFAULT_ACTION_ROUTER_ENABLED)
        return self.server_configs[gid]

    def get_sibo(self, guild_id):
        return self.get_config(guild_id).get("sibo")

    def set_sibo(self, guild_id, sibo_cfg):
        self.get_config(guild_id)["sibo"] = sibo_cfg
        self.save_config()

    def del_sibo(self, guild_id):
        cfg = self.get_config(guild_id)
        cfg.pop("sibo", None)
        self.save_config()

    def get_all_sibo(self):
        """모든 서버의 sibo 설정을 {gid: sibo_cfg} 딕셔너리로 반환."""
        result = {}
        for gid, cfg in self.server_configs.items():
            if gid.startswith("_") or not isinstance(cfg, dict):
                continue
            sibo = cfg.get("sibo")
            if sibo and isinstance(sibo, dict):
                result[gid] = sibo
        return result

    def get_global_config(self):
        global_cfg = self.server_configs.get("_global")
        if not isinstance(global_cfg, dict):
            global_cfg = {}
            self.server_configs["_global"] = global_cfg
        return global_cfg

    def get_lavalink_session_id(self):
        return self.get_global_config().get("lavalink_session_id")

    def set_lavalink_session_id(self, session_id):
        cfg = self.get_global_config()
        if cfg.get("lavalink_session_id") == session_id:
            return
        cfg["lavalink_session_id"] = session_id
        self.save_config()

    def set_sound_volume(self, volume):
        self.sound_volume = clamp_sound_volume(volume)
        self.get_global_config()["sound_volume"] = self.sound_volume
        self.save_config()
        print(f"[SOUND] 볼륨 저장: {int(self.sound_volume * 100)}% ({CONFIG_FILE})")

    def save_config(self):
        DataManager.save_json(CONFIG_FILE, self.server_configs)

    async def ensure_lavalink_node(self):
        if not LAVALINK_ENABLED:
            return None
        async with self.mafic_connect_lock:
            if self.mafic_node is not None:
                return self.mafic_node
            self.mafic_node = await self.mafic_pool.create_node(
                host=LAVALINK_HOST,
                port=LAVALINK_PORT,
                label=LAVALINK_LABEL,
                password=LAVALINK_PASSWORD,
                secure=LAVALINK_SECURE,
                timeout=LAVALINK_TIMEOUT,
                resuming_session_id=self.get_lavalink_session_id(),
                player_cls=mafic.Player,
            )
            if getattr(self.mafic_node, "session_id", None):
                self.set_lavalink_session_id(self.mafic_node.session_id)
            return self.mafic_node

    def get_channel_lock(self, channel_id):
        if channel_id not in self.channel_locks:
            self.channel_locks[channel_id] = asyncio.Lock()
        return self.channel_locks[channel_id]

    def get_voice_lock(self, guild_id):
        gid = str(guild_id)
        if gid not in self.voice_locks:
            self.voice_locks[gid] = asyncio.Lock()
        return self.voice_locks[gid]

    def get_voice_target(self, guild_id):
        cfg = self.get_config(guild_id)
        channel_id = cfg.get("voice_target_channel_id")
        if channel_id:
            return channel_id
        sibo = cfg.get("sibo")
        if isinstance(sibo, dict):
            return sibo.get("vc_channel")
        return None

    def set_voice_target(self, guild_id, channel_id):
        cfg = self.get_config(guild_id)
        if cfg.get("voice_target_channel_id") != channel_id:
            cfg["voice_target_channel_id"] = channel_id
            self.save_config()
            guild = self.get_guild(guild_id)
            guild_name = guild.name if guild else guild_id
            log_voice_debug(f"[VC] keepalive 대상 저장 ({guild_name}): channel_id={channel_id}")

    def clear_voice_target(self, guild_id):
        cfg = self.get_config(guild_id)
        if "voice_target_channel_id" in cfg:
            old_channel_id = cfg.get("voice_target_channel_id")
            cfg.pop("voice_target_channel_id", None)
            self.save_config()
            self.voice_bad_state_counts.pop(str(guild_id), None)
            self.voice_connected_at.pop(str(guild_id), None)
            self.voice_disconnected_at.pop(str(guild_id), None)
            guild = self.get_guild(guild_id)
            guild_name = guild.name if guild else guild_id
            log_voice_debug(f"[VC] keepalive 대상 해제 ({guild_name}): channel_id={old_channel_id}")

    def can_retry_voice(self, guild_id):
        now = time_module.monotonic()
        gid = str(guild_id)
        next_allowed = self.voice_retry_after.get(gid, 0.0)
        if now < next_allowed:
            return False
        self.voice_retry_after[gid] = now + VOICE_RECONNECT_COOLDOWN
        return True

    def reset_voice_retry(self, guild_id):
        self.voice_retry_after[str(guild_id)] = 0.0

    def mark_voice_connected(self, guild_id):
        gid = str(guild_id)
        self.voice_connected_at[gid] = time_module.monotonic()
        self.voice_disconnected_at.pop(gid, None)
        self.voice_bad_state_counts[gid] = 0

    def mark_voice_disconnected(self, guild_id):
        gid = str(guild_id)
        self.voice_connected_at.pop(gid, None)
        self.voice_disconnected_at[gid] = time_module.monotonic()
        self.voice_bad_state_counts[gid] = 0

    def get_voice_disconnected_age(self, guild_id):
        started = self.voice_disconnected_at.get(str(guild_id))
        if started is None:
            return None
        return time_module.monotonic() - started

    def get_voice_connected_age(self, guild_id):
        started = self.voice_connected_at.get(str(guild_id))
        if started is None:
            return None
        return time_module.monotonic() - started

    def record_voice_health(self, guild_id, unhealthy_reason):
        gid = str(guild_id)
        if unhealthy_reason:
            self.voice_bad_state_counts[gid] = self.voice_bad_state_counts.get(gid, 0) + 1
            return self.voice_bad_state_counts[gid]
        self.voice_bad_state_counts[gid] = 0
        return 0

bot = GeminiBot()

@tasks.loop(seconds=10)
async def voice_watchdog():
    """VC 연결 유지 감시: 끊긴 경우 자동 재접속 + 비정상 상태 강제 재연결."""
    for guild in bot.guilds:
        vc = guild.voice_client
        if vc and vc.is_connected():
            try:
                target_channel_id = bot.get_voice_target(guild.id)
                if VOICE_HEALTH_RECOVERY:
                    unhealthy_reason = get_voice_unhealthy_reason(vc)
                    connected_age = bot.get_voice_connected_age(guild.id)
                    bad_count = bot.record_voice_health(guild.id, unhealthy_reason)
                    if (
                        unhealthy_reason
                        and target_channel_id
                        and connected_age is not None
                        and connected_age >= VOICE_HEALTH_GRACE
                        and bad_count >= VOICE_HEALTH_BAD_LIMIT
                        and bot.can_retry_voice(guild.id)
                    ):
                        log_voice_snapshot(
                            "watchdog 강제 재연결",
                            guild,
                            vc,
                            extra=(
                                f"reason={unhealthy_reason} "
                                f"connected_age={connected_age:.1f}s "
                                f"bad_count={bad_count} "
                                f"target_channel_id={target_channel_id}"
                            ),
                        )
                        await cleanup_stale_voice_client(vc, guild)
                        vc = await ensure_voice_connection(
                            guild,
                            target_channel_id,
                            move=True,
                            reason="watchdog_unhealthy",
                        )
                        if vc:
                            bot.reset_voice_retry(guild.id)
                        continue
                else:
                    bot.record_voice_health(guild.id, None)
                if target_channel_id:
                    await ensure_silence_keepalive(guild, vc, reason="watchdog_idle")
            except Exception as e:
                log_voice_debug(f"[VC] watchdog 재시작 실패 ({guild.name}): {e}")
            continue

        target_channel_id = bot.get_voice_target(guild.id)
        disconnected_age = bot.get_voice_disconnected_age(guild.id)
        if (
            not target_channel_id
            or disconnected_age is None
            or disconnected_age < VOICE_LIBRARY_RECOVERY_GRACE
            or not bot.can_retry_voice(guild.id)
        ):
            continue

        try:
            log_voice_snapshot(
                "watchdog 재접속 시도",
                guild,
                vc,
                extra=(
                    f"target_channel_id={target_channel_id} "
                    f"disconnected_age={disconnected_age:.1f}s"
                ),
            )
            vc = await ensure_voice_connection(
                guild,
                target_channel_id,
                move=True,
                reason="watchdog",
            )
            if vc:
                bot.reset_voice_retry(guild.id)
        except Exception as e:
            log_voice_debug(f"[VC] watchdog 재접속 실패 ({guild.name}): {e}")

@tasks.loop(seconds=5)
async def sound_watcher():
    """SOUND_DIR 변화 감지 → 자동 리스캔."""
    current_mtime = bot._get_sound_dir_mtime()
    if current_mtime != bot._sound_dir_mtime:
        bot._sound_dir_mtime = current_mtime
        old_count = len(bot.sound_files)
        bot.sound_files = scan_sound_files()
        new_count = len(bot.sound_files)
        if old_count != new_count:
            print(f"[SB] 사운드 자동 리스캔: {old_count} → {new_count}개")

@tasks.loop(seconds=LOOP_MONITOR_INTERVAL)
async def loop_watchdog():
    now = time_module.monotonic()
    expected = bot.loop_monitor_expected
    bot.loop_monitor_expected = now + LOOP_MONITOR_INTERVAL
    if expected is None:
        return

    lag = now - expected
    if lag < LOOP_MONITOR_WARN:
        return

    active_targets = sum(1 for guild in bot.guilds if bot.get_voice_target(guild.id))
    print(f"[LOOP] event loop lag 감지: {lag:.3f}s | active_voice_targets={active_targets}")
    for guild in bot.guilds:
        target_channel_id = bot.get_voice_target(guild.id)
        if guild.voice_client is not None or target_channel_id:
            log_voice_snapshot(
                "event loop lag 시점",
                guild,
                guild.voice_client,
                extra=f"lag={lag:.3f}s target_channel_id={target_channel_id}",
            )

_sibo_generation = 0  # 시보 루프 세대 카운터

def restart_sibo():
    """시보 스케줄러를 취소하고 재시작. 세대 카운터로 중복 방지."""
    global _sibo_generation
    _sibo_generation += 1
    current_gen = _sibo_generation
    # 기존 task 취소
    if hasattr(bot, '_sibo_task') and bot._sibo_task and not bot._sibo_task.done():
        bot._sibo_task.cancel()
    bot._sibo_task = bot.loop.create_task(sibo_loop(current_gen))
    print(f"[SIBO] 스케줄러 재시작 (gen {current_gen})")

async def sibo_loop(my_gen=None):
    """정밀 시보 스케줄러. 세대 불일치 시 자동 종료."""
    await bot.wait_until_ready()
    while not bot.is_closed():
        # 내가 최신 세대가 아니면 즉시 종료
        if my_gen is not None and my_gen != _sibo_generation:
            print(f"[SIBO] gen {my_gen} 종료 (최신: {_sibo_generation})")
            return
        all_sibo = bot.get_all_sibo()
        if not all_sibo:
            await asyncio.sleep(1)
            continue

        # 유효한 설정만 필터
        valid = {}
        for gid, cfg in all_sibo.items():
            sound_key = cfg.get("sound")
            if sound_key and sound_key in bot.sound_files:
                valid[gid] = cfg
        if not valid:
            await asyncio.sleep(1)
            continue

        # 가장 큰 offset 기준으로 sleep (가장 먼저 재생해야 하는 것)
        max_offset = max(c.get("offset", 0.0) for c in valid.values())
        total_advance = max_offset + SIBO_LATENCY_COMP

        now = datetime.now(KST)
        next_hour = now.replace(minute=0, second=0, microsecond=0) + timedelta(hours=1)
        play_time = next_hour - timedelta(seconds=total_advance)
        wait_seconds = (play_time - now).total_seconds()

        if wait_seconds < 0:
            next_hour += timedelta(hours=1)
            play_time = next_hour - timedelta(seconds=total_advance)
            wait_seconds = (play_time - now).total_seconds()

        target_hour = next_hour.strftime('%H:00:00')
        print(f"[SIBO] 목표 정각: {target_hour} | 최대 offset: {max_offset}s | {wait_seconds:.1f}초 후")

        if LAVALINK_ENABLED:
            try:
                await asyncio.sleep(wait_seconds)
            except asyncio.CancelledError:
                print(f"[SIBO] gen {my_gen} 취소됨")
                return

            if my_gen is not None and my_gen != _sibo_generation:
                print(f"[SIBO] gen {my_gen} 종료 (최신: {_sibo_generation})")
                return

            valid = {}
            for gid, cfg in bot.get_all_sibo().items():
                sound_key = cfg.get("sound")
                if sound_key and sound_key in bot.sound_files:
                    valid[gid] = cfg

            last_wait = 0.0
            for gid, cfg in sorted(valid.items(), key=lambda item: item[1].get("offset", 0.0), reverse=True):
                sound_key = cfg.get("sound")
                offset = cfg.get("offset", 0.0)
                filepath = bot.sound_files.get(sound_key)
                if not filepath:
                    continue

                wait_delta = max(0.0, last_wait - offset)
                if wait_delta > 0:
                    await asyncio.sleep(wait_delta)
                last_wait = offset

                guild = bot.get_guild(int(gid))
                if not guild:
                    continue

                vc = guild.voice_client
                if not vc or not voice_is_connected(vc):
                    vc_channel_id = cfg.get("vc_channel")
                    if not vc_channel_id:
                        continue
                    try:
                        vc = await ensure_voice_connection(
                            guild,
                            vc_channel_id,
                            move=True,
                            reason="sibo_lavalink_reconnect",
                        )
                    except Exception as e:
                        print(f"[SIBO] Guild {gid} Lavalink VC 재접속 실패: {e}")
                        continue

                if not vc:
                    continue

                try:
                    if voice_is_playing(vc):
                        await voice_stop(vc)
                    await play_lavalink_local_sound(
                        vc,
                        filepath,
                        bot.sound_volume,
                        replace=True,
                        lead_in_ms=LAVALINK_SIBO_PAD_MS,
                    )
                    play_now = datetime.now(KST)
                    print(f"[SIBO] {play_now.strftime('%H:%M:%S.%f')[:-4]} Guild {gid} 시보 재생(Lavalink): {sound_key}")
                except Exception as e:
                    print(f"[SIBO] Guild {gid} Lavalink 재생 실패: {e}")

            await asyncio.sleep(1)
            continue

        # 1) play_time 보다 PRELOAD_LEAD 일찍 깨서 ffmpeg 미리 준비
        pre_warm_wait = max(0.0, wait_seconds - SIBO_PRELOAD_LEAD)
        try:
            await asyncio.sleep(pre_warm_wait)
        except asyncio.CancelledError:
            print(f"[SIBO] gen {my_gen} 취소됨")
            return

        # 세대 체크 (pre-warm 진입 전)
        if my_gen is not None and my_gen != _sibo_generation:
            print(f"[SIBO] gen {my_gen} 종료 (최신: {_sibo_generation})")
            return

        # 변경됐을 수 있으니 최신 설정 다시 읽기
        valid = {}
        for gid, cfg in bot.get_all_sibo().items():
            sound_key = cfg.get("sound")
            if sound_key and sound_key in bot.sound_files:
                valid[gid] = cfg

        # 2) 사전 생성 패스: 전 서버 FFmpegPCMAudio 동시 spawn (파이프 버퍼 채우기 시작)
        guild_audio = {}
        for gid, cfg in valid.items():
            sound_key = cfg.get("sound")
            filepath = bot.sound_files.get(sound_key)
            if not filepath:
                continue
            try:
                guild_audio[gid] = discord.PCMVolumeTransformer(
                    discord.FFmpegPCMAudio(filepath, options="-af loudnorm=I=-16:TP=-1.5:LRA=11"),
                    volume=bot.sound_volume
                )
            except Exception as e:
                print(f"[SIBO] Guild {gid} 사전 생성 실패: {e}")

        # 3) 남은 PRELOAD_LEAD 만큼 대기 (ffmpeg가 파이프 버퍼 채울 시간) → play_time 도달
        if pre_warm_wait > 0:
            try:
                await asyncio.sleep(SIBO_PRELOAD_LEAD)
            except asyncio.CancelledError:
                for a in guild_audio.values():
                    try: a.cleanup()
                    except Exception: pass
                print(f"[SIBO] gen {my_gen} 취소됨")
                return

        # 세대 재체크 (취소/전환됐으면 pre-created subprocess 정리)
        if my_gen is not None and my_gen != _sibo_generation:
            for a in guild_audio.values():
                try: a.cleanup()
                except Exception: pass
            print(f"[SIBO] gen {my_gen} 종료 (최신: {_sibo_generation})")
            return

        # 4) 타이밍 + 재생 패스: offset별 순차 재생 (이미 준비된 audio 사용)
        max_offset = max((c.get("offset", 0.0) for c in valid.values()), default=0)
        for gid, cfg in valid.items():
            sound_key = cfg.get("sound")
            offset = cfg.get("offset", 0.0)

            # 이 서버의 재생까지 남은 추가 대기
            extra_wait = max_offset - offset
            if extra_wait > 0:
                await asyncio.sleep(extra_wait)

            audio = guild_audio.get(gid)
            if not audio:
                continue

            guild = bot.get_guild(int(gid))
            if not guild:
                try: audio.cleanup()
                except Exception: pass
                continue
            vc = guild.voice_client
            if not vc or not voice_is_connected(vc):
                # VC 끊겼으면 저장된 채널로 자동 재접속
                vc_channel_id = cfg.get("vc_channel")
                if vc_channel_id:
                    channel = guild.get_channel(vc_channel_id)
                    if channel:
                        try:
                            vc = await ensure_voice_connection(
                                guild,
                                vc_channel_id,
                                move=True,
                                reason="sibo_reconnect",
                            )
                            if not vc:
                                try: audio.cleanup()
                                except Exception: pass
                                continue
                            print(f"[SIBO] Guild {gid} VC 자동 재접속: {channel.name}")
                        except Exception as e:
                            print(f"[SIBO] Guild {gid} VC 재접속 실패: {e}")
                            try: audio.cleanup()
                            except Exception: pass
                            continue
                    else:
                        try: audio.cleanup()
                        except Exception: pass
                        continue
                else:
                    try: audio.cleanup()
                    except Exception: pass
                    continue

            if voice_is_playing(vc):
                await voice_stop(vc)
            def after_sibo_play(error, guild_id=int(gid), sound_name=sound_key):
                if error:
                    print(f"[SIBO] Guild {guild_id} 재생 종료 오류: {error}")
                schedule_silence_keepalive_resume(
                    guild_id,
                    reason=f"sibo:{sound_name}",
                )

            vc.play(audio, after=after_sibo_play)
            play_now = datetime.now(KST)
            print(f"[SIBO] {play_now.strftime('%H:%M:%S.%f')[:-4]} Guild {gid} 시보 재생: {sound_key}")

        await asyncio.sleep(1)

PANEL_EMBED_TITLE_PREFIXES = (
    "🎛 Gemini 패널",
    "🔊 사운드보드",
    "⏰ 시보 설정",
)


def is_panel_message(message):
    if message.author != bot.user:
        return False
    return any(
        (embed.title or "").startswith(prefix)
        for embed in message.embeds
        for prefix in PANEL_EMBED_TITLE_PREFIXES
    )


async def ensure_control_panel(guild, channel, *, save_target=False):
    """채널 전체를 지우지 않고 통합 패널 메시지 하나만 준비한다."""
    if save_target:
        config = bot.get_config(guild.id)
        config["sb_channel_id"] = channel.id
        bot.save_config()

    panel_view = ControlPanelView(guild.id)
    bot.active_views.append(panel_view)

    panel_messages = []
    try:
        async for msg in channel.history(limit=50):
            if is_panel_message(msg):
                panel_messages.append(msg)
    except Exception as e:
        print(f"[PANEL] {guild.name}: 기존 패널 확인 실패: {e}")

    if panel_messages:
        target = panel_messages[0]
        await target.edit(embed=panel_view._make_embed(), view=panel_view)
        for stale in panel_messages[1:4]:
            try:
                await stale.delete()
            except Exception as e:
                print(f"[PANEL] {guild.name}: 중복 패널 삭제 실패: {e}")
        print(f"[PANEL] {guild.name}: #{channel.name} 통합 패널 갱신 완료")
        return "updated"

    await channel.send(embed=panel_view._make_embed(), view=panel_view)
    print(f"[PANEL] {guild.name}: #{channel.name} 통합 패널 생성 완료")
    return "created"


async def setup_panel_channels():
    """봇 시작 시 저장된 채널의 통합 패널을 갱신한다."""
    for guild in bot.guilds:
        config = bot.get_config(guild.id)
        panel_channel_id = config.get("sb_channel_id")
        if not panel_channel_id:
            continue
        channel = guild.get_channel(panel_channel_id)
        if not channel:
            print(f"[PANEL] {guild.name}: 패널 채널 없음 (ID: {panel_channel_id})")
            continue
        try:
            await ensure_control_panel(guild, channel)
        except Exception as e:
            print(f"[PANEL] {guild.name}: 임베드 생성 실패: {e}")

@bot.event
async def on_ready():
    print(f"[OK] {bot.user} 로그인 완료!")
    try:
        synced = await bot.tree.sync()
        print(f"[OK] {len(synced)}개의 슬래시 커맨드 동기화 완료")
    except Exception as e:
        print(f"[ERROR] 슬래시 커맨드 동기화 실패: {e}")
    if LAVALINK_ENABLED:
        try:
            node = await bot.ensure_lavalink_node()
            session_id = getattr(node, "session_id", None)
            print(f"[OK] Lavalink 연결 완료: {LAVALINK_HOST}:{LAVALINK_PORT} label={LAVALINK_LABEL} session_id={session_id}")
        except Exception as e:
            print(f"[ERROR] Lavalink 연결 실패: {e}")
    if not voice_watchdog.is_running():
        voice_watchdog.start()
    if not sound_watcher.is_running():
        sound_watcher.start()
    if not loop_watchdog.is_running():
        bot.loop_monitor_expected = None
        loop_watchdog.start()
    if not hasattr(bot, '_sibo_task') or bot._sibo_task is None or bot._sibo_task.done():
        global _sibo_generation
        _sibo_generation += 1
        bot._sibo_task = bot.loop.create_task(sibo_loop(_sibo_generation))
    print(f"[OK] 사운드 감시 + 시보 스케줄러 시작")
    for guild in bot.guilds:
        target_channel_id = bot.get_voice_target(guild.id)
        if guild.voice_client and guild.voice_client.is_connected():
            bot.mark_voice_connected(guild.id)
            await ensure_silence_keepalive(guild, guild.voice_client, reason="on_ready")
        if target_channel_id:
            log_voice_snapshot(
                "on_ready keepalive 상태",
                guild,
                guild.voice_client,
                extra=f"target_channel_id={target_channel_id}",
            )
    await setup_panel_channels()

@bot.event
async def on_disconnect():
    active = sum(1 for guild in bot.guilds if guild.voice_client is not None)
    log_voice_debug(f"[GATEWAY] on_disconnect | active_voice_clients={active}")
    for guild in bot.guilds:
        target_channel_id = bot.get_voice_target(guild.id)
        if guild.voice_client is not None or target_channel_id:
            log_voice_snapshot(
                "gateway disconnect 시점",
                guild,
                guild.voice_client,
                extra=f"target_channel_id={target_channel_id}",
            )

@bot.event
async def on_resumed():
    active = sum(1 for guild in bot.guilds if guild.voice_client is not None)
    log_voice_debug(f"[GATEWAY] on_resumed | active_voice_clients={active}")
    for guild in bot.guilds:
        target_channel_id = bot.get_voice_target(guild.id)
        if guild.voice_client is not None or target_channel_id:
            log_voice_snapshot(
                "gateway resumed 시점",
                guild,
                guild.voice_client,
                extra=f"target_channel_id={target_channel_id}",
            )

@bot.event
async def on_node_ready(node):
    if getattr(node, "session_id", None):
        bot.set_lavalink_session_id(node.session_id)
    print(
        f"[LAVALINK] node ready: label={getattr(node, 'label', '-')}"
        f" session_id={getattr(node, 'session_id', None)}"
    )

@bot.event
async def on_node_unavailable(node):
    print(f"[LAVALINK] node unavailable: label={getattr(node, 'label', '-')}")

@bot.event
async def on_websocket_closed(event):
    player = getattr(event, "player", None)
    guild_id = getattr(getattr(player, "guild", None), "id", None)
    print(
        f"[LAVALINK] websocket closed: guild={guild_id} code={getattr(event, 'code', None)} "
        f"reason={getattr(event, 'reason', None)} by_discord={getattr(event, 'by_discord', None)}"
    )

@bot.event
async def on_track_exception(event):
    player = getattr(event, "player", None)
    guild_id = getattr(getattr(player, "guild", None), "id", None)
    print(f"[LAVALINK] track exception: guild={guild_id} error={getattr(event, 'exception', None)}")

@bot.event
async def on_voice_state_update(member, before, after):
    """봇의 VC 상태 변화 추적: 목표 채널 저장 + 끊긴 경우 빠른 재접속."""
    if member != bot.user:
        return

    log_voice_snapshot(
        "bot voice_state_update",
        member.guild,
        member.guild.voice_client,
        extra=f"before=({describe_voice_state(before)}) after=({describe_voice_state(after)}) target_channel_id={bot.get_voice_target(member.guild.id)}",
    )

    if after.channel is not None:
        bot.set_voice_target(member.guild.id, after.channel.id)
        bot.reset_voice_retry(member.guild.id)
        bot.mark_voice_connected(member.guild.id)
        await ensure_silence_keepalive(
            member.guild,
            member.guild.voice_client,
            reason="voice_state_update_connected",
        )
        return

    target_channel_id = bot.get_voice_target(member.guild.id)
    if not target_channel_id:
        return

    bot.mark_voice_disconnected(member.guild.id)
    log_voice_debug(
        f"[VC] bot 음성 채널 이탈 감지 ({member.guild.name}): "
        f"target_channel_id={target_channel_id}, "
        f"discord.py 내부 복구 우선, {VOICE_LIBRARY_RECOVERY_GRACE:.0f}초 후 watchdog 개입"
    )
    asyncio.create_task(attempt_voice_state_rejoin(member.guild.id, target_channel_id))

@bot.event
async def on_message(message):
    if message.author.bot: return
    if not message.guild: return

    # 사운드보드 업로드 채널 체크
    config = bot.get_config(message.guild.id)
    sb_channel_id = config.get("sb_channel_id")
    if sb_channel_id and message.channel.id == sb_channel_id and message.attachments:
        dl_dir = os.path.join(SOUND_DIR, "discord")
        os.makedirs(dl_dir, exist_ok=True)
        # 메시지 텍스트가 있으면 파일명으로 사용 (한글 파일명 깨짐 방지)
        custom_names = [n.strip() for n in message.content.split(",") if n.strip()] if message.content.strip() else []
        added = []
        for i, att in enumerate(message.attachments):
            ext = os.path.splitext(att.filename)[1].lower()
            if ext in SOUND_EXTENSIONS:
                if i < len(custom_names):
                    filename = custom_names[i] + ext
                else:
                    filename = att.filename
                filename = unicodedata.normalize("NFC", filename)
                save_path = os.path.join(dl_dir, filename)
                await att.save(save_path)
                name = "discord/" + os.path.splitext(filename)[0].lower()
                added.append(name)
        if added:
            bot.sound_files = scan_sound_files()
            tip = "\n💡 파일명 지정: 메시지에 이름을 적으면 그 이름으로 저장됩니다." if not custom_names else ""
            reply = await message.reply(f"사운드 추가됨: **{', '.join(added)}**\n현재 {len(bot.sound_files)}개{tip}", mention_author=False)
            # 5초 후 원본 메시지 + 확인 메시지 삭제
            await asyncio.sleep(5)
            try:
                await message.delete()
            except Exception as e:
                print(f"[SB] 업로드 메시지 삭제 실패: {e}")
            try:
                await reply.delete()
            except Exception as e:
                print(f"[SB] 응답 메시지 삭제 실패: {e}")
        return

    if message.channel.id != config.get("channel_id"): return

    current_display_name = get_server_display_name(message.guild, message.author)
    log_time = datetime.now().strftime('%H:%M:%S')
    lock = bot.get_channel_lock(message.channel.id)
    print(f"[{log_time}] [{message.guild.name}] {current_display_name}: {message.content}")
    print(
        f"[DISPATCH] msg={message.id} attachments={len(message.attachments)} "
        f"queued={lock.locked()}"
    )
    lock_wait_started = time_module.perf_counter()
    async with lock:
        lock_wait_ms = round((time_module.perf_counter() - lock_wait_started) * 1000)
        if lock_wait_ms >= 100:
            print(f"[DISPATCH] msg={message.id} lock acquired after {lock_wait_ms}ms")
        async with message.channel.typing():
            try:
                server_text_model = config.get("text_model", TEXT_MODEL)
                server_image_model = config.get("image_model", IMAGE_MODEL)
                history_limit = get_history_limit(config)
                image_history_limit = get_image_history_limit(config)
                history_scan_limit = max(history_limit, image_history_limit)
                reset_ts = datetime.fromisoformat(config["reset_time"]) if config.get("reset_time") else None
                reply_target_id = message.reference.message_id if message.reference and message.reference.message_id else None

                history_entries = []
                history_texts = []
                history_position = 0
                history_inline_images = 0
                history_ref_images = []

                async for msg in message.channel.history(limit=history_scan_limit + 10):
                    if msg.id == message.id: continue
                    if msg.id > message.id and msg.author != bot.user: continue
                    if reset_ts and msg.created_at < reset_ts: break
                    interaction_output = msg.author == bot.user and getattr(msg, "interaction_metadata", None) is not None
                    has_image_output = any(is_image_attachment(att) for att in msg.attachments)
                    if interaction_output and not has_image_output:
                        continue
                    if msg.author == bot.user and not msg.content and not msg.attachments:
                        continue

                    within_text_history = history_position < history_limit
                    within_image_history = history_position < image_history_limit

                    cached_doc_blocks = []
                    if within_text_history and msg.author != bot.user:
                        replay_full_docs = (reply_target_id == msg.id) or (history_position < DOC_CACHE_AUTO_REPLAY_WINDOW)
                        cached_doc_blocks = bot.doc_cache.render_blocks(msg.id, full_content=replay_full_docs)

                    image_parts = []
                    if msg.author == bot.user:
                        content = msg.content if within_text_history else ""
                        if msg.attachments:
                            for att in msg.attachments:
                                if is_image_attachment(att):
                                    if within_image_history:
                                        d, m = await download_media(att.url)
                                        if d:
                                            raw = d.getvalue()
                                            m = infer_attachment_mime(att, m)
                                            image_parts.append(types.Part(inline_data=types.Blob(mime_type=m, data=raw)))
                                            history_inline_images += 1
                                            try:
                                                history_ref_images.append(Image.open(io.BytesIO(raw)).copy())
                                            except Exception:
                                                pass
                    else:
                        content = msg.content if within_text_history else ""
                        if msg.attachments:
                            for att in msg.attachments:
                                if is_image_attachment(att):
                                    if within_image_history:
                                        d, m = await download_media(att.url)
                                        if d:
                                            raw = d.getvalue()
                                            m = infer_attachment_mime(att, m)
                                            image_parts.append(types.Part(inline_data=types.Blob(mime_type=m, data=raw)))
                                            history_inline_images += 1
                                            try:
                                                history_ref_images.append(Image.open(io.BytesIO(raw)).copy())
                                            except Exception:
                                                pass
                                elif within_text_history:
                                    file_text, file_kind = await extract_attachment_text(att)
                                    if file_text and not bot.doc_cache.has_source(msg.id, "text_file", att.filename):
                                        bot.doc_cache.store_document("text_file", att.filename, file_text, msg.id, msg.channel.id)
                                        content += f'\n\nAttached file "{att.filename}":\n\n{file_text}'

                    if cached_doc_blocks:
                        content += "\n\n" + "\n\n".join(cached_doc_blocks)

                    role = "model" if msg.author == bot.user else "user"
                    parts = []
                    if content.strip():
                        parts.append(types.Part(text=content))
                    parts.extend(image_parts)
                    if not parts:
                        history_position += 1
                        if history_position >= history_scan_limit:
                            break
                        continue
                    history_entries.insert(
                        0,
                        (
                            types.Content(role=role, parts=parts),
                            msg.author.id if role == "user" else None,
                            get_server_display_name(message.guild, msg.author) if role == "user" else None,
                        ),
                    )
                    history_texts.insert(0, content)
                    history_position += 1
                    if history_position >= history_scan_limit:
                        break

                print(
                    f"[HISTORY] count={len(history_entries)} text_window={history_limit} "
                    f"image_window={image_history_limit} inline_images={history_inline_images}"
                )

                current_text = message.content.strip()
                current_parts = [types.Part(text=current_text)] if current_text else []
                current_ref_images = []
                has_current_images = False
                failed_attachments = []

                urls = re.findall(r'(https?://\S+)', message.content)
                for url in urls:
                    web_content = await scrap_webpage(url)
                    if web_content:
                        doc_id = bot.doc_cache.store_document("webpage", url, web_content, message.id, message.channel.id)
                        print(f"[CACHE] webpage 저장: {doc_id} ({url})")
                        current_parts.append(types.Part(text=f'Webpage "{url}":\n\n{web_content}'))
                        print(f"[WEB] 웹사이트 내용 읽음: {url}")

                if message.attachments:
                    for att in message.attachments:
                        if is_image_attachment(att):
                            d, m = await download_media(att.url)
                            if d:
                                raw = d.getvalue()
                                m = infer_attachment_mime(att, m)
                                current_parts.append(types.Part(inline_data=types.Blob(mime_type=m, data=raw)))
                                has_current_images = True
                                try:
                                    current_ref_images.append(Image.open(io.BytesIO(raw)).copy())
                                except Exception:
                                    pass
                            else:
                                failed_attachments.append(att.filename)
                        elif is_audio_or_video_attachment(att):
                            d, m = await download_media(att.url)
                            if d:
                                file_bytes = d.getvalue()
                                media_type = attachment_media_type(att)
                                m = infer_attachment_mime(att, m)
                                if len(file_bytes) < 15 * 1024 * 1024:
                                    current_parts.append(types.Part(inline_data=types.Blob(mime_type=m, data=file_bytes)))
                                else:
                                    file_ref = await upload_file_api(file_bytes, m)
                                    if file_ref:
                                        current_parts.append(types.Part(file_data=types.FileData(file_uri=file_ref.uri, mime_type=m)))
                                    else:
                                        failed_attachments.append(att.filename)
                                print(f"[{media_type.upper()}] {media_type} 첨부: {att.filename} ({len(file_bytes)} bytes)")
                            else:
                                failed_attachments.append(att.filename)
                        else:
                            file_content, file_kind = await extract_attachment_text(att)
                            if file_content:
                                doc_id = bot.doc_cache.store_document("text_file", att.filename, file_content, message.id, message.channel.id)
                                print(f"[CACHE] {file_kind or 'text'} file 저장: {doc_id} ({att.filename})")
                                current_parts.append(types.Part(text=f'Attached file "{att.filename}":\n\n{file_content}'))
                            else:
                                failed_attachments.append(att.filename)

                if failed_attachments:
                    print(f"[WARN] 첨부 파일 읽기 실패: {', '.join(failed_attachments)}")
                if not current_parts:
                    await message.channel.send("첨부 파일을 읽지 못했습니다. 파일을 다시 올려주세요.")
                    return

                route = {
                    "action": "respond_text",
                    "refined_prompt": "",
                    "use_recent_images": False,
                    "reason": "action_router_skipped",
                }
                action_router_enabled = is_enabled(config.get("action_router", DEFAULT_ACTION_ROUTER_ENABLED))
                recent_images_available = bool(history_ref_images)
                recent_context = build_recent_conversation_excerpt(history_texts)
                if action_router_enabled and should_call_action_router(
                    message.content,
                    has_current_images=has_current_images,
                    has_recent_images=recent_images_available,
                    allow_image_generation=bool(server_image_model),
                    allow_recent_images=True,
                ):
                    route = await route_message_action(
                        message.content,
                        has_current_images=has_current_images,
                        has_recent_images=recent_images_available,
                        text_model=server_text_model,
                        recent_context=recent_context,
                    )

                if route.get("action") == "generate_image":
                    display_prompt = route.get("refined_prompt") or message.content
                    generation_prompt = build_contextual_image_prompt(route.get("refined_prompt"), message.content, history_texts)
                    ref_images = []
                    if current_ref_images:
                        ref_images = current_ref_images
                    elif route.get("use_recent_images"):
                        ref_images = list(history_ref_images)

                    print(f"[IMG] 라우팅 생성 요청: '{display_prompt}' (현재첨부={len(current_ref_images)} 최근참조={len(ref_images)} ctx={len(recent_context)})")
                    status_embed = discord.Embed(
                        title="이미지 생성 중",
                        description=compact_prompt_for_status(display_prompt),
                    )
                    status_msg = await message.channel.send(embed=status_embed)
                    req_content = ref_images + [generation_prompt]
                    try:
                        img_data = await call_nano_banana(req_content, image_model=server_image_model)
                        with io.BytesIO(img_data) as binary:
                            binary.seek(0)
                            fn = "edited.png" if ref_images else "generated.png"
                            await message.channel.send(file=discord.File(fp=binary, filename=fn))
                        status_embed.title = "이미지 생성 완료"
                        await status_msg.edit(embed=status_embed)
                        print(f"[BOT]({server_image_model}): [이미지 생성 완료]")
                    except Exception as e:
                        err_msg = f"이미지 생성 실패: {str(e)}"
                        print(err_msg)
                        status_embed.title = "이미지 생성 실패"
                        status_embed.description = str(e)
                        await status_msg.edit(embed=status_embed)
                    return

                conversation_entries = history_entries + [
                    (
                        types.Content(role="user", parts=current_parts),
                        message.author.id,
                        current_display_name,
                    )
                ]
                full_contents = merge_adjacent_contents(
                    label_discord_user_contents(conversation_entries)
                )

                response, token_usage, call_trace = await call_gemini(
                    full_contents,
                    text_model=server_text_model,
                    thinking_level=config.get("thinking_level"),
                )

                status = get_interaction_status(response)
                if status not in {"completed", "incomplete"}:
                    error = get_interaction_error_message(response)
                    print(
                        "[WARN] Gemini non-completed response "
                        f"response={summarize_interaction_response(response)} "
                        f"call={json.dumps(call_trace, ensure_ascii=False, default=str, separators=(',', ':'))}"
                    )
                    detail = f": {error}" if error else ""
                    await message.channel.send(f"응답 생성이 중단되었습니다 (status={status}){detail}")
                    return
                if status == "incomplete":
                    print("[WARN] Gemini interaction이 incomplete 상태로 종료되어 부분 응답을 전송합니다.")

                text_reply = extract_interaction_text(response)
                if not text_reply:
                    error = get_interaction_error_message(response)
                    failure_context = {
                        "message_id": message.id,
                        "guild_id": message.guild.id,
                        "channel_id": message.channel.id,
                        "history_entries": len(history_entries),
                        "history_inline_images": history_inline_images,
                        "current_images": len(current_ref_images),
                        "route": route.get("action"),
                        "call": call_trace,
                    }
                    print(
                        "[WARN] Gemini content-free response "
                        f"response={summarize_interaction_response(response)} "
                        f"context={json.dumps(failure_context, ensure_ascii=False, default=str, separators=(',', ':'))}"
                    )
                    detail = f" ({error})" if error else ""
                    await message.channel.send(f"Gemini가 텍스트 응답을 생성하지 못했습니다{detail}")
                    return

                print(f"[BOT]({server_text_model}): {text_reply} (Tokens: {token_usage}, Status: {status})")
                await send_long_message(message.channel, text_reply)

                # video_calls 처리 — 비용 절감으로 임시 비활성화

                print("-------")

            except Exception as e:
                error_detail = describe_gemini_exception(e)
                print(f"[CRITICAL] Error: {error_detail}")
                traceback.print_exc()
                await message.channel.send(f"오류가 발생했습니다: {error_detail[:1500]}")

# ================= 5. 서버 설정 명령어 =================

@bot.tree.command(name="set", description="현재 채널을 봇 활성화 채널로 설정")
async def set_ch(interaction: discord.Interaction):
    if not interaction.guild:
        await temp_reply(interaction, "[ERROR] 서버에서만 사용 가능합니다.")
        return
    config = bot.get_config(interaction.guild_id)
    config["channel_id"] = interaction.channel_id
    bot.save_config()
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    await interaction.response.send_message(
        f"[OK] {actor_name}님이 봇 활성화 채널을 {interaction.channel.mention} 으로 설정했습니다."
    )


@bot.tree.command(name="history", description="Gemini 대화/이미지 히스토리 참조 개수 설정")
@app_commands.describe(
    count=f"텍스트로 참고할 최근 메시지 개수 ({MIN_HISTORY_LIMIT}~{MAX_HISTORY_LIMIT}, 사용자+봇 합산)",
    image_count=f"이미지를 실제 재첨부할 최근 메시지 개수 ({MIN_IMAGE_HISTORY_LIMIT}~{MAX_IMAGE_HISTORY_LIMIT}, 0=끄기)",
)
async def set_history(
    interaction: discord.Interaction,
    count: int | None = None,
    image_count: int | None = None,
):
    if not interaction.guild:
        await temp_reply(interaction, "[ERROR] 서버에서만 사용 가능합니다.")
        return
    config = bot.get_config(interaction.guild_id)
    if count is None and image_count is None:
        await interaction.response.send_message(
            f"현재 Gemini 원문 히스토리 참조: **{get_history_limit(config)}개** "
            f"(범위 {MIN_HISTORY_LIMIT}~{MAX_HISTORY_LIMIT}, 사용자+봇 합산)\n"
            f"현재 Gemini 이미지 히스토리 재첨부: **{get_image_history_limit(config)}개** "
            f"(범위 {MIN_IMAGE_HISTORY_LIMIT}~{MAX_IMAGE_HISTORY_LIMIT}, 0=끄기)"
        )
        return
    changed = []
    if count is not None:
        limit = clamp_history_limit(count)
        config["history_limit"] = limit
        suffix = "" if limit == count else f" (요청값 {count}에서 허용 범위로 보정)"
        changed.append(f"원문 히스토리 **{limit}개**{suffix}")
    if image_count is not None:
        image_limit = clamp_image_history_limit(image_count)
        config["image_history_limit"] = image_limit
        suffix = "" if image_limit == image_count else f" (요청값 {image_count}에서 허용 범위로 보정)"
        changed.append(f"이미지 히스토리 **{image_limit}개**{suffix}")
    bot.save_config()
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    await interaction.response.send_message(
        f"[OK] {actor_name}님이 Gemini " + ", ".join(changed) + "로 설정했습니다."
    )

@bot.tree.command(name="reset", description="히스토리 초기화")
async def reset_history(interaction: discord.Interaction):
    if not interaction.guild:
        await temp_reply(interaction, "[ERROR] 서버에서만 사용 가능합니다.")
        return
    config = bot.get_config(interaction.guild_id)
    config["reset_time"] = datetime.now(timezone.utc).isoformat()
    bot.save_config()
    await interaction.response.send_message("[OK] 히스토리 초기화")

# ================= 6. User Install 슬래시 커맨드 =================

@bot.tree.command(name="chat", description="AI와 대화합니다 (이미지 분석 가능)")
@app_commands.allowed_installs(guilds=True, users=True)
@app_commands.allowed_contexts(guilds=True, dms=True, private_channels=True)
@app_commands.describe(
    message="질문 또는 메시지",
    image="분석할 이미지 파일 (선택)",
    image_url="분석할 이미지 URL (선택, Discord에 업로드 후 링크 복사)"
)
async def chat_command(interaction: discord.Interaction, message: str, image: discord.Attachment = None, image_url: str = None):
    await interaction.response.defer()

    log_time = datetime.now().strftime('%H:%M:%S')
    location = f"[{interaction.guild.name}]" if interaction.guild else "[DM]"
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    print(f"[{log_time}] {location} [/chat] {actor_name}: {message}")

    try:
        parts = [types.Part(text=message)]

        img_data = None
        mime_type = None
        image_source = None

        if image and is_image_attachment(image):
            img_data, mime_type = await download_media(image.url)
            if img_data:
                mime_type = infer_attachment_mime(image, mime_type)
                parts.append(types.Part(inline_data=types.Blob(mime_type=mime_type, data=img_data.getvalue())))
                image_source = image.url
                print(f"  [IMG] 이미지 첨부: {image.url}")
        elif image_url:
            img_data, mime_type = await download_media(image_url)
            if img_data:
                parts.append(types.Part(inline_data=types.Blob(mime_type=mime_type, data=img_data.getvalue())))
                image_source = image_url
                print(f"  [IMG] 이미지 URL: {image_url}")

        # 사용자 입력 표시
        user_display = f"**{actor_name}**: {message}"
        if image_source:
            embed = discord.Embed(description=user_display)
            embed.set_image(url=image_source)
            await interaction.followup.send(embed=embed)
        else:
            await interaction.followup.send(user_display)

        contents = [types.Content(role="user", parts=parts)]

        response = await create_model_interaction(
            model=TEXT_MODEL,
            input_data=gemini_contents_to_interaction_steps(contents),
            system_instruction=CHAT_SYSTEM_PROMPT,
            tools=[GOOGLE_SEARCH_TOOL],
            thinking_level=MAX_THINKING_LEVEL,
        )
        token_usage = get_interaction_usage(response)

        status = get_interaction_status(response)
        if status not in {"completed", "incomplete"}:
            await interaction.followup.send(f"[ERROR] 응답 생성 중단 (status={status})")
            return
        if status == "incomplete":
            print("  [WARN] Gemini interaction이 incomplete 상태로 종료되어 부분 응답을 전송합니다.")

        text_reply = extract_interaction_text(response)

        if text_reply.strip():
            for chunk in split_discord_message(text_reply):
                await interaction.followup.send(chunk)

            print(f"  [BOT]({TEXT_MODEL}): {text_reply}")
            print(f"  [TOKEN]: {token_usage}")
        else:
            await interaction.followup.send("[WARN] 응답을 생성하지 못했습니다.")

        print("-------")

    except Exception as e:
        print(f"  [ERROR]: {e}")
        print("-------")
        await interaction.followup.send(f"[ERROR] 오류 발생: {e}")


@bot.tree.command(name="image", description="이미지를 생성하거나 편집합니다")
@app_commands.allowed_installs(guilds=False, users=True)
@app_commands.allowed_contexts(guilds=True, dms=True, private_channels=True)
@app_commands.describe(
    prompt="생성할 이미지에 대한 설명",
    reference="편집할 원본 이미지 파일 (선택)",
    reference_url="편집할 이미지 URL (선택, Discord에 업로드 후 링크 복사)"
)
async def image_command(interaction: discord.Interaction, prompt: str, reference: discord.Attachment = None, reference_url: str = None):
    await interaction.response.defer()

    log_time = datetime.now().strftime('%H:%M:%S')
    location = f"[{interaction.guild.name}]" if interaction.guild else "[DM]"
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    print(f"[{log_time}] {location} [/image] {actor_name}: {prompt}")

    try:
        ref_images = []
        has_reference = False
        reference_source = None

        if reference and is_image_attachment(reference):
            img_data, _ = await download_media(reference.url)
            if img_data:
                ref_images.append(Image.open(img_data))
                has_reference = True
                reference_source = reference.url
                print(f"  [IMG] 참조 이미지: {reference.url}")
        elif reference_url:
            img_data, _ = await download_media(reference_url)
            if img_data:
                ref_images.append(Image.open(img_data))
                has_reference = True
                reference_source = reference_url
                print(f"  [IMG] 참조 이미지 URL: {reference_url}")

        user_display = f"**{actor_name}**: {prompt}"
        if reference_source:
            embed = discord.Embed(description=user_display)
            embed.set_image(url=reference_source)
            embed.set_footer(text="원본 이미지")
            await interaction.followup.send(embed=embed)
        else:
            await interaction.followup.send(user_display)

        user_message = f"User request: {prompt}"
        if has_reference:
            user_message += "\nA reference image is provided for editing."
        user_message += "\n\nReturn the refined image prompt and an optional brief user-facing message."

        response = await create_model_interaction(
            model=TEXT_MODEL,
            input_data=user_message,
            system_instruction=IMAGE_SYSTEM_PROMPT,
            response_schema=IMAGE_PROMPT_SCHEMA,
            thinking_level=MAX_THINKING_LEVEL,
        )

        status = get_interaction_status(response)
        if status not in {"completed", "incomplete"}:
            await interaction.followup.send(f"[ERROR] 이미지 프롬프트 생성 중단 (status={status})")
            return
        if status == "incomplete":
            print("  [WARN] 이미지 프롬프트 interaction이 incomplete 상태로 종료되었습니다.")

        prompt_payload = coerce_json_object(extract_interaction_text(response)) or {}
        generated_prompt = (prompt_payload.get("prompt") or prompt).strip()
        text_reply = (prompt_payload.get("message") or "").strip()

        print(f"  [PROMPT] 정제된 프롬프트: {generated_prompt}")

        if text_reply.strip():
            if len(text_reply) > 1900:
                text_reply = text_reply[:1900] + "..."
            await interaction.followup.send(text_reply)
            print(f"  [BOT]({TEXT_MODEL}): {text_reply}")

        status_msg = await interaction.followup.send(f"**Generating:** `{generated_prompt}`")

        req_content = ref_images + [generated_prompt] if ref_images else [generated_prompt]
        img_data = await call_nano_banana(req_content)
        with io.BytesIO(img_data) as binary:
            binary.seek(0)
            fn = "edited.png" if ref_images else "generated.png"
            await interaction.followup.send(file=discord.File(fp=binary, filename=fn))

        await status_msg.edit(content=f"[OK] **Generated:** `{generated_prompt}`")
        print(f"  [BOT]({IMAGE_MODEL}): [이미지 생성 완료]")
        print("-------")

    except Exception as e:
        print(f"  [ERROR]: {e}")
        print("-------")
        await interaction.followup.send(f"[ERROR] 이미지 생성 실패: {e}")

# ================= 7. 비디오 생성 (User Install) =================

# ================= /video 커맨드 — 비용 절감으로 임시 비활성화 =================
# 복원 시 VIDEO_GEN_MODEL 주석 해제 및 아래 코드 복원 필요

# ================= 8. 사운드보드 =================

SOUNDS_PER_PAGE = 25

SB_PLAY_LOG = []  # [{"user": name, "sound": key, "time": str}, ...] 최근 20개
SB_LOG_MAX = 20

class VolumeModal(discord.ui.Modal, title="사운드보드 볼륨 설정"):
    def __init__(self, soundboard_view):
        super().__init__()
        self.soundboard_view = soundboard_view
        self.value = discord.ui.TextInput(
            label="볼륨 (1~100%)",
            placeholder=str(int(bot.sound_volume * 100)),
            max_length=3,
        )
        self.add_item(self.value)

    async def on_submit(self, interaction: discord.Interaction):
        try:
            percent = max(1, min(100, int(self.value.value)))
        except ValueError:
            await temp_reply(interaction, "숫자를 입력해주세요.")
            return
        bot.set_sound_volume(percent / 100.0)
        self.soundboard_view._refresh_keys()
        self.soundboard_view._update_components()
        await interaction.response.edit_message(embed=self.soundboard_view._make_embed(), view=self.soundboard_view)


@bot.tree.command(name="sb", description="사운드보드 (이름 입력시 재생, 빈 입력시 현재 채널에 통합 패널 생성)")
@app_commands.allowed_installs(guilds=True, users=False)
@app_commands.allowed_contexts(guilds=True, dms=False, private_channels=False)
@app_commands.describe(name="재생할 사운드 이름 (비우면 현재 채널을 패널 채널로 설정)")
async def soundboard(interaction: discord.Interaction, name: str = None):
    if not interaction.guild:
        await temp_reply(interaction, "서버에서만 사용 가능합니다.")
        return

    # 인자 없으면 현재 채널을 패널 채널로 설정하고 통합 패널 준비
    if not name:
        await interaction.response.defer()
        try:
            panel_status = await ensure_control_panel(interaction.guild, interaction.channel, save_target=True)
            await temp_followup(
                interaction,
                f"{interaction.channel.mention}에 통합 패널을 {'갱신' if panel_status == 'updated' else '생성'}했습니다.",
            )
        except Exception as e:
            await temp_followup(interaction, f"[ERROR] 통합 패널 생성 실패: {e}")
        return

    # 인자 있으면 바로 재생
    name = unicodedata.normalize("NFC", name.replace("📁 ", "").strip())

    member = interaction.guild.get_member(interaction.user.id)
    if not member or not member.voice:
        await temp_reply(interaction, "음성 채널에 먼저 접속해주세요.")
        return

    await interaction.response.defer()
    try:
        success, msg = await play_sound_in_vc(interaction.guild, member, name, bot)
        if success:
            SB_PLAY_LOG.append({
                "user": get_server_display_name(interaction.guild, interaction.user),
                "sound": name,
                "time": datetime.now(KST).strftime("%H:%M:%S")
            })
            if len(SB_PLAY_LOG) > SB_LOG_MAX:
                SB_PLAY_LOG.pop(0)
        await temp_followup(interaction, msg)
    except Exception as e:
        await temp_followup(interaction, f"[ERROR] {e}")

@soundboard.autocomplete('name')
async def sb_autocomplete(interaction: discord.Interaction, current: str):
    current = unicodedata.normalize("NFC", current.lower())
    all_keys = sorted(bot.sound_files.keys())
    choices = []
    for key in all_keys:
        if current and current not in key:
            continue
        display = key if len(key) <= 100 else "..." + key[-97:]
        choices.append(app_commands.Choice(name=display, value=key))
    return choices[:25]

class OffsetModal(discord.ui.Modal, title="오프셋 설정"):
    def __init__(self, sibo_view):
        super().__init__()
        self.sibo_view = sibo_view
        cfg = sibo_view._get_cfg()
        current = str(cfg.get("offset", 0.0)) if cfg else "0"
        self.value = discord.ui.TextInput(label="정각 몇 초 전에 재생 (소수점 가능, 0~30)", placeholder=current, max_length=10)
        self.add_item(self.value)

    async def on_submit(self, interaction: discord.Interaction):
        try:
            offset = max(0.0, min(30.0, float(self.value.value)))
            gid = self.sibo_view.guild_id
            cfg = bot.get_sibo(gid)
            if not cfg or not cfg.get("sound"):
                await temp_reply(interaction, "먼저 음원을 선택해주세요.")
                return
            cfg["offset"] = offset
            bot.save_config()
            restart_sibo()
            self.sibo_view._update_components()
            await interaction.response.edit_message(embed=self.sibo_view._make_embed(), view=self.sibo_view)
        except ValueError:
            await temp_reply(interaction, "숫자를 입력해주세요.")

class CompModal(discord.ui.Modal, title="딜레이 보정값 설정"):
    def __init__(self, sibo_view):
        super().__init__()
        self.sibo_view = sibo_view
        self.value = discord.ui.TextInput(label="보정 초 (소수점 가능, 음수도 가능)", placeholder=str(SIBO_LATENCY_COMP), max_length=10)
        self.add_item(self.value)

    async def on_submit(self, interaction: discord.Interaction):
        global SIBO_LATENCY_COMP
        try:
            SIBO_LATENCY_COMP = float(self.value.value)
            g = bot.server_configs.setdefault("_global", {})
            g["sibo_comp"] = SIBO_LATENCY_COMP
            bot.save_config()
            # comp는 다음 사이클에서 자연 반영, restart 불필요
            self.sibo_view._update_components()
            await interaction.response.edit_message(embed=self.sibo_view._make_embed(), view=self.sibo_view)
        except ValueError:
            await temp_reply(interaction, "숫자를 입력해주세요.")


class ControlPanelView(discord.ui.View):
    def __init__(self, guild_id, mode="soundboard"):
        super().__init__(timeout=None)
        self.guild_id = str(guild_id)
        self.mode = mode
        self.folder = ""
        self.search = ""
        self.page = 0
        self.show_log = False
        self.sibo_folder = ""
        self.sound_page = 0
        self._refresh_keys()
        self._rebuild_sound_keys()
        self._update_components()

    def _get_cfg(self):
        return bot.get_sibo(self.guild_id)

    def _refresh_keys(self):
        all_keys = sorted(bot.sound_files.keys())
        if self.folder:
            all_keys = [k for k in all_keys if k.startswith(self.folder + "/")]
        if self.search:
            all_keys = [k for k in all_keys if self.search in k]
        self.original_keys = sorted(bot.sound_files.keys())
        self.all_keys = all_keys
        self.total_pages = max(1, (len(all_keys) + SOUNDS_PER_PAGE - 1) // SOUNDS_PER_PAGE)
        self.page = min(self.page, self.total_pages - 1)

    def _page_keys(self):
        start = self.page * SOUNDS_PER_PAGE
        return self.all_keys[start:start + SOUNDS_PER_PAGE]

    def _rebuild_sound_keys(self):
        all_keys = sorted(bot.sound_files.keys())
        if self.sibo_folder:
            self.sound_keys = [k for k in all_keys if k.startswith(self.sibo_folder + "/")]
        else:
            self.sound_keys = all_keys
        self.sound_total_pages = max(1, (len(self.sound_keys) + 24) // 25)
        self.sound_page = min(self.sound_page, self.sound_total_pages - 1)

    def _page_sound_keys(self):
        start = self.sound_page * 25
        return self.sound_keys[start:start + 25]

    def _make_embed(self):
        if self.mode == "sibo":
            return self._make_sibo_embed()
        return self._make_soundboard_embed()

    def _make_soundboard_embed(self):
        if self.show_log:
            embed = discord.Embed(title="🎛 Gemini 패널 — 사운드보드 로그", color=0x5865F2)
            if SB_PLAY_LOG:
                lines = [f"`{e['time']}` **{e['user']}** → `{e['sound']}`" for e in reversed(SB_PLAY_LOG)]
                embed.description = "\n".join(lines)
            else:
                embed.description = "(재생 기록 없음)"
            embed.set_footer(text=f"최근 {len(SB_PLAY_LOG)}개 | 시보 설정 버튼으로 전환")
            return embed

        page_keys = self._page_keys()
        folder_label = self.folder or "전체"
        listing = "\n".join(f"`{k}`" for k in page_keys) or "(비어있음)"
        embed = discord.Embed(title=f"🎛 Gemini 패널 — 사운드보드/{folder_label}", description=listing, color=0x5865F2)
        if SB_PLAY_LOG:
            last = SB_PLAY_LOG[-1]
            embed.set_footer(text=f"총 {len(self.all_keys)}개 | 페이지 {self.page + 1}/{self.total_pages} | 최근: {last['user']} → {last['sound']}")
        else:
            embed.set_footer(text=f"총 {len(self.all_keys)}개 | 페이지 {self.page + 1}/{self.total_pages}")
        return embed

    def _make_sibo_embed(self):
        cfg = self._get_cfg()
        embed = discord.Embed(title="🎛 Gemini 패널 — 시보 설정", color=0xF1C40F)
        if cfg:
            sound = cfg.get("sound", "없음")
            offset = cfg.get("offset", 0)
            vc_id = cfg.get("vc_channel")
            vc_text = f"<#{vc_id}>" if vc_id else "미설정"
            embed.add_field(name="음원", value=f"`{sound}`", inline=True)
            embed.add_field(name="오프셋", value=f"정각 {offset}초 전", inline=True)
            embed.add_field(name="VC 채널", value=vc_text, inline=True)
            embed.add_field(name="딜레이 보정", value=f"{SIBO_LATENCY_COMP}초", inline=True)
            embed.add_field(name="상태", value="✅ 활성", inline=True)
        else:
            embed.description = "시보가 설정되어 있지 않습니다."
            embed.add_field(name="딜레이 보정", value=f"{SIBO_LATENCY_COMP}초", inline=True)
        folder_label = self.sibo_folder or "전체"
        embed.set_footer(text=f"시보 음원 폴더: {folder_label} | {len(self.sound_keys)}개 | 페이지 {self.sound_page + 1}/{self.sound_total_pages}")
        return embed

    def _update_components(self):
        self.clear_items()
        if self.mode == "sibo":
            self._add_sibo_components()
        else:
            self._add_soundboard_components()

    def _folder_options(self, current_folder):
        folders = sorted(set(k.split("/")[0] for k in bot.sound_files.keys() if "/" in k))
        if not folders:
            return []
        options = [discord.SelectOption(label="📂 전체", value="_all", default=current_folder == "")]
        for folder in folders[:24]:
            options.append(discord.SelectOption(label=f"📁 {folder}", value=folder, default=current_folder == folder))
        return options

    def _add_soundboard_components(self):
        folder_options = self._folder_options(self.folder)
        if folder_options:
            folder_select = discord.ui.Select(placeholder="📁 폴더 선택", options=folder_options, row=0)
            folder_select.callback = self._soundboard_folder_callback
            self.add_item(folder_select)

        page_keys = self._page_keys()
        if page_keys:
            sound_options = []
            for key in page_keys:
                display = key.split("/")[-1] if "/" in key else key
                sound_options.append(discord.SelectOption(label=display[:100], value=key, description=key[:100] if "/" in key else None))
            sound_select = discord.ui.Select(placeholder="🔊 재생할 사운드 선택", options=sound_options, row=1)
            sound_select.callback = self._play_callback
            self.add_item(sound_select)

        if self.total_pages > 1:
            prev_btn = discord.ui.Button(label="◀", style=discord.ButtonStyle.secondary, disabled=self.page == 0, row=2)
            prev_btn.callback = self._prev_callback
            self.add_item(prev_btn)
            page_btn = discord.ui.Button(label=f"{self.page + 1}/{self.total_pages}", style=discord.ButtonStyle.primary, disabled=True, row=2)
            self.add_item(page_btn)
            next_btn = discord.ui.Button(label="▶", style=discord.ButtonStyle.secondary, disabled=self.page >= self.total_pages - 1, row=2)
            next_btn.callback = self._next_callback
            self.add_item(next_btn)

        leave_btn = discord.ui.Button(label="🚪 퇴장", style=discord.ButtonStyle.danger, row=2)
        leave_btn.callback = self._leave_callback
        self.add_item(leave_btn)

        log_label = "📋 목록" if self.show_log else "📋 로그"
        log_btn = discord.ui.Button(label=log_label, style=discord.ButtonStyle.primary if self.show_log else discord.ButtonStyle.secondary, row=3)
        log_btn.callback = self._log_callback
        self.add_item(log_btn)

        volume_btn = discord.ui.Button(label=f"🔉 볼륨: {int(bot.sound_volume * 100)}%", style=discord.ButtonStyle.secondary, row=3)
        volume_btn.callback = self._volume_callback
        self.add_item(volume_btn)

        sibo_btn = discord.ui.Button(label="⏰ 시보 설정", style=discord.ButtonStyle.secondary, row=4)
        sibo_btn.callback = self._show_sibo_callback
        self.add_item(sibo_btn)

    def _add_sibo_components(self):
        cfg = self._get_cfg()
        folder_options = self._folder_options(self.sibo_folder)
        if folder_options:
            folder_select = discord.ui.Select(placeholder="📁 폴더 필터", options=folder_options, row=0)
            folder_select.callback = self._sibo_folder_callback
            self.add_item(folder_select)

        page_keys = self._page_sound_keys()
        if page_keys:
            current_sound = cfg.get("sound") if cfg else None
            options = []
            for key in page_keys:
                desc = key[:100] if "/" in key else None
                short = key.split("/")[-1] if "/" in key else key
                options.append(discord.SelectOption(label=short[:100], value=key, description=desc, default=key == current_sound))
            sound_select = discord.ui.Select(placeholder="🔊 시보 음원 선택", options=options, row=1)
            sound_select.callback = self._sound_callback
            self.add_item(sound_select)

        if self.sound_total_pages > 1:
            prev_btn = discord.ui.Button(label="◀", style=discord.ButtonStyle.secondary, disabled=self.sound_page == 0, row=2)
            prev_btn.callback = self._sound_prev
            self.add_item(prev_btn)
            page_btn = discord.ui.Button(label=f"{self.sound_page + 1}/{self.sound_total_pages}", style=discord.ButtonStyle.primary, disabled=True, row=2)
            self.add_item(page_btn)
            next_btn = discord.ui.Button(label="▶", style=discord.ButtonStyle.secondary, disabled=self.sound_page >= self.sound_total_pages - 1, row=2)
            next_btn.callback = self._sound_next
            self.add_item(next_btn)

        current_offset = cfg.get("offset", 0) if cfg else 0
        offset_btn = discord.ui.Button(label=f"⏱️ 오프셋: {current_offset}s", style=discord.ButtonStyle.primary, row=3)
        offset_btn.callback = self._offset_modal_callback
        self.add_item(offset_btn)

        comp_btn = discord.ui.Button(label=f"⚙️ 보정: {SIBO_LATENCY_COMP}s", style=discord.ButtonStyle.secondary, row=3)
        comp_btn.callback = self._comp_callback
        self.add_item(comp_btn)

        vc_btn = discord.ui.Button(label="🔗 현재 VC 연결", style=discord.ButtonStyle.success, row=3)
        vc_btn.callback = self._vc_callback
        self.add_item(vc_btn)

        if cfg:
            off_btn = discord.ui.Button(label="❌ 비활성화", style=discord.ButtonStyle.danger, row=3)
            off_btn.callback = self._off_callback
            self.add_item(off_btn)

        sb_btn = discord.ui.Button(label="🔊 사운드보드", style=discord.ButtonStyle.secondary, row=4)
        sb_btn.callback = self._show_soundboard_callback
        self.add_item(sb_btn)

    async def _show_sibo_callback(self, interaction: discord.Interaction):
        self.mode = "sibo"
        self.show_log = False
        self._rebuild_sound_keys()
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _show_soundboard_callback(self, interaction: discord.Interaction):
        self.mode = "soundboard"
        self._refresh_keys()
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _soundboard_folder_callback(self, interaction: discord.Interaction):
        folder = interaction.data["values"][0]
        self.search = ""
        self.folder = "" if folder == "_all" else folder
        self.page = 0
        self._refresh_keys()
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _play_callback(self, interaction: discord.Interaction):
        sound_key = interaction.data["values"][0]
        member = interaction.guild.get_member(interaction.user.id)
        if not member or not member.voice:
            await temp_reply(interaction, "음성 채널에 먼저 접속해주세요.")
            return
        try:
            success, _ = await play_sound_in_vc(interaction.guild, member, sound_key, bot)
            if success:
                SB_PLAY_LOG.append({
                    "user": get_server_display_name(interaction.guild, interaction.user),
                    "sound": sound_key,
                    "time": datetime.now(KST).strftime("%H:%M:%S")
                })
                if len(SB_PLAY_LOG) > SB_LOG_MAX:
                    SB_PLAY_LOG.pop(0)
            self._refresh_keys()
            self._update_components()
            await interaction.response.edit_message(embed=self._make_embed(), view=self)
        except Exception as e:
            await temp_reply(interaction, f"[ERROR] {e}")

    async def _log_callback(self, interaction: discord.Interaction):
        self.show_log = not self.show_log
        self._refresh_keys()
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _prev_callback(self, interaction: discord.Interaction):
        self._refresh_keys()
        self.page = max(0, self.page - 1)
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _next_callback(self, interaction: discord.Interaction):
        self._refresh_keys()
        self.page = min(self.total_pages - 1, self.page + 1)
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _leave_callback(self, interaction: discord.Interaction):
        vc = interaction.guild.voice_client
        if vc and vc.is_connected():
            actor_name = get_server_display_name(interaction.guild, interaction.user)
            log_voice_snapshot("수동 퇴장 요청", interaction.guild, vc, extra=f"user={actor_name}")
            bot.clear_voice_target(interaction.guild.id)
            await vc.disconnect()
            await interaction.response.edit_message(embed=self._make_embed(), view=self)
        else:
            await temp_reply(interaction, "현재 음성 채널에 접속되어 있지 않습니다.")

    async def _volume_callback(self, interaction: discord.Interaction):
        modal = VolumeModal(self)
        await interaction.response.send_modal(modal)

    async def _sibo_folder_callback(self, interaction: discord.Interaction):
        folder = interaction.data["values"][0]
        self.sibo_folder = "" if folder == "_all" else folder
        self.sound_page = 0
        self._rebuild_sound_keys()
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _sound_prev(self, interaction: discord.Interaction):
        self.sound_page = max(0, self.sound_page - 1)
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _sound_next(self, interaction: discord.Interaction):
        self.sound_page = min(self.sound_total_pages - 1, self.sound_page + 1)
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _sound_callback(self, interaction: discord.Interaction):
        sound_key = interaction.data["values"][0]
        cfg = bot.get_sibo(self.guild_id) or {}
        vc_channel = cfg.get("vc_channel")
        offset = cfg.get("offset", 0.0)
        if not vc_channel:
            member = interaction.guild.get_member(interaction.user.id)
            if member and member.voice:
                vc_channel = member.voice.channel.id
        bot.set_sibo(self.guild_id, {"sound": sound_key, "offset": offset, "vc_channel": vc_channel})
        restart_sibo()
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _offset_modal_callback(self, interaction: discord.Interaction):
        modal = OffsetModal(self)
        await interaction.response.send_modal(modal)

    async def _comp_callback(self, interaction: discord.Interaction):
        modal = CompModal(self)
        await interaction.response.send_modal(modal)

    async def _vc_callback(self, interaction: discord.Interaction):
        cfg = bot.get_sibo(self.guild_id)
        if not cfg:
            await temp_reply(interaction, "먼저 음원을 선택해주세요.")
            return
        member = interaction.guild.get_member(interaction.user.id)
        if not member or not member.voice:
            await temp_reply(interaction, "음성 채널에 먼저 접속해주세요.")
            return
        cfg["vc_channel"] = member.voice.channel.id
        bot.save_config()
        restart_sibo()
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def _off_callback(self, interaction: discord.Interaction):
        bot.del_sibo(self.guild_id)
        restart_sibo()
        self._update_components()
        await interaction.response.edit_message(embed=self._make_embed(), view=self)

    async def on_timeout(self):
        self.clear_items()


# ================= Deep Research 명령어 =================

async def run_deep_research_command(interaction: discord.Interaction, prompt: str, agent_key: str):
    """대화 컨텍스트 없이 입력 프롬프트만 단독 리서치. 채널 락을 잡지 않으므로 일반 대화와 병렬로 돈다."""
    agent_id = DEEP_RESEARCH_AGENTS[agent_key]
    log_time = datetime.now().strftime('%H:%M:%S')
    location = f"[{interaction.guild.name}]" if interaction.guild else "[DM]"
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    print(f"[{log_time}] {location} [/{agent_key}] {actor_name}: {prompt}")

    await interaction.response.send_message(
        f"[Deep Research] `{agent_id}` 작업을 시작했습니다.\n"
        f"**주제:** {prompt[:300]}\n"
        f"작업 시간이 길어져도 완료될 때까지 추적하며, 완료되면 멘션으로 알려드립니다."
    )

    target_channel = interaction.channel
    target_channel_id = interaction.channel_id

    async def send_result(content, *, attachment=None, filename=None):
        channel = target_channel
        if channel is None and target_channel_id is not None:
            channel = bot.get_channel(target_channel_id)
        if channel is None and target_channel_id is not None:
            channel = await bot.fetch_channel(target_channel_id)
        if channel is None:
            raise RuntimeError("Deep Research 결과를 전송할 채널을 찾지 못했습니다.")
        if attachment:
            await channel.send(content, file=discord.File(fp=io.BytesIO(attachment), filename=filename))
        else:
            await channel.send(content)

    try:
        result = await call_deep_research(prompt, agent_id)
        mention = interaction.user.mention

        if isinstance(result, dict) and result.get("kind") == "deep_research_report":
            await send_result(
                f"{mention} {result.get('summary') or '[Deep Research] 완료했습니다.'}",
                attachment=result.get("html_bytes"),
                filename=result.get("filename") or "deep_research_report.html",
            )
            print(f"  [BOT]({agent_id}): [Deep Research HTML]")
        else:
            text = f"{mention} {result if isinstance(result, str) else str(result)}"
            for chunk in [text[i:i+1900] for i in range(0, len(text), 1900)]:
                await send_result(chunk)
            print(f"  [BOT]({agent_id}): {str(result)[:200]}")
        print("-------")
    except Exception as e:
        print(f"  [ERROR] /{agent_key}: {e}")
        print("-------")


@bot.tree.command(name="deepo", description="Deep Research — 입력한 주제만 단독 리서치 (약 10분, 오프라인 HTML)")
@app_commands.allowed_installs(guilds=True, users=True)
@app_commands.allowed_contexts(guilds=True, dms=True, private_channels=True)
@app_commands.describe(prompt="리서치할 주제/질문 (대화 컨텍스트는 참조하지 않음)")
async def deepo_command(interaction: discord.Interaction, prompt: str):
    await run_deep_research_command(interaction, prompt, "deepo")


@bot.tree.command(name="deepm", description="Deep Research Max — 입력한 주제만 단독 리서치 (더 깊게, 더 오래 걸림)")
@app_commands.allowed_installs(guilds=True, users=True)
@app_commands.allowed_contexts(guilds=True, dms=True, private_channels=True)
@app_commands.describe(prompt="리서치할 주제/질문 (대화 컨텍스트는 참조하지 않음)")
async def deepm_command(interaction: discord.Interaction, prompt: str):
    await run_deep_research_command(interaction, prompt, "deepm")


# ================= 모델 전환 =================

@bot.tree.command(name="model", description="AI 모델 전환")
@app_commands.allowed_installs(guilds=True, users=True)
@app_commands.allowed_contexts(guilds=True, dms=True, private_channels=True)
@app_commands.describe(name="모델 선택")
@app_commands.choices(name=[
    app_commands.Choice(name=info["display_name"], value=key)
    for key, info in AVAILABLE_MODELS.items()
])
async def switch_model(interaction: discord.Interaction, name: str = None):
    global TEXT_MODEL, IMAGE_MODEL
    if name is None:
        if interaction.guild:
            cfg = bot.get_config(interaction.guild_id)
            server_tm = cfg.get("text_model", TEXT_MODEL)
            server_im = cfg.get("image_model", IMAGE_MODEL)
            server_hist = get_history_limit(cfg)
            server_image_hist = get_image_history_limit(cfg)
            await interaction.response.send_message(
                f"이 서버 모델: `{server_tm}` (텍스트, 히스토리 {server_hist}개, 이미지 {server_image_hist}개) / `{server_im}` (이미지)\n"
                f"글로벌 기본값: `{TEXT_MODEL}` / `{IMAGE_MODEL}`"
            )
        else:
            await interaction.response.send_message(
                f"글로벌 기본 모델: `{TEXT_MODEL}` / `{IMAGE_MODEL}` (이미지)"
            )
        return
    if name not in AVAILABLE_MODELS:
        available = ', '.join(AVAILABLE_MODELS.keys())
        await temp_reply(interaction, f"'{name}' 모델을 찾을 수 없습니다. 사용 가능: {available}")
        return
    TEXT_MODEL = AVAILABLE_MODELS[name]["text"]
    IMAGE_MODEL = AVAILABLE_MODELS[name]["image"]
    thinking_level = AVAILABLE_MODELS[name]["thinking_level"]
    if interaction.guild:
        cfg = bot.get_config(interaction.guild_id)
        cfg["text_model"] = TEXT_MODEL
        cfg["image_model"] = IMAGE_MODEL
        cfg["thinking_level"] = thinking_level
        bot.save_config()
    history_limit = get_history_limit(bot.get_config(interaction.guild_id)) if interaction.guild else DEFAULT_HISTORY_LIMIT
    image_history_limit = get_image_history_limit(bot.get_config(interaction.guild_id)) if interaction.guild else DEFAULT_IMAGE_HISTORY_LIMIT
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    thinking_label = f"`{thinking_level}`"
    await interaction.response.send_message(
        f"모델 변경됨: **{name}** by {actor_name}\n"
        f"텍스트: `{TEXT_MODEL}`\n이미지: `{IMAGE_MODEL}`\n"
        f"사고 수준: {thinking_label}\n히스토리 참조: 텍스트 {history_limit}개 / 이미지 {image_history_limit}개"
    )

if __name__ == "__main__":
    bot.run(DISCORD_TOKEN)
