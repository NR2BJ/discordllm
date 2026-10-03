import discord
from discord import app_commands
from discord.ext import commands, tasks
import csv
import json
import os
import aiohttp
import asyncio
import base64
import bisect
import copy
import io
import re
import time
from datetime import datetime, timedelta, timezone
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import quote
from bs4 import BeautifulSoup

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

ENABLE_UVLOOP = os.environ.get("LLM_ENABLE_UVLOOP", "0") == "1"
if ENABLE_UVLOOP:
    try:
        import uvloop
    except Exception:
        uvloop = None
    else:
        uvloop.install()
else:
    uvloop = None

# ================= 1. 설정 및 상수 =================
DISCORD_TOKEN = os.environ["LLM_DISCORD_TOKEN"]
OPENROUTER_API_KEY = os.environ["OPENROUTER_API_KEY"]
# 웹 검색·링크 본문 추출은 OpenRouter 서버 도구(유료) 대신 Tavily를 직접 호출한다.
TAVILY_API_KEY = os.environ.get("TAVILY_API_KEY")
TAVILY_SEARCH_URL = "https://api.tavily.com/search"
TAVILY_EXTRACT_URL = "https://api.tavily.com/extract"
TAVILY_MAX_RESULTS = 5  # 모델이 결과 수를 정하지 않았을 때와 사전 검색에서 쓰는 값
TAVILY_MAX_RESULTS_CAP = 10  # 모델이 고를 수 있는 결과 수 상한 (토큰 안전장치)
TAVILY_RESULT_CHARS = 2000  # 결과 하나당 본문 최대 길이 (토큰 안전장치, 실측 최대 약 1,600자)
WEB_SEARCH_MAX_CALLS = 5  # 한 메시지에서 Tavily 검색 호출 총 상한 (크레딧 상한)
WEB_SEARCH_MAX_ROUNDS = WEB_SEARCH_MAX_CALLS  # 한 번에 하나씩 검색해도 한도를 다 쓸 수 있게 왕복 수를 맞춘다


def is_enabled(value):
    if isinstance(value, str):
        return value.strip().lower() in ("1", "true", "yes", "on")
    return bool(value)

# 🔢 [수치 조절 구역]
# 히스토리는 블록 단위로 자른다: 창이 HISTORY_LIMIT개에서 시작해 HISTORY_LIMIT + 블록 - 1개까지 뒤로만 늘고,
# 그다음 메시지에서 오래된 블록 하나를 한 번에 버린다. 블록 동안 앞부분이 그대로라 프롬프트 캐시가 맞는다.
HISTORY_LIMIT = 20
MIN_HISTORY_LIMIT = 1
MAX_HISTORY_LIMIT = 100
HISTORY_BLOCK = 20
MIN_HISTORY_BLOCK = 1
MAX_HISTORY_BLOCK = 100
# 창 안 미디어(이미지·영상·오디오)는 전부 다시 보낸다. 이 개수를 넘으면 블록을 일찍 교체한다 (0이면 히스토리 미디어 끔).
IMAGE_HISTORY_LIMIT = int(os.environ.get("LLM_IMAGE_HISTORY_LIMIT", "10"))
MIN_IMAGE_HISTORY_LIMIT = 0
MAX_IMAGE_HISTORY_LIMIT = 50
HISTORY_MEDIA_BYTES_CAP = 14 * 1024 * 1024  # 히스토리로 다시 보내는 미디어 원본 합계 상한 (요청 크기 한도 대비)
MAX_MEDIA_BASE64_MB = 20
MAX_DOC_CHARS = 12000
MEMORY_EXCERPT_MAX_CHARS = 5000

# 읽을 수 있는 텍스트 파일 확장자
SUPPORTED_EXTENSIONS = ('.txt', '.py', '.c', '.cpp', '.java', '.js', '.html', 
                        '.css', '.md', '.json', '.xml', '.yaml', '.yml', '.toml', '.ini',
                        '.sh', '.rb', '.go', '.rs', '.csv', '.tsv', '.log')

DEFAULT_MODEL = "google/gemini-3-flash-preview"
DEFAULT_IMAGE_MODEL = os.environ.get("LLM_DEFAULT_IMAGE_MODEL")
DEFAULT_IMAGE_PROMPT_MODEL = os.environ.get("LLM_IMAGE_PROMPT_MODEL") or os.environ.get("LLM_ROUTE_MODEL") or os.environ.get("LLM_MEMORY_MODEL")
DEFAULT_PROVIDER_PRIVACY = os.environ.get("LLM_PROVIDER_PRIVACY", "allow")
MEMORY_UPDATE_WINDOW = 8
# Conservative token-ish guardrail for very large prompts only. Set to 0 to disable compression entirely.
CONTEXT_COMPRESSION_MIN_TOKENS = int(os.environ.get("LLM_CONTEXT_COMPRESSION_MIN_TOKENS", "200000"))
TIMING_DEBUG_LOGS = is_enabled(os.environ.get("LLM_TIMING_DEBUG", "1"))
PROVIDER_PRIVACY_MODES = ("allow", "deny")

IMAGE_ATTACHMENT_EXTENSIONS = (".png", ".jpg", ".jpeg", ".webp", ".gif")

_COMMON_RULES = """## General Rules
- Output response content directly. No prefixes like 'Bot:' or 'Model:'.
- Respond in the user's language and mirror their tone (casual, formal, vulgar — match them).
- If the user sets up a specific roleplay scenario (character, world, setting), maintain it consistently within that conversation. Do NOT treat casual chat as "established context" that must persist.
- If you see `[System Error]` or `[Error]` in history, YOUR previous action failed.
- When a URL is in the message, the system fetches it and attaches `[Webpage Content]` — use that content directly.
- Analyze attached files (images/video/audio) when supported.
- Keep responses concise and Discord-appropriate.
- If web search is enabled for the request, use the available search tool for up-to-date information when needed. Otherwise rely on your training knowledge."""

SYSTEM_PROMPT = f"""
You are an AI member of this Discord server.

## User Identification
- Each historical user message starts with `[from display_name (id:numeric_id)]`.
- The CURRENT message you should respond to starts with `[CURRENT MESSAGE from display_name (id:numeric_id)]` — focus your reply on this user.
- Before that marker, the bot automatically inserts a `[System time: YYYY-MM-DD (Day) HH:MM KST — …]` line. The user did not write it. It is the real current date and time; your training data is older, so do not assume an earlier year.
- Past bot replies start with `[Bot replied to display_name (id:numeric_id)]`, indicating which user that reply addressed.
- The numeric ID is a UNIQUE identifier. Two users with similar-looking names are ALWAYS different people if their IDs differ.
- NEVER merge or conflate messages from different IDs, even if topics are similar or names look alike.
- Track each user's statements separately. What User A said is not what User B said.

## Conversation Context
- Chat history is for CONTEXT ONLY — to understand the flow, not to recycle.
- When the topic changes, focus ONLY on the current topic. Do NOT drag in old topics.
- Do NOT reuse specific words, jokes, memes, or references from previous conversations to fake familiarity. Every new topic gets a fresh start.
- User's background info (job, interests, etc.) is for matching TONE only. Do NOT shoehorn it into answers unless the user explicitly brings it up or it has direct logical relevance to the current question.
- Answer new questions cleanly and on-point. Do NOT over-decorate answers by mixing in past context.
- You may briefly reference a previous topic AT MOST once if naturally relevant, then move on entirely.

{_COMMON_RULES}
"""

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.environ.get("CONFIG_FILE", os.path.join(BASE_DIR, "server.json"))
KST = timezone(timedelta(hours=9))
DOC_CACHE_DIR = os.environ.get("LLM_DOC_CACHE_DIR", os.path.join(os.path.dirname(CONFIG_FILE), "llm_doc_cache"))


def format_current_time_marker():
    """현재 메시지 맨 앞에 봇이 자동으로 붙이는 현재 시각 줄 (KST).

    모델이 학습 시점을 오늘로 믿고 검색어에 옛 날짜를 넣는 문제가 실측됐다 (gemini-3-flash-preview가
    오늘을 2024-05-23으로 믿고 2024년 뉴스를 검색). 매 분 바뀌는 값이라 시스템 프롬프트(히스토리 앞)에 두면
    히스토리 캐시가 깨져서, 매번 새로 바뀌는 현재 메시지에 붙인다. 이 줄의 뜻은 시스템 프롬프트가 설명한다.
    "실제 현재 날짜" 문장을 줄 안에 같이 둬야 검색어 연도가 맞았다 (날짜만: 첫 검색어 2026 1/4, 문장 포함: 4/4).
    """
    now = datetime.now(KST)
    return (
        f"[System time: {now.strftime('%Y-%m-%d (%a) %H:%M KST')} — this is the real current date, "
        f"so do not assume an earlier year than {now.year}]"
    )

# ================= 2. 데이터 관리자 =================
class DataManager:
    @staticmethod
    def load_json(filename, default_data):
        if os.path.exists(filename):
            try:
                with open(filename, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except:
                return default_data
        return default_data

    @staticmethod
    def save_json(filename, data):
        try:
            with open(filename, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
        except:
            pass


class DocumentCache:
    def __init__(self, cache_dir):
        self.cache_dir = cache_dir
        self.docs_dir = os.path.join(cache_dir, "docs")
        self.index_file = os.path.join(cache_dir, "index.json")
        os.makedirs(self.docs_dir, exist_ok=True)
        self.index = self._load_index()

    def _default_index(self):
        return {"version": 1, "next_id": 1, "docs": {}, "message_to_docs": {}}

    def _load_index(self):
        if os.path.exists(self.index_file):
            try:
                with open(self.index_file, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                if isinstance(data, dict):
                    data.setdefault("version", 1)
                    data.setdefault("next_id", 1)
                    data.setdefault("docs", {})
                    data.setdefault("message_to_docs", {})
                    return data
            except Exception:
                pass
        return self._default_index()

    def _save_index(self):
        with open(self.index_file, 'w', encoding='utf-8') as f:
            json.dump(self.index, f, indent=2, ensure_ascii=False)

    def _make_doc_id(self):
        doc_id = f"doc_{self.index['next_id']:06d}"
        self.index["next_id"] += 1
        return doc_id

    def _doc_path(self, doc_id):
        return os.path.join(self.docs_dir, f"{doc_id}.txt")

    def _write_doc(self, doc_id, content):
        with open(self._doc_path(doc_id), 'w', encoding='utf-8') as f:
            f.write(content)

    def store_document(self, doc_type, source_label, content, message_id, channel_id):
        message_key = str(message_id)
        docs_for_message = self.index["message_to_docs"].setdefault(message_key, [])
        now = datetime.now(timezone.utc).isoformat()

        for doc_id in docs_for_message:
            doc = self.index["docs"].get(doc_id)
            if doc and doc.get("type") == doc_type and doc.get("source_label") == source_label:
                self._write_doc(doc_id, content)
                doc.update({
                    "channel_id": str(channel_id),
                    "updated_at": now,
                    "size": len(content),
                })
                self._save_index()
                return doc_id

        doc_id = self._make_doc_id()
        self._write_doc(doc_id, content)
        self.index["docs"][doc_id] = {
            "id": doc_id,
            "type": doc_type,
            "source_label": source_label,
            "message_id": message_key,
            "channel_id": str(channel_id),
            "created_at": now,
            "updated_at": now,
            "size": len(content),
        }
        docs_for_message.append(doc_id)
        self._save_index()
        return doc_id

    def get_docs_for_message(self, message_id):
        message_key = str(message_id)
        result = []
        for doc_id in self.index["message_to_docs"].get(message_key, []):
            doc = self.index["docs"].get(doc_id)
            if doc:
                result.append(doc)
        return result

    def has_source(self, message_id, doc_type, source_label):
        for doc in self.get_docs_for_message(message_id):
            if doc.get("type") == doc_type and doc.get("source_label") == source_label:
                return True
        return False

    def get_content(self, doc_id):
        path = self._doc_path(doc_id)
        if not os.path.exists(path):
            return None
        try:
            with open(path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception:
            return None

    def render_blocks(self, message_id, full_content=False):
        blocks = []
        for doc in self.get_docs_for_message(message_id):
            doc_id = doc["id"]
            label = doc.get("source_label", "")
            if full_content:
                content = self.get_content(doc_id)
                if content is None:
                    continue
                title = "Cached Webpage Content" if doc.get("type") == "webpage" else "Cached File Content"
                blocks.append(f"[{title}: {doc_id} | {label}]\n```\n{trim_document_text(content)}\n```")
            else:
                title = "Cached Webpage" if doc.get("type") == "webpage" else "Cached File"
                blocks.append(f"[{title}: {doc_id} | {label}]")
        return blocks


def trim_document_text(text, limit=MAX_DOC_CHARS):
    cleaned = (text or "").strip()
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[:limit].rstrip() + "\n...[truncated]..."


def decode_text_bytes(raw_bytes):
    for encoding in ("utf-8", "cp949", "latin1"):
        try:
            return raw_bytes.decode(encoding)
        except Exception:
            continue
    return None


def extract_pdf_text(raw_bytes):
    if PdfReader is None:
        return None
    try:
        reader = PdfReader(io.BytesIO(raw_bytes))
        parts = []
        for page in reader.pages[:20]:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return trim_document_text("\n\n".join(part for part in parts if part.strip()))
    except Exception:
        return None


def extract_docx_text(raw_bytes):
    if DocxDocument is None:
        return None
    try:
        document = DocxDocument(io.BytesIO(raw_bytes))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
        return trim_document_text("\n".join(paragraphs))
    except Exception:
        return None


def extract_xlsx_text(raw_bytes):
    if load_workbook is None:
        return None
    try:
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        blocks = []
        for sheet in workbook.worksheets[:5]:
            rows = []
            for row in sheet.iter_rows(max_row=50, values_only=True):
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
        workbook.close()
        return trim_document_text("\n\n".join(blocks))
    except Exception:
        return None


def extract_csv_text(raw_bytes, delimiter=","):
    decoded = decode_text_bytes(raw_bytes)
    if not decoded:
        return None
    try:
        reader = csv.reader(io.StringIO(decoded), delimiter=delimiter)
        rows = []
        for idx, row in enumerate(reader):
            if idx >= 100:
                break
            cells = [cell.strip() for cell in row if cell and cell.strip()]
            if cells:
                rows.append(" | ".join(cells))
        return trim_document_text("\n".join(rows))
    except Exception:
        return trim_document_text(decoded)


async def extract_attachment_text(att):
    filename = (att.filename or "").lower()
    raw_bytes = await att.read()

    if filename.endswith(".pdf"):
        text = extract_pdf_text(raw_bytes)
        return text, "pdf"
    if filename.endswith(".docx"):
        text = extract_docx_text(raw_bytes)
        return text, "docx"
    if filename.endswith(".xlsx"):
        text = extract_xlsx_text(raw_bytes)
        return text, "xlsx"
    if filename.endswith(".csv"):
        text = extract_csv_text(raw_bytes, delimiter=",")
        return text, "csv"
    if filename.endswith(".tsv"):
        text = extract_csv_text(raw_bytes, delimiter="\t")
        return text, "tsv"
    if filename.endswith(SUPPORTED_EXTENSIONS):
        decoded = decode_text_bytes(raw_bytes)
        return trim_document_text(decoded), "text"
    return None, None


def sanitize_for_memory_excerpt(text):
    cleaned = re.sub(
        r"\[(?:Attached File Content|Current File|Webpage Content):[^\]]+\]\s*```.*?```",
        "[Large attached content omitted]",
        text or "",
        flags=re.DOTALL,
    )
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned[:MEMORY_EXCERPT_MAX_CHARS]


def clamp_history_limit(value):
    try:
        limit = int(value)
    except (TypeError, ValueError):
        limit = HISTORY_LIMIT
    return max(MIN_HISTORY_LIMIT, min(MAX_HISTORY_LIMIT, limit))


def get_history_limit(config):
    return clamp_history_limit((config or {}).get("history_limit", HISTORY_LIMIT))


def clamp_image_history_limit(value):
    try:
        limit = int(value)
    except (TypeError, ValueError):
        limit = IMAGE_HISTORY_LIMIT
    return max(MIN_IMAGE_HISTORY_LIMIT, min(MAX_IMAGE_HISTORY_LIMIT, limit))


def get_image_history_limit(config):
    return clamp_image_history_limit((config or {}).get("image_history_limit", IMAGE_HISTORY_LIMIT))


def clamp_history_block(value):
    try:
        block = int(value)
    except (TypeError, ValueError):
        block = HISTORY_BLOCK
    return max(MIN_HISTORY_BLOCK, min(MAX_HISTORY_BLOCK, block))


def get_history_block(config):
    return clamp_history_block((config or {}).get("history_block", HISTORY_BLOCK))


def select_history_window(
    message_ids,
    *,
    anchor_id,
    covered,
    base,
    block,
    media_counts,
    media_bytes,
    media_cap,
    media_bytes_cap,
):
    """블록 단위 히스토리 창의 시작 위치를 고른다. 반환: (시작 인덱스, 이유)

    message_ids: 히스토리 후보 ID (오래된 → 최신, 현재 메시지 제외)
    anchor_id: 이 채널의 현재 블록 시작 메시지 ID (봇 재시작 직후 등은 None)
    covered: 후보 목록이 anchor까지 거슬러 올라갔는지 (아니면 블록을 새로 시작)
    창 길이는 base ~ base + block - 1 사이를 오간다. 미디어 개수·용량이 상한을 넘으면
    블록을 일찍 교체해 최근 base개로 줄인다. 최근 base개만으로도 상한을 넘으면 교체해도 소용없고
    메시지마다 교체가 반복돼 캐시만 깨지므로 그때는 교체하지 않는다.
    이유: keep(블록 유지) | new(새 블록) | rollover(정기 교체) | media(미디어 때문에 조기 교체)
    """
    total = len(message_ids)
    if total == 0:
        return 0, "new"
    if anchor_id is not None and covered:
        start = bisect.bisect_left(message_ids, anchor_id)
        reason = "keep"
    else:
        start = max(0, total - base)
        reason = "new"
    if total - start < base and start > 0:
        start, reason = max(0, total - base), "new"
    window = total - start
    if window >= base + block:
        start += ((window - base) // block) * block
        reason = "rollover"
    if media_cap > 0:
        early_start = max(start, total - base)
        over_now = sum(media_counts[start:]) > media_cap or sum(media_bytes[start:]) > media_bytes_cap
        fits_after = (
            sum(media_counts[early_start:]) <= media_cap and sum(media_bytes[early_start:]) <= media_bytes_cap
        )
        if over_now and early_start > start and fits_after:
            start, reason = early_start, "media"
    return start, reason


def plan_history_media(msg, *, is_bot, enabled, can_see_image, can_see_video, can_see_audio):
    """히스토리 메시지에서 다시 보낼 미디어 첨부 목록 (다운로드 전, 첨부 메타데이터만 본다).

    LLM 봇은 봇이 보낸 이미지는 다시 보내지 않고 텍스트 표시만 남긴다.
    """
    if not enabled or is_bot:
        return []
    planned = []
    for att in msg.attachments:
        if (att.size or 0) > MAX_MEDIA_BASE64_MB * 1024 * 1024:
            continue
        ct = att.content_type or ""
        if is_image_attachment(att):
            if can_see_image:
                planned.append(att)
        elif ct.startswith("video"):
            if can_see_video:
                planned.append(att)
        elif ct.startswith("audio"):
            if can_see_audio:
                planned.append(att)
    return planned


def fit_history_media_to_budget(media_plans, bytes_cap):
    """창 안 미디어 원본 합계가 상한을 넘으면 최신부터 담고, 넘치는 오래된 미디어는 뺀다."""
    fitted = [[] for _ in media_plans]
    used = 0
    for index in range(len(media_plans) - 1, -1, -1):
        for att in media_plans[index]:
            size = att.size or 0
            if used + size > bytes_cap:
                continue
            used += size
            fitted[index].append(att)
    return fitted


def get_bot_display_name():
    user = getattr(bot, "user", None)
    return getattr(user, "display_name", None) or getattr(user, "name", None) or "Bot"


def get_bot_username():
    user = getattr(bot, "user", None)
    return getattr(user, "name", None) or get_bot_display_name()


def get_server_display_name(guild, user):
    user_id = getattr(user, "id", None)
    member = guild.get_member(user_id) if guild and user_id is not None else None
    if member is not None:
        return getattr(member, "display_name", None) or getattr(member, "name", None) or str(user_id)
    if guild and isinstance(user, discord.Member) and getattr(user.guild, "id", None) == guild.id:
        return getattr(user, "display_name", None) or getattr(user, "name", None) or str(user_id)
    if guild:
        return str(user_id or "unknown")
    return getattr(user, "display_name", None) or getattr(user, "name", None) or str(user_id or "unknown")


def format_user_identity(guild, user, kind="history"):
    """사용자 식별 마커 생성. kind: 'history' | 'current'"""
    name = get_server_display_name(guild, user)
    uid = str(getattr(user, "id", "unknown"))
    if kind == "current":
        return f"[CURRENT MESSAGE from {name} (id:{uid})]"
    return f"[from {name} (id:{uid})]"


def format_bot_target_marker(display_name, user_id):
    """봇 응답이 어떤 사용자에게 답한 것인지 표시하는 prefix 마커."""
    return f"[Bot replied to {display_name} (id:{user_id})]"


def prepend_user_identity(text, guild, user, kind="history"):
    """사용자 메시지 본문 앞에 식별 마커를 prefix로 붙임."""
    body = (text or "").lstrip()
    marker = format_user_identity(guild, user, kind=kind)
    if not body:
        return marker
    return f"{marker}\n{body}"


def prepend_bot_target(text, display_name, user_id):
    """봇 응답 본문 앞에 응답 대상 사용자 마커를 prefix로 붙임."""
    body = (text or "").lstrip()
    marker = format_bot_target_marker(display_name, user_id)
    if not body:
        return marker
    return f"{marker}\n{body}"


def clean_assistant_response_text(text):
    if not isinstance(text, str):
        return text
    # 신규 형식: 모델이 마커를 모방해 응답 첫 줄에 붙인 경우 제거
    text = re.sub(r"^\s*\[System time: [^\]\n]+\]\s*", "", text)
    text = re.sub(r"^\s*\[CURRENT MESSAGE from [^\]\n]+\]\s*", "", text)
    text = re.sub(r"^\s*\[Bot replied to [^\]\n]+\]\s*", "", text)
    text = re.sub(r"^\s*\[from [^\]\n]+\]\s*", "", text)
    # 레거시 형식 호환
    text = re.sub(r"^\s*\[사용자:[^\]\n]+\]\s*(?:\n\s*\[ID:[^\]\n]+\]\s*)?", "", text)
    text = re.sub(r"^\s*\[ID:[^\]\n]+\]\s*", "", text)
    text = re.sub(r"^\s*discord_author=\{[^\n]*\}\s*", "", text)
    text = re.sub(r"\s*discord_author=\{[^\n]*\}\s*$", "", text)
    return text.lstrip()


def normalize_model_id(model_id):
    return (model_id or "").strip()

# ================= 3. OpenRouter API 클라이언트 =================
class CompletionContentError(Exception):
    def __init__(self, message, *, reasoning_only=False, finish_reason="unknown", retryable=True, provider=None):
        super().__init__(message)
        self.reasoning_only = reasoning_only
        self.finish_reason = finish_reason
        self.retryable = retryable
        # 빈 content를 반환한 provider name (있을 때만). reasoning-only 재시도에서 회전 대상으로 사용.
        self.provider = provider


class APIStatusError(Exception):
    """HTTP non-200 응답. fallback 분기 위해 상태코드/본문/매칭된 provider를 보존."""
    def __init__(self, status, body, provider=None):
        super().__init__(f"API Error ({status}): {body}")
        self.status = status
        self.body = body or ""
        self.provider = provider  # OpenRouter 응답 헤더/본문에서 추출된 provider name (있을 때만)


class OpenRouterClient:
    def __init__(self, api_key):
        self.api_key = api_key
        self.base_url = "https://openrouter.ai/api/v1"
        self.headers = {
            "Authorization": f"Bearer {api_key}",
            "HTTP-Referer": "https://discord.com",
            "X-Title": "Discord Bot",
            "Content-Type": "application/json",
            "User-Agent": "DiscordBot (OpenRouterClient, 1.0)"
        }
        self.models_list = []
        self.models_source = "unknown"

    async def fetch_models_metadata(self):
        try:
            print("⏳ OpenRouter 모델 리스트 동기화 중...")
            async with aiohttp.ClientSession() as session:
                # 전체 카탈로그(/models) 우선 — user-filtered는 privacy/guardrail 동의 안 한 모델을
                # 누락시키므로 신규 이미지 모델(Grok Imagine, Recraft 등)이 자동완성에 안 뜸.
                # 호출 단계에서 활성화 안 된 모델은 OpenRouter가 에러로 안내.
                endpoints = [
                    (f"{self.base_url}/models?output_modalities=all", "global-all"),
                    (f"{self.base_url}/models", "global"),
                    (f"{self.base_url}/models/user?output_modalities=all", "user-filtered-all-fallback"),
                    (f"{self.base_url}/models/user", "user-filtered-fallback"),
                ]
                for url, source in endpoints:
                    async with session.get(url, headers=self.headers) as resp:
                        if resp.status == 200:
                            data = await resp.json()
                            raw_list = data.get('data', [])
                            raw_list.sort(key=lambda x: x.get('created') or 0, reverse=True)
                            self.models_list = raw_list
                            self.models_source = source
                            print(f"📚 {len(self.models_list)}개 모델 로드 완료 ({source})")
                            return
                        print(f"⚠️ 모델 정보 로드 실패 ({source}): {resp.status}")
        except Exception as e:
            print(f"⚠️ 모델 정보 로드 중 오류: {e}")

    def get_model_info(self, model_id):
        clean_id = normalize_model_id(model_id)
        for m in self.models_list:
            if m['id'] == clean_id: return m
        return None

    async def resolve_model_info(self, model_id):
        clean_id = normalize_model_id(model_id)
        model = self.get_model_info(clean_id)
        if model:
            return model
        detailed = await self.fetch_model_details(clean_id)
        if detailed:
            return detailed
        return {"id": clean_id, "name": clean_id}

    def _capabilities_from_model(self, model):
        if not model:
            return {"input": ["text"], "output": ["text"]}

        input_mod = model.get("input_modalities")
        output_mod = model.get("output_modalities")

        if input_mod and output_mod:
            return {"input": input_mod, "output": output_mod}

        arch = model.get('architecture', {})
        arch_input_mod = arch.get("input_modalities")
        arch_output_mod = arch.get("output_modalities")
        if arch_input_mod and arch_output_mod:
            return {"input": arch_input_mod, "output": arch_output_mod}

        modality = arch.get('modality', 'text->text').lower()
        parts = modality.split('->')
        inp = parts[0] if parts else 'text'
        out = parts[1] if len(parts) > 1 else 'text'
        input_list = [m.strip() for m in inp.split('+')]
        output_list = [m.strip() for m in out.split('+')]
        return {"input": input_list, "output": output_list}

    async def fetch_model_details(self, model_id):
        clean_id = normalize_model_id(model_id)
        if "/" not in clean_id:
            return None

        author, slug = clean_id.split("/", 1)
        url = f"{self.base_url}/models/{quote(author, safe='')}/{quote(slug, safe='')}/endpoints"

        try:
            async with aiohttp.ClientSession() as session:
                async with session.get(url, headers=self.headers) as resp:
                    if resp.status != 200:
                        return None
                    data = await resp.json()
                    return data.get("data")
        except Exception:
            return None

    async def fetch_model_endpoints(self, model_id):
        details = await self.fetch_model_details(model_id)
        if not details:
            return None
        return details.get("endpoints", [])

    def analyze_model_capabilities(self, model_id, model_info=None):
        """모델의 입출력 capability 반환: {"input": [...], "output": [...]}"""
        clean_id = normalize_model_id(model_id)
        model = model_info or self.get_model_info(clean_id)
        return self._capabilities_from_model(model)

    def analyze_model_type(self, model_id, model_info=None):
        caps = self.analyze_model_capabilities(model_id, model_info=model_info)
        if "image" in caps["input"]:
            return 'multimodal'
        return 'text'

    def get_type_string(self, model_id, model_info=None):
        caps = self.analyze_model_capabilities(model_id, model_info=model_info)
        parts = []
        if "image" in caps["input"]: parts.append("Vision")
        if "video" in caps["input"]: parts.append("Video")
        if "audio" in caps["input"]: parts.append("Audio")
        if "image" in caps["output"]: parts.append("Image Gen")
        if "audio" in caps["output"]: parts.append("Audio Gen")
        if parts:
            return "Multimodal (" + ", ".join(parts) + ")"
        return "Text Generation"

    def _parse_completion_output(self, data, raw_text=""):
        import ast  # 문자열 형태의 리스트 파싱을 위해 추가

        # 200 OK인데 body가 null이거나 choices가 비정상인 경우 방어
        # (예: minimax 등 일부 provider가 타임아웃/블록 시 빈 body를 200으로 반환 → 기존엔 NoneType 크래시)
        def _malformed(reason_label):
            snippet = (raw_text or "").strip() or "(빈 응답 본문)"
            if len(snippet) > 1200:
                snippet = snippet[:1200] + " …(이하 생략)"
            return CompletionContentError(
                f"모델이 200 OK로 비정상 응답을 반환했습니다 ({reason_label}).\n[응답 원문]\n{snippet}",
                finish_reason="malformed_body",
                retryable=False,
            )

        if not isinstance(data, dict) or not isinstance(data.get('choices'), list) or not data['choices']:
            raise _malformed("choices 없음")
        choice = data['choices'][0]
        if not isinstance(choice, dict):
            raise _malformed("choices[0]가 비정상")
        usage = data.get('usage', {})
        message = choice.get('message', {})
        raw_content = message.get('content', '')
        msg_keys = list(message.keys())
        data_keys = [k for k in data.keys() if k not in ('id', 'object', 'created', 'model', 'choices', 'usage')]
        print(f"[DEBUG] message keys: {msg_keys}, content type: {type(raw_content).__name__}, extra data keys: {data_keys}")
        if raw_content is None:
            print(f"[DEBUG] content is None, full message: {json.dumps({k: str(v)[:200] for k, v in message.items()}, ensure_ascii=False)}")
        final_content = ""
        images = []

        if raw_content is None:
            raw_content = ''

        def _extract_images(source, keys):
            found = []
            for key in keys:
                extra_imgs = source.get(key, [])
                if isinstance(extra_imgs, list):
                    for item in extra_imgs:
                        if isinstance(item, str):
                            found.append(item)
                        elif isinstance(item, dict):
                            if item.get('type') == 'image_url':
                                img_url = item.get('image_url', {}).get('url', '')
                                if img_url:
                                    found.append(img_url)
                            else:
                                url = item.get('url') or item.get('b64_json') or item.get('data', '')
                                if url:
                                    found.append(url)
                elif isinstance(extra_imgs, str) and extra_imgs:
                    found.append(extra_imgs)
            return found

        images.extend(_extract_images(message, ['images', 'image']))
        images.extend(_extract_images(data, ['data', 'images']))

        if isinstance(raw_content, list):
            for part in raw_content:
                if isinstance(part, dict):
                    if part.get('type') in ('text', 'output_text'):
                        final_content += part.get('text', '')
                    elif part.get('type') == 'image_url':
                        img_url = part.get('image_url', {}).get('url', '')
                        if img_url:
                            images.append(img_url)
                    elif part.get('type') == 'image':
                        src = part.get('source', {})
                        if src.get('data'):
                            mt = src.get('media_type', 'image/png')
                            images.append(f"data:{mt};base64,{src['data']}")
        elif isinstance(raw_content, str):
            stripped = raw_content.strip()
            if stripped.startswith('[') and stripped.endswith(']'):
                try:
                    parsed_list = ast.literal_eval(stripped)
                    if isinstance(parsed_list, list):
                        for part in parsed_list:
                            if isinstance(part, dict):
                                if part.get('type') in ('text', 'output_text'):
                                    final_content += part.get('text', '')
                                elif part.get('type') == 'image_url':
                                    img_url = part.get('image_url', {}).get('url', '')
                                    if img_url:
                                        images.append(img_url)
                        if not final_content and not images:
                            final_content = raw_content
                    else:
                        final_content = raw_content
                except Exception:
                    final_content = raw_content
            else:
                final_content = raw_content
        else:
            final_content = str(raw_content)

        # 도구 호출 응답은 content가 비어 있는 게 정상이다. tool_calls 처리는 호출자(도구 루프)가 맡는다.
        if message.get("tool_calls"):
            return final_content, usage, images, data

        if not final_content.strip() and not images:
            finish_reason = choice.get("finish_reason") or data.get("finish_reason") or "unknown"
            has_reasoning = bool(message.get("reasoning") or message.get("reasoning_details"))
            if has_reasoning:
                raise CompletionContentError(
                    "모델이 최종 답변 없이 reasoning만 반환했습니다 "
                    f"(finish_reason={finish_reason}). 다시 시도하거나 다른 provider/model을 선택해 주세요.",
                    reasoning_only=True,
                    finish_reason=finish_reason,
                )
            raise CompletionContentError(
                "모델이 빈 응답을 반환했습니다 "
                f"(finish_reason={finish_reason}). 다시 시도하거나 다른 provider/model을 선택해 주세요.",
                finish_reason=finish_reason,
            )

        return final_content, usage, images, data

    def _build_content_retry_payload(self, payload):
        retry_payload = copy.deepcopy(payload)
        retry_instruction = (
            "The previous completion returned no visible final answer in message.content. "
            "Retry and put the user-visible final answer in message.content. "
            "Do not return only reasoning/thinking. "
            "If a JSON response format is requested, return the required JSON as the visible final answer."
        )
        retry_messages = copy.deepcopy(payload.get("messages", []))
        if retry_messages and retry_messages[0].get("role") == "system" and isinstance(retry_messages[0].get("content"), str):
            retry_messages[0]["content"] += f"\n\n[Retry instruction]\n{retry_instruction}"
        else:
            retry_messages.insert(0, {"role": "system", "content": retry_instruction})
        retry_payload["messages"] = retry_messages

        # Compression occasionally interacts badly with reasoning/final-answer separation.
        plugins = retry_payload.get("plugins")
        if isinstance(plugins, list):
            plugins = [p for p in plugins if not (isinstance(p, dict) and p.get("id") == "context-compression")]
            if plugins:
                retry_payload["plugins"] = plugins
            else:
                retry_payload.pop("plugins", None)
        return retry_payload

    async def chat_completion(
        self,
        model,
        messages,
        capabilities=None,
        image_config=None,
        *,
        tools=None,
        plugins=None,
        response_format=None,
        provider=None,
        cache_control=None,
        session_id=None,
    ):
        payload = {
            "model": model,
            "messages": messages,
        }
        # 같은 세션(채널) 요청을 같은 provider로 보내 그쪽 프롬프트 캐시를 이어 쓰게 한다 (OpenRouter sticky routing).
        if session_id:
            payload["session_id"] = session_id

        if tools:
            payload["tools"] = copy.deepcopy(tools)
            payload["tool_choice"] = "auto"

        plugin_list = []
        if should_use_context_compression(messages):
            plugin_list.append({"id": "context-compression"})
        if plugins:
            plugin_list.extend(copy.deepcopy(plugins))
        if plugin_list:
            payload["plugins"] = plugin_list

        if capabilities and "image" in capabilities.get("output", []):
            if "text" in capabilities.get("output", []):
                payload["modalities"] = ["image", "text"]
            else:
                payload["modalities"] = ["image"]

        if image_config:
            payload["image_config"] = image_config
        if response_format:
            payload["response_format"] = response_format
        if provider:
            payload["provider"] = copy.deepcopy(provider)
        if cache_control:
            payload["cache_control"] = cache_control

        # 추론형 모델: 토큰 수를 직접 박지 않고 effort로 reasoning 예산을 자동 스케일한다.
        # effort="high"는 모델별 한도의 약 0.8 비율로 OpenRouter가 변환하므로, 토큰 수를
        # 하드코딩하지 않고도 "허용 최대선에서 자동 조절(과도 제한 없음)"에 해당한다.
        # reasoning은 best-effort 옵션이라 provider.require_parameters와 함께 보내지 않는다
        # (함께 보내면 라우팅이 "reasoning 지원 endpoint"로 좁혀져 endpoint가 0개가 될 수 있다).
        # reasoning 미지원 모델에도 주입하지 않아 불필요한 endpoint 필터링을 피한다.
        model_meta = self.get_model_info(model)
        supported = (model_meta or {}).get("supported_parameters") or []
        if isinstance(supported, list) and any(str(s).lower() == "reasoning" for s in supported):
            payload.setdefault("reasoning", {"effort": "high"})

        async with aiohttp.ClientSession() as session:
            async def _post_completion(request_payload):
                async with session.post(f"{self.base_url}/chat/completions", headers=self.headers, json=request_payload) as resp:
                    # provider 추출: 헤더 우선, 200 body의 provider 필드, 5xx body의 provider 필드 순
                    provider_used = (
                        resp.headers.get("X-OpenRouter-Provider")
                        or resp.headers.get("X-Provider")
                        or resp.headers.get("X-Provider-Name")
                    )
                    if resp.status == 200:
                        raw_text = await resp.text()
                        try:
                            data = json.loads(raw_text) if raw_text and raw_text.strip() else None
                        except Exception:
                            data = None
                        if not provider_used and isinstance(data, dict):
                            provider_used = data.get("provider")
                        try:
                            final_content, usage, images, full_data = self._parse_completion_output(data, raw_text=raw_text)
                        except CompletionContentError as content_err:
                            # 빈 content를 반환한 provider를 에러에 실어 보내 재시도에서 회전할 수 있게 한다.
                            if content_err.provider is None:
                                content_err.provider = provider_used
                            raise
                        return final_content, usage, images, full_data
                    error_text = await resp.text()
                    if not provider_used:
                        try:
                            err_json = json.loads(error_text)
                            if isinstance(err_json, dict):
                                provider_used = (
                                    (err_json.get("error") or {}).get("provider")
                                    if isinstance(err_json.get("error"), dict) else None
                                ) or err_json.get("provider")
                        except Exception:
                            pass
                    raise APIStatusError(resp.status, error_text, provider=provider_used)

            try:
                return await _post_completion(payload)
            except APIStatusError as api_err:
                raise Exception(f"API Error ({api_err.status}): {api_err.body}") from api_err
            except CompletionContentError as first_error:
                if not first_error.retryable:
                    # 비정상 body 등 재시도 무의미 → 사용자에게 응답 원문 그대로 노출
                    print(f"[WARN] malformed completion (no retry): model={model} finish_reason={first_error.finish_reason}")
                    raise
                # 빈 content / reasoning-only 응답은 같은 provider로 재시도하면 동일하게 실패하는 경우가 잦다.
                # → (1) content 강조 지시문 주입 + (2) 직전 실패 provider를 ignore에 추가해 다른 provider로 회전.
                # 도구 정의(tools)는 그대로 유지한다.
                print(
                    "[WARN] completion content empty; retrying with provider rotation "
                    f"(model={model}, reasoning_only={first_error.reasoning_only}, "
                    f"finish_reason={first_error.finish_reason}, provider={first_error.provider or '?'})"
                )
                MAX_CONTENT_RETRIES = 2
                # _build_content_retry_payload가 이미 deepcopy를 반환하므로 루프에선 provider만
                # 새로 만들어 끼운다(전체 재-deepcopy 불필요). ignore는 호출 간 누적한다.
                retry_payload = self._build_content_retry_payload(payload)
                ignore_acc = []  # OpenRouter가 보고한 원본 케이스 그대로 유지(대소문자 무시 dedup)

                def _add_ignore(name):
                    if name and name.lower() not in {x.lower() for x in ignore_acc}:
                        ignore_acc.append(name)

                _add_ignore(first_error.provider)
                last_content_err = first_error
                last_api_err = None
                for attempt in range(1, MAX_CONTENT_RETRIES + 1):
                    if ignore_acc:
                        prov = dict(retry_payload.get("provider") or {})
                        existing = list(prov.get("ignore") or [])
                        existing_lower = {x.lower() for x in existing if isinstance(x, str)}
                        for name in ignore_acc:
                            if name.lower() not in existing_lower:
                                existing.append(name)
                        prov["ignore"] = existing
                        # provider_order가 고정(allow_fallbacks=False)돼 있으면 회전이 막히므로
                        # 실패 회복 경로에서는 fallback을 허용해 다른 provider를 시도할 수 있게 한다.
                        if prov.get("allow_fallbacks") is False:
                            prov["allow_fallbacks"] = True
                        retry_payload["provider"] = prov
                        print(f"[WARN] content 재시도 #{attempt}/{MAX_CONTENT_RETRIES}: ignore={existing}")
                    else:
                        print(f"[WARN] content 재시도 #{attempt}/{MAX_CONTENT_RETRIES}: provider 미상 → 동일 라우팅 재시도")
                    try:
                        result = await _post_completion(retry_payload)
                        print(f"[INFO] content 재시도 #{attempt} 성공")
                        return result
                    except CompletionContentError as retry_error:
                        last_content_err = retry_error
                        last_api_err = None
                        if not retry_error.retryable:
                            raise
                        # 회전할 provider 정보가 없고 누적 ignore도 없으면 더 돌려도 동일 → 중단
                        if not retry_error.provider and not ignore_acc:
                            break
                        _add_ignore(retry_error.provider)
                        continue
                    except APIStatusError as api_retry_err:
                        # 4xx 등은 즉시 깔끔한 에러. 5xx + provider면 그 provider도 ignore에 넣고
                        # 도구 정의는 유지한 채 회전을 계속한다.
                        if api_retry_err.status < 500 or not api_retry_err.provider:
                            raise Exception(f"API Error ({api_retry_err.status}): {api_retry_err.body}") from api_retry_err
                        last_api_err = api_retry_err
                        _add_ignore(api_retry_err.provider)
                        continue
                # 재시도 모두 소진 — 마지막 실패 유형에 맞춰 에러 형식 유지
                if last_api_err is not None:
                    raise Exception(f"API Error ({last_api_err.status}): {last_api_err.body}") from last_api_err
                raise CompletionContentError(
                    str(last_content_err).replace("다시 시도하거나 ", "자동 재시도 후에도 실패했습니다. "),
                    reasoning_only=last_content_err.reasoning_only,
                    finish_reason=last_content_err.finish_reason,
                    provider=last_content_err.provider,
                ) from last_content_err

    async def chat_completion_with_tools(
        self,
        model,
        messages,
        *,
        tools,
        tool_handlers,
        max_rounds=WEB_SEARCH_MAX_ROUNDS,
        **kwargs,
    ):
        """모델이 도구를 호출하면 봇이 실행해 결과를 돌려주고, 최종 답변이 나올 때까지 반복한다.

        assistant의 tool_calls 메시지는 reasoning_details까지 받은 그대로 되돌려야 한다.
        추론 블록이 빠지면 DeepSeek V4 thinking, Gemini thought signature 등은 다음 요청을 거부한다(OpenRouter 문서 요구).
        max_rounds번 도구를 쓰고도 더 호출하려 하면, 도구 대화를 걷어내고 모은 결과를 마지막 user 메시지에
        텍스트로 붙여 도구 없이 답을 받는다. Gemini는 도구 대화가 이미 있으면 tool_choice="none"을 무시한다.
        """
        convo = list(messages)
        tool_outputs = []
        totals = {"total_tokens": 0, "prompt_tokens": 0, "cached_tokens": 0}

        def add_usage(usage):
            if not isinstance(usage, dict):
                return
            totals["total_tokens"] += usage.get("total_tokens") or 0
            totals["prompt_tokens"] += usage.get("prompt_tokens") or 0
            details = usage.get("prompt_tokens_details")
            if isinstance(details, dict):
                totals["cached_tokens"] += details.get("cached_tokens") or 0

        def finish(content, usage, images, data):
            if not (content or "").strip() and not images:
                raise CompletionContentError(
                    "모델이 도구 호출 뒤 최종 답변을 만들지 못했습니다. 다시 시도해 주세요.",
                    finish_reason="tool_loop_no_answer",
                )
            # 왕복마다 받은 사용량을 합쳐 돌려준다 (입력·캐시 토큰 로그용)
            merged_usage = dict(usage) if isinstance(usage, dict) else {}
            merged_usage["total_tokens"] = totals["total_tokens"]
            merged_usage["prompt_tokens"] = totals["prompt_tokens"]
            details = merged_usage.get("prompt_tokens_details")
            merged_usage["prompt_tokens_details"] = {
                **(details if isinstance(details, dict) else {}),
                "cached_tokens": totals["cached_tokens"],
            }
            return content, merged_usage, images, data

        for _ in range(max_rounds):
            content, usage, images, data = await self.chat_completion(model, convo, tools=tools, **kwargs)
            add_usage(usage)
            choices = (data or {}).get("choices") or [{}]
            message = (choices[0] or {}).get("message") or {}
            tool_calls = message.get("tool_calls") or []
            if not tool_calls:
                return finish(content, usage, images, data)

            assistant_msg = {"role": "assistant", "content": message.get("content"), "tool_calls": tool_calls}
            if message.get("reasoning_details") is not None:
                assistant_msg["reasoning_details"] = message["reasoning_details"]
            elif message.get("reasoning") is not None:
                assistant_msg["reasoning"] = message["reasoning"]
            convo.append(assistant_msg)

            for call in tool_calls:
                fn = call.get("function") or {}
                name = fn.get("name") or ""
                raw_args = fn.get("arguments") or "{}"
                try:
                    args = json.loads(raw_args) if isinstance(raw_args, str) else dict(raw_args)
                except (ValueError, TypeError):
                    args = {}
                handler = tool_handlers.get(name)
                result = await handler(args) if handler else f"[오류] 알 수 없는 도구: {name}"
                tool_outputs.append(result)
                convo.append({"role": "tool", "tool_call_id": call.get("id"), "content": result})

        print(f"[SEARCH] 도구 왕복 한도({max_rounds}회) 도달 → 모은 결과로 도구 없이 최종 답변 요청")
        final_messages = copy.deepcopy(messages)
        append_to_last_user_text(
            final_messages,
            "\n\n".join(tool_outputs) + "\n\n(System: 검색 한도에 도달했다. 더 검색하지 말고 위 결과만으로 지금 답하라.)",
        )
        content, usage, images, data = await self.chat_completion(model, final_messages, **kwargs)
        add_usage(usage)
        return finish(content, usage, images, data)

# ================= 4. 봇 클래스 =================
class OpenRouterBot(commands.Bot):
    def __init__(self):
        intents = discord.Intents.default()
        intents.message_content = True
        super().__init__(command_prefix="!", intents=intents)
        self.or_client = OpenRouterClient(OPENROUTER_API_KEY)
        self.server_configs = DataManager.load_json(CONFIG_FILE, {})
        self.channel_locks = {}
        self.history_anchors = {}  # 채널별 현재 히스토리 블록의 시작 메시지 ID (재시작하면 새 블록부터)
        self.doc_cache = DocumentCache(DOC_CACHE_DIR)
        self._migrate_legacy_router_keys()

    def _migrate_legacy_router_keys(self):
        """라우터 제거에 따른 cfg 마이그레이션:
        - 옛 `route_model` → 새 `image_prompt_model` (있으면)
        - 옛 `action_router` 키 제거
        """
        changed = False
        for gid, cfg in self.server_configs.items():
            if not isinstance(cfg, dict):
                continue
            if "route_model" in cfg:
                if "image_prompt_model" not in cfg and cfg.get("route_model"):
                    cfg["image_prompt_model"] = cfg["route_model"]
                    print(f"[MIGRATE] guild {gid}: route_model → image_prompt_model='{cfg['image_prompt_model']}'")
                cfg.pop("route_model", None)
                changed = True
            if "action_router" in cfg:
                cfg.pop("action_router", None)
                changed = True
        if changed:
            DataManager.save_json(CONFIG_FILE, self.server_configs)

    async def setup_hook(self):
        await self.or_client.fetch_models_metadata()
        await self.tree.sync()
        print("✅ OpenRouter Bot 준비 완료")

    def get_config(self, guild_id):
        gid = str(guild_id)
        if gid not in self.server_configs:
            self.server_configs[gid] = {
                "channel_id": None,
                "model": DEFAULT_MODEL,
                "fallback_model": DEFAULT_IMAGE_MODEL,
                "reset_time": None,
                "search": False,
                "provider_privacy": DEFAULT_PROVIDER_PRIVACY,
                "provider_order": [],
                "history_limit": HISTORY_LIMIT,
                "history_block": HISTORY_BLOCK,
                "image_history_limit": IMAGE_HISTORY_LIMIT,
                "image_prompt_model": DEFAULT_IMAGE_PROMPT_MODEL,
            }
        self.server_configs[gid].setdefault("fallback_model", DEFAULT_IMAGE_MODEL)
        self.server_configs[gid].setdefault("provider_privacy", DEFAULT_PROVIDER_PRIVACY)
        self.server_configs[gid].setdefault("provider_order", [])
        self.server_configs[gid].setdefault("history_limit", HISTORY_LIMIT)
        self.server_configs[gid].setdefault("history_block", HISTORY_BLOCK)
        self.server_configs[gid].setdefault("image_history_limit", IMAGE_HISTORY_LIMIT)
        self.server_configs[gid].setdefault("image_prompt_model", DEFAULT_IMAGE_PROMPT_MODEL)
        self.server_configs[gid].pop("provider_policy", None)
        self.server_configs[gid].pop("provider_allow_fallbacks", None)
        # 옛 라우터 키 정리 (재진입 시 대비)
        self.server_configs[gid].pop("route_model", None)
        self.server_configs[gid].pop("action_router", None)
        return self.server_configs[gid]

    def get_channel_lock(self, channel_id):
        if channel_id not in self.channel_locks:
            self.channel_locks[channel_id] = asyncio.Lock()
        return self.channel_locks[channel_id]

    def save_config(self):
        DataManager.save_json(CONFIG_FILE, self.server_configs)

bot = OpenRouterBot()

# ================= 5. 유틸리티 함수 =================
async def scrap_webpage(url):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"
    }
    try:
        async with aiohttp.ClientSession(headers=headers) as session:
            async with session.get(url, timeout=5) as resp:
                if resp.status == 200:
                    html = await resp.text()
                    soup = BeautifulSoup(html, 'html.parser')
                    for script in soup(["script", "style", "nav", "footer", "iframe"]):
                        script.extract()
                    text = soup.get_text(separator='\n')
                    clean_text = '\n'.join([line.strip() for line in text.splitlines() if line.strip()])
                    return clean_text
    except Exception as e:
        print(f"[WARN] URL 읽기 실패: {e}")
    return None


# ----------------- Tavily 웹 검색 / 링크 본문 추출 -----------------
# 검색 방법(검색어·언어·출처·기간·깊이·횟수)은 모델이 정한다. 설명에는 각 옵션이 하는 일과 비용만 적는다.
TAVILY_TOPICS = ("general", "news")
TAVILY_TIME_RANGES = ("day", "week", "month", "year")
TAVILY_SEARCH_DEPTHS = ("basic", "advanced", "fast", "ultra-fast")
ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

WEB_SEARCH_TOOL = {
    "type": "function",
    "function": {
        "name": "web_search",
        "description": (
            "Search the web (Tavily). Each result has a title, URL, published date when known, and relevant "
            f"text from the page. You decide how to search. Up to {WEB_SEARCH_MAX_CALLS} searches per message."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "query": {"type": "string", "description": "Search query."},
                "topic": {
                    "type": "string",
                    "enum": list(TAVILY_TOPICS),
                    "description": "general (default): any web page. news: news articles.",
                },
                "time_range": {
                    "type": "string",
                    "enum": list(TAVILY_TIME_RANGES),
                    "description": (
                        "Only results published within this period before now. "
                        "Cannot be combined with start_date/end_date."
                    ),
                },
                "start_date": {
                    "type": "string",
                    "description": "Only results published on or after this date (YYYY-MM-DD).",
                },
                "end_date": {
                    "type": "string",
                    "description": "Only results published on or before this date (YYYY-MM-DD).",
                },
                "search_depth": {
                    "type": "string",
                    "enum": list(TAVILY_SEARCH_DEPTHS),
                    "description": (
                        "basic (default): balanced relevance and speed. advanced: highest relevance, slower, "
                        "costs 2 credits instead of 1. fast: lower latency. ultra-fast: lowest latency, "
                        "one summary per page."
                    ),
                },
                "max_results": {
                    "type": "integer",
                    "minimum": 1,
                    "maximum": TAVILY_MAX_RESULTS_CAP,
                    "description": f"Number of results (default {TAVILY_MAX_RESULTS}).",
                },
                "country": {
                    "type": "string",
                    "description": (
                        "Boost results from this country, as a lowercase English country name such as "
                        "'south korea', 'united states', 'japan'. Applies to topic general."
                    ),
                },
                "include_domains": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "Only return results from these domains.",
                },
                "exclude_domains": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "Never return results from these domains.",
                },
            },
            "required": ["query"],
        },
    },
}


async def _tavily_post(url, body, timeout_seconds):
    """Tavily API POST. 성공 시 (data, None), 실패 시 (None, 오류 문자열)."""
    if not TAVILY_API_KEY:
        return None, "TAVILY_API_KEY가 설정되지 않았습니다."
    headers = {"Authorization": f"Bearer {TAVILY_API_KEY}", "Content-Type": "application/json"}
    try:
        async with aiohttp.ClientSession() as session:
            async with session.post(
                url, headers=headers, json=body, timeout=aiohttp.ClientTimeout(total=timeout_seconds)
            ) as resp:
                text = await resp.text()
                if resp.status != 200:
                    return None, f"HTTP {resp.status}: {text[:300]}"
                return json.loads(text), None
    except Exception as e:
        return None, f"{type(e).__name__}: {e}"


def tavily_search_options(options):
    """모델이 고른 검색 옵션 중 형식이 맞는 것만 Tavily 요청 필드로 옮긴다.

    옵션 조합은 판단하지 않는다. 잘못된 조합(예: time_range + start_date)이나 없는 국가명은
    Tavily가 400으로 이유를 돌려주고, 그 문구가 모델에게 그대로 가서 모델이 고쳐 다시 검색할 수 있다.
    """
    options = options or {}
    picked = {}
    if options.get("topic") in TAVILY_TOPICS:
        picked["topic"] = options["topic"]
    if options.get("time_range") in TAVILY_TIME_RANGES:
        picked["time_range"] = options["time_range"]
    for key in ("start_date", "end_date"):
        value = str(options.get(key) or "").strip()
        if ISO_DATE_RE.match(value):
            picked[key] = value
    if options.get("search_depth") in TAVILY_SEARCH_DEPTHS:
        picked["search_depth"] = options["search_depth"]
    try:
        max_results = int(options.get("max_results") or TAVILY_MAX_RESULTS)
    except (TypeError, ValueError):
        max_results = TAVILY_MAX_RESULTS
    picked["max_results"] = min(max(max_results, 1), TAVILY_MAX_RESULTS_CAP)
    country = str(options.get("country") or "").strip().lower()
    if country:
        picked["country"] = country
    for key in ("include_domains", "exclude_domains"):
        domains = options.get(key)
        if isinstance(domains, str):
            domains = [domains]
        if isinstance(domains, list):
            cleaned = [str(d).strip() for d in domains if str(d).strip()][:20]
            if cleaned:
                picked[key] = cleaned
    return picked


async def tavily_search(query, options=None):
    """Tavily 검색. 반환: {"ok", "results", "credits", "error", "options"(실제로 보낸 검색 옵션)}"""
    picked = tavily_search_options(options)
    body = {"query": query, **picked, "include_published_date": True, "include_usage": True}
    data, error = await _tavily_post(TAVILY_SEARCH_URL, body, timeout_seconds=40)
    if error:
        return {"ok": False, "results": [], "credits": None, "error": error, "options": picked}
    results = [
        {
            "title": item.get("title") or "",
            "url": item["url"],
            "content": (item.get("content") or "")[:TAVILY_RESULT_CHARS],
            "published_date": item.get("published_date"),
        }
        for item in data.get("results") or []
        if item.get("url")
    ]
    credits = (data.get("usage") or {}).get("credits")
    return {"ok": True, "results": results, "credits": credits, "error": None, "options": picked}


def format_search_results(query, outcome):
    """모델에게 넘길 검색 결과 텍스트. 실패도 숨기지 않고 그대로 알린다.

    검색 시각을 같이 적는다. 모델이 학습 시점을 현재로 착각해 최신 결과를 '미래 날짜라 이상하다'며
    버리는 일이 실측됐다(DeepSeek V4: "2026-10-01 기준이라 현재 시점(2025년)과 맞지 않습니다").
    """
    if not outcome.get("ok"):
        return (
            f"[web_search 실패] query={query!r}: {outcome.get('error')}\n"
            "옵션을 고쳐 다시 검색해도 된다. 검색 없이 답한다면 최신 정보가 필요한 부분은 확인이 필요하다고 밝혀라."
        )
    if not outcome.get("results"):
        return f"[web_search 결과 없음] query={query!r}"
    searched_at = datetime.now(KST).strftime("%Y-%m-%d %H:%M KST")
    lines = [f"[web_search 결과] query={query!r} (검색 시각: {searched_at})"]
    for i, item in enumerate(outcome["results"], start=1):
        date = f" ({item['published_date']})" if item.get("published_date") else ""
        lines.append(f"{i}. {item['title']}{date}\n   URL: {item['url']}\n   {item['content']}")
    return "\n".join(lines)


def make_web_search_handler(max_calls=WEB_SEARCH_MAX_CALLS):
    """메시지 하나 동안 쓰는 web_search 실행기. Tavily 호출 횟수를 max_calls로 제한한다."""
    state = {"calls": 0}

    async def handler(args):
        args = args or {}
        query = str(args.get("query") or "").strip()
        if not query:
            return "[web_search 실패] 검색어가 비어 있습니다."
        if state["calls"] >= max_calls:
            return f"[web_search 중단] 이번 메시지의 검색 한도({max_calls}회)에 도달했습니다. 지금까지 찾은 결과로 답하세요."
        state["calls"] += 1
        outcome = await tavily_search(query, args)
        count = len(outcome.get("results") or [])
        detail = f"credits={outcome.get('credits')}" if outcome.get("ok") else f"error={outcome.get('error')}"
        print(
            f"[SEARCH] #{state['calls']} query={query!r} opts={outcome.get('options')} "
            f"ok={outcome.get('ok')} results={count} {detail}"
        )
        return format_search_results(query, outcome)

    return handler


async def tavily_extract(urls):
    """Tavily Extract로 URL 본문을 읽는다. 반환: ({url: 본문}, {url: 오류})"""
    urls = [u for u in urls if u][:20]
    if not urls:
        return {}, {}
    data, error = await _tavily_post(
        TAVILY_EXTRACT_URL,
        {"urls": urls, "extract_depth": "basic", "format": "markdown", "include_usage": True},
        timeout_seconds=60,
    )
    if error:
        print(f"[WEB] Tavily Extract 실패 urls={len(urls)} error={error}")
        return {}, {u: error for u in urls}
    contents = {
        item["url"]: item["raw_content"]
        for item in data.get("results") or []
        if item.get("url") and (item.get("raw_content") or "").strip()
    }
    failed = {
        item["url"]: item.get("error") or "추출 실패"
        for item in data.get("failed_results") or []
        if item.get("url")
    }
    credits = (data.get("usage") or {}).get("credits")
    print(f"[WEB] Tavily Extract urls={len(urls)} ok={len(contents)} failed={len(failed)} credits={credits}")
    return contents, failed


def select_search_mode(config, model_info):
    """검색 방식: off / tools(모델이 필요할 때 web_search 호출) / presearch(도구 미지원 모델용 사전 검색)."""
    if not config.get("search"):
        return "off"
    if not TAVILY_API_KEY:
        print("[SEARCH] TAVILY_API_KEY 미설정 → 검색 비활성")
        return "off"
    supported = (model_info or {}).get("supported_parameters")
    if isinstance(supported, list) and supported and not any(str(s).lower() == "tools" for s in supported):
        return "presearch"
    return "tools"


def describe_search_mode(config, model_info):
    """대시보드·/search 표시용 검색 상태 문구."""
    if not config.get("search"):
        return "OFF"
    if not TAVILY_API_KEY:
        return "ON (TAVILY_API_KEY 미설정 → 동작 안 함)"
    if select_search_mode(config, model_info) == "presearch":
        return "ON (Tavily · 매 메시지 사전 검색 — 도구 미지원 모델)"
    return "ON (Tavily · 모델이 필요할 때 검색)"


def append_to_last_user_text(messages, block):
    """마지막 user 메시지의 텍스트 끝에 block을 덧붙인다."""
    for msg in reversed(messages):
        if msg.get("role") != "user":
            continue
        content = msg.get("content")
        if isinstance(content, str):
            msg["content"] = f"{content}\n\n{block}"
        elif isinstance(content, list):
            for part in reversed(content):
                if isinstance(part, dict) and part.get("type") == "text":
                    part["text"] = f"{part.get('text') or ''}\n\n{block}"
                    break
            else:
                content.append({"type": "text", "text": block})
        return


async def apply_presearch(messages, query):
    """도구 미지원 모델용: 현재 메시지로 한 번 검색해 결과를 마지막 user 메시지에 붙인다."""
    query = (query or "").strip()[:400]
    if not query:
        return
    outcome = await tavily_search(query)
    count = len(outcome.get("results") or [])
    detail = f"credits={outcome.get('credits')}" if outcome.get("ok") else f"error={outcome.get('error')}"
    print(f"[SEARCH] 사전 검색 query={query!r} ok={outcome.get('ok')} results={count} {detail}")
    append_to_last_user_text(messages, format_search_results(query, outcome))


def is_tool_unsupported_error(error):
    text = str(error).lower()
    return "tool" in text and ("no endpoints found" in text or "does not support" in text)


async def run_chat(model, messages, *, search_mode, search_query, **chat_kwargs):
    """검색 방식에 맞춰 채팅 호출. tools 방식에서 도구 미지원 오류가 나면 사전 검색으로 전환한다."""
    if search_mode != "tools":
        return await bot.or_client.chat_completion(model, messages, **chat_kwargs)
    try:
        return await bot.or_client.chat_completion_with_tools(
            model,
            messages,
            tools=[WEB_SEARCH_TOOL],
            tool_handlers={"web_search": make_web_search_handler()},
            **chat_kwargs,
        )
    except Exception as e:
        if not is_tool_unsupported_error(e):
            raise
        print(f"[SEARCH] {model}: 도구 호출 미지원 응답 → 사전 검색으로 전환 ({str(e)[:150]})")
        await apply_presearch(messages, search_query)
        return await bot.or_client.chat_completion(model, messages, **chat_kwargs)


async def encode_image(url):
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(url) as resp:
                if resp.status == 200:
                    data = await resp.read()
                    return base64.b64encode(data).decode('utf-8'), resp.headers.get('Content-Type')
    except: pass
    return None, None


async def download_binary(url):
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(url) as resp:
                if resp.status == 200:
                    return await resp.read(), resp.headers.get("Content-Type")
    except Exception:
        pass
    return None, None


def guess_image_mime(filename):
    ext = os.path.splitext(filename or "")[1].lower()
    return {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
        ".gif": "image/gif",
    }.get(ext, "image/png")


def is_image_attachment(attachment):
    content_type = (getattr(attachment, "content_type", None) or "").lower()
    filename = (getattr(attachment, "filename", None) or "").lower()
    return content_type.startswith("image/") or filename.endswith(IMAGE_ATTACHMENT_EXTENSIONS)


def is_audio_or_video_attachment(attachment):
    content_type = (getattr(attachment, "content_type", None) or "").lower()
    return content_type.startswith("video/") or content_type.startswith("audio/")


def attachment_media_type(attachment):
    content_type = (getattr(attachment, "content_type", None) or "").lower()
    if "/" in content_type:
        return content_type.split("/", 1)[0]
    if is_image_attachment(attachment):
        return "image"
    return "file"


def infer_attachment_mime(attachment, detected_mime=None):
    if detected_mime:
        return detected_mime.split(";", 1)[0].strip()
    content_type = getattr(attachment, "content_type", None)
    if content_type:
        return content_type.split(";", 1)[0].strip()
    if is_image_attachment(attachment):
        return guess_image_mime(getattr(attachment, "filename", None))
    return "application/octet-stream"


async def send_long_message(channel, content):
    if len(content) <= 2000:
        await channel.send(content)
        return
    
    chunks = []
    while content:
        if len(content) <= 1900:
            chunks.append(content)
            break
        split_index = content.rfind('\n', 0, 1900)
        if split_index == -1: split_index = 1900
        chunks.append(content[:split_index])
        content = content[split_index:].lstrip()

    for chunk in chunks:
        await channel.send(chunk)
        await asyncio.sleep(0.5)

async def send_generated_images(send_func, images, filename_prefix="generated", on_sent=None):
    sent = 0
    for idx, img_data_url in enumerate(images):
        try:
            if img_data_url.startswith("http://") or img_data_url.startswith("https://"):
                async with aiohttp.ClientSession() as session:
                    async with session.get(img_data_url) as resp:
                        if resp.status != 200:
                            raise Exception(f"remote image fetch failed: {resp.status}")
                        binary = await resp.read()
                        mime = resp.headers.get("Content-Type", "image/png")
                        ext = mime.split("/")[-1] if "/" in mime else "png"
            elif img_data_url.startswith("data:"):
                header, b64_data = img_data_url.split(",", 1)
                mime = header.split(":")[1].split(";")[0] if ":" in header else "image/png"
                ext = mime.split("/")[-1] if "/" in mime else "png"
                binary = base64.b64decode(b64_data)
            else:
                b64_data = img_data_url
                ext = "png"
                binary = base64.b64decode(b64_data)
            await send_func(file=discord.File(fp=io.BytesIO(binary), filename=f"{filename_prefix}_{idx}.{ext}"))
            if on_sent:
                await on_sent(binary, ext, idx)
            sent += 1
        except Exception as img_err:
            print(f"⚠️ 이미지 전송 실패: {img_err}")
    return sent

def format_tokens(n):
    if n is None: return "N/A"
    try:
        n = int(n)
        if n >= 1_000_000: return f"{n/1_000_000:.1f}M"
        if n >= 1_000: return f"{n/1_000:.0f}K"
        return str(n)
    except: return "N/A"

def format_date(timestamp):
    if not timestamp: return "N/A"
    try:
        dt = datetime.fromtimestamp(int(timestamp))
        return dt.strftime("%Y-%m-%d")
    except: return "N/A"

async def ensure_model_cache_loaded():
    if not bot.or_client.models_list:
        await bot.or_client.fetch_models_metadata()

async def build_model_choices(current: str, predicate=None) -> List[app_commands.Choice[str]]:
    await ensure_model_cache_loaded()

    models = bot.or_client.models_list
    current_lower = current.lower()
    filtered = [
        m for m in models
        if (current_lower in m['id'].lower() or current_lower in m['name'].lower())
        and (predicate(m) if predicate else True)
    ]

    choices = []
    for m in filtered[:25]:
        name_display = f"{m['name']} ({format_date(m.get('created'))})"
        if len(name_display) > 100:
            name_display = name_display[:100]
        choices.append(app_commands.Choice(name=name_display, value=m['id']))
    return choices

def count_image_generation_models():
    return sum(
        1
        for model in bot.or_client.models_list
        if "image" in bot.or_client.analyze_model_capabilities(model["id"], model_info=model)["output"]
    )

async def resolve_model_bundle(model_id):
    info = await bot.or_client.resolve_model_info(model_id)
    endpoints = await bot.or_client.fetch_model_endpoints(model_id)
    caps = bot.or_client.analyze_model_capabilities(model_id, model_info=info)
    return info, endpoints, caps

async def validate_model_for_use(model_id, *, require_input=None, require_output=None, require_endpoint_check=False):
    clean_id = normalize_model_id(model_id)
    info, endpoints, caps = await resolve_model_bundle(clean_id)

    if require_endpoint_check and endpoints is None:
        return None, endpoints, caps, f"`{clean_id}` 모델의 endpoint 정보를 확인하지 못했습니다. 모델 목록이 오래됐거나 OpenRouter 응답이 일시적으로 불안정할 수 있습니다."

    if endpoints is not None and len(endpoints) == 0:
        return None, endpoints, caps, f"`{clean_id}` 모델은 현재 이 API 키 기준으로 사용 가능한 OpenRouter endpoint가 없습니다."

    if require_input and require_input not in caps["input"]:
        return None, endpoints, caps, f"`{clean_id}` 모델은 `{require_input}` 입력을 지원하지 않습니다."

    if require_output and require_output not in caps["output"]:
        return None, endpoints, caps, f"`{clean_id}` 모델은 `{require_output}` 출력을 지원하지 않습니다."

    return info, endpoints, caps, None

def get_active_generation_model(config, base_caps):
    if "image" in base_caps["output"]:
        return config["model"], "current-model"
    fallback_model = config.get("fallback_model") or DEFAULT_IMAGE_MODEL
    if fallback_model:
        return fallback_model, "configured-fallback"
    return None, None


def normalize_provider_order(value):
    if isinstance(value, str):
        raw = value.split(",")
    elif isinstance(value, list):
        raw = value
    else:
        raw = []
    return [str(item).strip().lower() for item in raw if str(item).strip()]


def normalize_provider_privacy(value):
    privacy = str(value or DEFAULT_PROVIDER_PRIVACY).strip().lower()
    return privacy if privacy in PROVIDER_PRIVACY_MODES else "allow"


def build_provider_preferences(config):
    provider_order = normalize_provider_order(config.get("provider_order"))
    privacy = normalize_provider_privacy(config.get("provider_privacy"))
    provider = {}
    if provider_order:
        provider["order"] = provider_order
        provider["allow_fallbacks"] = False
    if privacy == "deny":
        provider["data_collection"] = "deny"
    return provider or None


def endpoint_tag(endpoint):
    tag = endpoint.get("tag") or endpoint.get("provider_slug") or endpoint.get("slug")
    if tag:
        return str(tag).strip().lower()
    name = endpoint.get("provider_name") or endpoint.get("name") or ""
    if ":" in name:
        name = name.split(":", 1)[0]
    return re.sub(r"[^a-z0-9/._-]+", "-", name.strip().lower()).strip("-")


def endpoint_display_name(endpoint):
    return endpoint.get("provider_name") or endpoint.get("name") or endpoint_tag(endpoint) or "unknown"


def format_endpoint_status(endpoint):
    status = endpoint.get("status")
    if status in (0, "0", "operational", "online", "ok"):
        return "OK"
    if status is None:
        return "N/A"
    return str(status)


def format_uptime(value):
    try:
        return f"{float(value):.1f}%"
    except (TypeError, ValueError):
        return "N/A"


def endpoint_metric(endpoint, *keys):
    current = endpoint
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def format_latency(endpoint):
    latency = (
        endpoint.get("latency")
        or endpoint.get("latency_ms")
        or endpoint.get("latency_last_5m")
        or endpoint_metric(endpoint, "latency_last_30m", "p50")
        or endpoint_metric(endpoint, "latency_last_30m", "median")
        or endpoint_metric(endpoint, "latency", "p50")
        or endpoint_metric(endpoint, "performance", "latency")
        or endpoint_metric(endpoint, "performance", "latency_p50")
    )
    if latency is None:
        return None
    try:
        numeric = float(latency)
        if numeric > 100:
            return f"{numeric:.0f}ms"
        return f"{numeric:.2f}s"
    except (TypeError, ValueError):
        return str(latency)


def format_throughput(endpoint):
    throughput = (
        endpoint.get("throughput")
        or endpoint.get("throughput_tps")
        or endpoint.get("throughput_last_5m")
        or endpoint_metric(endpoint, "throughput_last_30m", "p50")
        or endpoint_metric(endpoint, "throughput_last_30m", "median")
        or endpoint_metric(endpoint, "throughput", "p50")
        or endpoint_metric(endpoint, "performance", "throughput")
        or endpoint_metric(endpoint, "performance", "throughput_p50")
    )
    if throughput is None:
        return None
    try:
        return f"{float(throughput):.0f}tps"
    except (TypeError, ValueError):
        return str(throughput)


def endpoint_token_value(endpoint, *keys):
    for key in keys:
        value = endpoint.get(key)
        if value is not None:
            return value
    return None


def format_endpoint_context(endpoint):
    context = endpoint_token_value(endpoint, "context_length", "context", "max_context")
    max_output = endpoint_token_value(
        endpoint,
        "max_completion_tokens",
        "max_output",
        "max_output_tokens",
        "max_tokens",
    )
    if context is None and max_output is None:
        return "ctx N/A"
    parts = []
    if context is not None:
        parts.append(f"ctx {format_tokens(context)}")
    if max_output is not None:
        parts.append(f"out {format_tokens(max_output)}")
    return "/".join(parts)


def format_endpoint_quantization(endpoint):
    quant = endpoint.get("quantization") or endpoint.get("quantization_level")
    if isinstance(quant, list):
        quant = ", ".join(str(item) for item in quant if item)
    return str(quant) if quant else None


def format_endpoint_data_policy(endpoint):
    value = (
        endpoint.get("data_collection")
        or endpoint.get("data_policy")
        or endpoint.get("dataCollection")
        or endpoint.get("zdr")
    )
    if value is True:
        return "zdr"
    if value is False or value is None or str(value).lower() in ("unknown", "n/a"):
        return None
    return str(value)


def format_endpoint_price(endpoint):
    pricing = endpoint.get("pricing") or {}
    try:
        in_p = float(pricing.get("prompt", 0)) * 1_000_000
        out_p = float(pricing.get("completion", 0)) * 1_000_000
    except (TypeError, ValueError):
        return "N/A"

    def _fmt(v):
        if v == 0:
            return "Free"
        if v < 0.001:
            return f"${v:.6f}"
        if v < 0.01:
            return f"${v:.4f}"
        if v < 1:
            return f"${v:.3f}"
        return f"${v:.2f}"

    return f"{_fmt(in_p)}/{_fmt(out_p)}"


def format_provider_mode(config):
    order = normalize_provider_order(config.get("provider_order"))
    privacy = normalize_provider_privacy(config.get("provider_privacy"))
    privacy_label = "data_collection deny" if privacy == "deny" else "data_collection allow"
    if order:
        return f"direct `{', '.join(order)}` | fallback OFF | {privacy_label}"
    return f"OpenRouter auto | {privacy_label}"


def format_endpoint_lines(endpoints, *, selected_tags=None, limit=10):
    selected_tags = set(selected_tags or [])
    if endpoints is None:
        return "endpoint 정보를 가져오지 못했습니다."
    if not endpoints:
        return "사용 가능한 endpoint가 없습니다."

    def _sort_key(endpoint):
        uptime = endpoint.get("uptime_last_30m")
        try:
            uptime_value = float(uptime)
        except (TypeError, ValueError):
            uptime_value = -1.0
        selected = endpoint_tag(endpoint) in selected_tags
        return (selected, uptime_value)

    lines = []
    for endpoint in sorted(endpoints, key=_sort_key, reverse=True)[:limit]:
        tag = endpoint_tag(endpoint) or "unknown"
        marker = "★ " if tag in selected_tags else ""
        extras = [
            format_endpoint_context(endpoint),
            format_endpoint_quantization(endpoint),
            format_endpoint_data_policy(endpoint),
        ]
        extras = [item for item in extras if item]
        latency = format_latency(endpoint)
        throughput = format_throughput(endpoint)
        metrics = [
            f"up {format_uptime(endpoint.get('uptime_last_30m'))}",
            f"lat {latency}" if latency else None,
            throughput,
            f"status {format_endpoint_status(endpoint)}",
            format_endpoint_price(endpoint),
            *extras,
        ]
        metrics = [item for item in metrics if item]
        lines.append(f"{marker}`{tag}` | {endpoint_display_name(endpoint)} | " + " | ".join(metrics))
    if len(endpoints) > limit:
        lines.append(f"... 외 {len(endpoints) - limit}개")
    return "\n".join(lines)


def build_prompt_cache_control(model_id, messages):
    clean_id = normalize_model_id(model_id)
    if not clean_id.startswith("anthropic/"):
        return None
    total_text = estimate_message_text_chars(messages)
    if total_text < 2048:
        return None
    return {"type": "ephemeral"}


def estimate_message_text_chars(messages):
    total_text = 0
    for message in messages or []:
        content = message.get("content") if isinstance(message, dict) else None
        if isinstance(content, str):
            total_text += len(content)
        elif isinstance(content, list):
            for part in content:
                if isinstance(part, dict) and part.get("type") == "text":
                    total_text += len(part.get("text", ""))
    return total_text


def estimate_text_tokens_conservative(text):
    # No single tokenizer matches all OpenRouter models. This intentionally errs high:
    # Korean/CJK/non-ASCII text is close to 1 char/token, while ASCII word runs average lower.
    token_count = 0
    ascii_run = 0
    for char in text or "":
        if char.isascii() and (char.isalnum() or char == "_"):
            ascii_run += 1
            continue
        if ascii_run:
            token_count += max(1, (ascii_run + 3) // 4)
            ascii_run = 0
        if char.isspace():
            continue
        token_count += 1
    if ascii_run:
        token_count += max(1, (ascii_run + 3) // 4)
    return token_count


def estimate_message_text_tokens(messages):
    total_tokens = 0
    for message in messages or []:
        content = message.get("content") if isinstance(message, dict) else None
        if isinstance(content, str):
            total_tokens += estimate_text_tokens_conservative(content)
        elif isinstance(content, list):
            for part in content:
                if isinstance(part, dict) and part.get("type") == "text":
                    total_tokens += estimate_text_tokens_conservative(part.get("text", ""))
    return total_tokens


def should_use_context_compression(messages):
    return CONTEXT_COMPRESSION_MIN_TOKENS > 0 and estimate_message_text_tokens(messages) >= CONTEXT_COMPRESSION_MIN_TOKENS


def coerce_json_object(text):
    if isinstance(text, dict):
        return text
    if not isinstance(text, str):
        return None
    candidate = text.strip()
    if not candidate:
        return None
    if "```" in candidate:
        match = re.search(r"```(?:json)?\s*(\{.*\})\s*```", candidate, re.DOTALL)
        if match:
            candidate = match.group(1).strip()
    start = candidate.find("{")
    end = candidate.rfind("}")
    if start >= 0 and end > start:
        candidate = candidate[start:end + 1]
    try:
        return json.loads(candidate)
    except Exception:
        return None


# ================= 6. 이벤트 핸들러 =================
@tasks.loop(hours=1)
async def refresh_models_hourly():
    """매 시간 OpenRouter 모델 리스트 자동 갱신."""
    print("[HOURLY] 모델 리스트 자동 갱신 시작")
    await bot.or_client.fetch_models_metadata()

@refresh_models_hourly.before_loop
async def _before_refresh_models_hourly():
    await bot.wait_until_ready()

@bot.event
async def on_ready():
    print(f"🤖 로그인: {bot.user}")
    if not refresh_models_hourly.is_running():
        refresh_models_hourly.start()
        print("[HOURLY] 모델 자동 갱신 스케줄러 시작 (매 1시간)")

@bot.event
async def on_message(message):
    if message.author.bot:
        return
    if not message.guild:
        return
    config = bot.get_config(message.guild.id)
    if message.channel.id != config.get("channel_id"):
        return

    lock = bot.get_channel_lock(message.channel.id)
    async with lock:
        async with message.channel.typing():
            try:
                timing_start = time.perf_counter()
                timing_last = timing_start

                def timing_mark(label):
                    nonlocal timing_last
                    if not TIMING_DEBUG_LOGS:
                        return
                    now = time.perf_counter()
                    print(
                        f"[TIMING] msg={message.id} {label} "
                        f"+{(now - timing_last) * 1000:.0f}ms total={(now - timing_start) * 1000:.0f}ms"
                    )
                    timing_last = now

                base_model = config["model"]
                base_model_info = await bot.or_client.resolve_model_info(base_model)
                base_caps = bot.or_client.analyze_model_capabilities(base_model, model_info=base_model_info)
                current_model = base_model
                caps = base_caps
                can_see_image = "image" in caps["input"]
                can_see_video = "video" in caps["input"]
                can_see_audio = "audio" in caps["input"]
                can_gen_image = "image" in caps["output"]
                current_display_name = get_server_display_name(message.guild, message.author)
                search_mode = select_search_mode(config, base_model_info)
                print(f"[CAPS] {current_model}: in={caps['input']} out={caps['output']} search={search_mode}")
                timing_mark("caps")

                sys_prompt = SYSTEM_PROMPT
                capability_hints = []
                if can_see_image:
                    capability_hints.append("You can analyze images attached by users.")
                if can_see_video:
                    capability_hints.append("You can analyze video files attached by users.")
                if can_see_audio:
                    capability_hints.append("You can analyze audio files attached by users.")
                if search_mode == "tools":
                    capability_hints.append(
                        "You have a `web_search` tool. For anything time-sensitive — recent events, news, "
                        "product/model releases, version numbers, prices, exchange rates, or facts that may have "
                        "changed after your training cutoff — call web_search before answering instead of relying "
                        "on possibly-outdated training knowledge. Base your answer on what you actually found and "
                        "cite sources inline as Markdown links. Do not search for casual chat."
                    )
                elif search_mode == "presearch":
                    capability_hints.append(
                        "Web search results for the user's current message are attached at the end of that "
                        "message as [web_search 결과]. For time-sensitive facts, rely on those results over your "
                        "training knowledge and cite sources inline as Markdown links. Ignore results that are "
                        "irrelevant to the question."
                    )
                capability_hints.append("Image generation is only available via the `/image` slash command, not in chat.")

                cannot_hints = []
                if not can_see_image:
                    cannot_hints.append("analyze images")
                if not can_see_video:
                    cannot_hints.append("analyze videos")
                if not can_see_audio:
                    cannot_hints.append("analyze audio")
                if cannot_hints:
                    capability_hints.append(
                        f"You CANNOT {', '.join(cannot_hints)}. If asked, politely explain the limitation and point the user to the relevant slash command if helpful."
                    )
                if capability_hints:
                    sys_prompt += "\n" + "\n".join(f"- {hint}" for hint in capability_hints)

                messages = [{"role": "system", "content": sys_prompt}]

                reset_dt = datetime.fromisoformat(config["reset_time"]) if config.get("reset_time") else None
                history_limit = get_history_limit(config)
                history_block = get_history_block(config)
                media_cap = get_image_history_limit(config)

                # 1) 히스토리 후보를 모은다 (최신 → 과거). 창 범위를 먼저 정하고 나서 내용을 만든다.
                candidates = []
                scan_limit = history_limit + history_block + 20
                scanned = 0
                oldest_scanned_id = None
                hit_reset = False
                async for msg in message.channel.history(limit=scan_limit):
                    scanned += 1
                    oldest_scanned_id = msg.id
                    if msg.id == message.id:
                        continue
                    if msg.id > message.id and msg.author != bot.user:
                        continue
                    if reset_dt and msg.created_at < reset_dt:
                        hit_reset = True
                        break
                    if msg.author == bot.user and getattr(msg, "interaction_metadata", None) is not None:
                        continue
                    if msg.author == bot.user and not (msg.content or "").strip() and not msg.attachments:
                        continue
                    candidates.append(msg)
                candidates.reverse()

                # 2) 블록 창을 고른다. 블록 동안 시작점이 고정돼 앞부분이 요청마다 같게 유지된다.
                media_plans = [
                    plan_history_media(
                        msg,
                        is_bot=msg.author == bot.user,
                        enabled=media_cap > 0,
                        can_see_image=can_see_image,
                        can_see_video=can_see_video,
                        can_see_audio=can_see_audio,
                    )
                    for msg in candidates
                ]
                anchor_id = bot.history_anchors.get(message.channel.id)
                covered = hit_reset or scanned < scan_limit or (
                    anchor_id is not None and oldest_scanned_id is not None and oldest_scanned_id <= anchor_id
                )
                window_start, window_reason = select_history_window(
                    [msg.id for msg in candidates],
                    anchor_id=anchor_id,
                    covered=covered,
                    base=history_limit,
                    block=history_block,
                    media_counts=[len(plan) for plan in media_plans],
                    media_bytes=[sum(att.size or 0 for att in plan) for plan in media_plans],
                    media_cap=media_cap,
                    media_bytes_cap=HISTORY_MEDIA_BYTES_CAP,
                )
                window = candidates[window_start:]
                window_media = fit_history_media_to_budget(media_plans[window_start:], HISTORY_MEDIA_BYTES_CAP)
                bot.history_anchors[message.channel.id] = window[0].id if window else None

                # 3) 창 안 메시지는 위치와 상관없이 같은 모양으로 넣는다: 미디어 전부, 문서는 캐시 기준 전문.
                history_buffer = []
                history_inline_media = 0
                for msg, planned in zip(window, window_media):
                    is_bot = msg.author == bot.user
                    role = "assistant" if is_bot else "user"
                    planned_ids = {att.id for att in planned}
                    content_list = []
                    if is_bot:
                        text_content = msg.content or ""
                        if ":" in text_content[:20]:
                            parts = text_content.split(": ", 1)
                            if len(parts) > 1 and parts[0] in [get_bot_display_name(), get_bot_username(), "Bot", "Model"]:
                                text_content = parts[1]
                        for att in msg.attachments:
                            ct = att.content_type or ""
                            if is_image_attachment(att):
                                text_content += f"\n(Bot sent an image: {att.filename})"
                            elif ct.startswith("video"):
                                text_content += "\n(Bot generated a video here)"
                    else:
                        text_content = msg.content or ""
                        for att in msg.attachments:
                            ct = att.content_type or ""
                            too_large = (att.size or 0) > MAX_MEDIA_BASE64_MB * 1024 * 1024
                            if is_image_attachment(att):
                                text_content += f"\n(User attached an image: {att.filename})"
                                if att.id in planned_ids:
                                    raw, mime = await download_binary(att.url)
                                    if raw:
                                        mime = infer_attachment_mime(att, mime)
                                        b64 = base64.b64encode(raw).decode("utf-8")
                                        content_list.append({"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}})
                                        history_inline_media += 1
                                    else:
                                        text_content += " (image download failed)"
                            elif ct.startswith("video") or ct.startswith("audio"):
                                is_video = ct.startswith("video")
                                inlined = False
                                if att.id in planned_ids:
                                    b64, mime = await encode_image(att.url)
                                    if b64:
                                        if is_video:
                                            content_list.append({"type": "video_url", "video_url": {"url": f"data:{mime};base64,{b64}"}})
                                        else:
                                            fmt = att.filename.rsplit(".", 1)[-1].lower() if "." in att.filename else "mp3"
                                            content_list.append({"type": "input_audio", "input_audio": {"data": b64, "format": fmt}})
                                        history_inline_media += 1
                                        inlined = True
                                if not inlined:
                                    label = "a video" if is_video else "an audio file"
                                    text_content += f"\n(User attached {label}: {att.filename})"
                            elif not bot.doc_cache.has_source(msg.id, "text_file", att.filename):
                                # 문서는 캐시에 저장만 하고, 아래에서 캐시 기준으로 넣어 요청마다 같은 모양이 되게 한다
                                extracted_text, doc_type = await extract_attachment_text(att)
                                if extracted_text:
                                    bot.doc_cache.store_document("text_file", att.filename, extracted_text, msg.id, msg.channel.id)
                                elif too_large:
                                    media_type = attachment_media_type(att)
                                    text_content += f"\n(User attached a {media_type} file: {att.filename}, but it was too large to inline.)"
                                elif doc_type:
                                    text_content += f"\n(User attached a {doc_type} file: {att.filename})"
                        cached_doc_blocks = bot.doc_cache.render_blocks(msg.id, full_content=True)
                        if cached_doc_blocks:
                            text_content += "\n\n" + "\n\n".join(cached_doc_blocks)
                        text_content = prepend_user_identity(text_content, message.guild, msg.author, kind="history")

                    if text_content.strip():
                        content_list.append({"type": "text", "text": text_content})
                    if not content_list:
                        continue
                    # author 메타: 봇 응답에 응답 대상 마커를 붙이기 위한 후처리용
                    item_meta = {
                        "is_bot": is_bot,
                        "user_info": (None if is_bot else (get_server_display_name(message.guild, msg.author), str(msg.author.id))),
                    }
                    history_buffer.append({"role": role, "content": content_list, "_meta": item_meta})
                print(
                    f"[HISTORY] window={len(window)} (base {history_limit} + block {history_block}) "
                    f"reason={window_reason} media={history_inline_media}/{media_cap}"
                )

                # 봇 응답에 [Bot replied to ...] 마커 prefix 추가
                # history_buffer는 시간순(oldest → newest). 직전(과거) user 메시지를 그 대상으로 추정
                last_user_info = None
                for item in history_buffer:
                    meta = item.get("_meta", {})
                    if meta.get("is_bot"):
                        if last_user_info:
                            for part in item["content"]:
                                if part.get("type") == "text":
                                    part["text"] = prepend_bot_target(part["text"], *last_user_info)
                                    break
                    else:
                        if meta.get("user_info"):
                            last_user_info = meta["user_info"]
                # _meta는 OpenAI API에 보낼 필요 없으니 제거
                for item in history_buffer:
                    item.pop("_meta", None)

                messages.extend(history_buffer)
                timing_mark(f"history count={len(history_buffer)} inline_media={history_inline_media}")

                current_content = []
                current_image_inputs = []
                user_text = message.content
                unresolved_urls = []

                urls = re.findall(r"(https?://\S+)", message.content)
                for url in urls:
                    if url.lower().endswith(".pdf"):
                        unresolved_urls.append(url)
                        continue
                    web_content = await scrap_webpage(url)
                    if web_content:
                        doc_id = bot.doc_cache.store_document("webpage", url, web_content, message.id, message.channel.id)
                        print(f"[CACHE] webpage 저장: {doc_id} ({url})")
                        user_text += f"\n\n[Webpage Content: {url}]\n```\n{trim_document_text(web_content)}\n```"
                        print(f"[WEB] 웹사이트 내용 읽음: {url}")
                    else:
                        unresolved_urls.append(url)
                if unresolved_urls:
                    # PDF와 직접 긁기에 실패한 링크는 Tavily Extract로 읽는다. 못 읽은 링크는 모델에게 그대로 알린다.
                    extracted, extract_failed = await tavily_extract(unresolved_urls)
                    extracted_by_url = {u.rstrip("/"): text for u, text in extracted.items()}
                    failed_by_url = {u.rstrip("/"): reason for u, reason in extract_failed.items()}
                    for url in unresolved_urls:
                        web_content = extracted.get(url) or extracted_by_url.get(url.rstrip("/"))
                        if web_content:
                            doc_id = bot.doc_cache.store_document("webpage", url, web_content, message.id, message.channel.id)
                            print(f"[CACHE] webpage 저장: {doc_id} ({url})")
                            user_text += f"\n\n[Webpage Content: {url}]\n```\n{trim_document_text(web_content)}\n```"
                            print(f"[WEB] Tavily Extract로 읽음: {url}")
                        else:
                            reason = extract_failed.get(url) or failed_by_url.get(url.rstrip("/")) or "본문 없음"
                            print(f"[WEB] 링크 읽기 실패: {url} ({reason})")
                            user_text += f"\n\n(System: 이 링크의 본문을 읽지 못했습니다: {url} — 내용을 추측하지 말 것)"

                has_media = False
                if message.attachments:
                    for att in message.attachments:
                        ct = att.content_type or ""
                        file_size = att.size or 0
                        too_large = file_size > MAX_MEDIA_BASE64_MB * 1024 * 1024
                        print(f"[ATT] {att.filename} ct={ct} size={file_size} can_see_image={can_see_image} too_large={too_large}")

                        if is_image_attachment(att) and not too_large:
                            raw, mime = await download_binary(att.url)
                            if raw:
                                mime = infer_attachment_mime(att, mime)
                                b64 = base64.b64encode(raw).decode("utf-8")
                                image_part = {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}
                                current_image_inputs.append(image_part)
                                if can_see_image:
                                    current_content.append(image_part)
                                    has_media = True
                                print(f"[ATT] 이미지 base64 인코딩 완료: {att.filename} ({len(b64)} chars)")
                                if not can_see_image:
                                    user_text += f"\n(User attached an image: {att.filename})"
                            else:
                                user_text += f"\n(User attached an image: {att.filename})"
                        elif ct.startswith("video") and can_see_video and not too_large:
                            b64, mime = await encode_image(att.url)
                            if b64:
                                current_content.append({"type": "video_url", "video_url": {"url": f"data:{mime};base64,{b64}"}})
                                has_media = True
                            else:
                                user_text += f"\n(User attached a video: {att.filename})"
                        elif ct.startswith("audio") and can_see_audio and not too_large:
                            b64, _ = await encode_image(att.url)
                            if b64:
                                fmt = att.filename.rsplit(".", 1)[-1].lower() if "." in att.filename else "mp3"
                                current_content.append({"type": "input_audio", "input_audio": {"data": b64, "format": fmt}})
                                has_media = True
                            else:
                                user_text += f"\n(User attached an audio file: {att.filename})"
                        else:
                            extracted_text, doc_type = await extract_attachment_text(att)
                            if extracted_text:
                                doc_id = bot.doc_cache.store_document("text_file", att.filename, extracted_text, message.id, message.channel.id)
                                print(f"[CACHE] {doc_type or 'text'} file 저장: {doc_id} ({att.filename})")
                                user_text += f"\n\n[Attached File Content: {att.filename}]\n```\n{extracted_text}\n```"
                            elif too_large:
                                media_type = attachment_media_type(att)
                                user_text += f"\n(User attached a {media_type} file: {att.filename}, but it's too large to process: {file_size // (1024 * 1024)}MB)"
                            elif is_image_attachment(att):
                                user_text += f"\n(User attached an image: {att.filename})"
                timing_mark(f"attachments current_images={len(current_image_inputs)}")

                # 현재 메시지에 이미지 첨부됐는데 모델이 비전 미지원이면 거부
                if current_image_inputs and not can_see_image:
                    await message.channel.send(
                        f"❌ `{current_model}` 모델은 이미지 분석을 지원하지 않습니다. 비전 입력을 지원하는 `/model set` 모델로 바꿔보세요."
                    )
                    return

                current_text = prepend_user_identity(user_text, message.guild, message.author, kind="current")
                current_content.append({"type": "text", "text": f"{format_current_time_marker()}\n{current_text}"})
                messages.append({"role": "user", "content": current_content})

                if search_mode == "presearch":
                    await apply_presearch(messages, message.content)
                    timing_mark("presearch")

                provider_prefs = build_provider_preferences(config)
                cache_control = build_prompt_cache_control(current_model, messages)
                timing_mark(
                    f"pre_chat chars={estimate_message_text_chars(messages)} "
                    f"approx_tokens={estimate_message_text_tokens(messages)} "
                    f"plugins={'compression' if should_use_context_compression(messages) else 'none'}"
                )

                session_id = f"discord-channel-{message.channel.id}"
                try:
                    response, usage, images, full_data = await run_chat(
                        current_model,
                        messages,
                        search_mode=search_mode,
                        search_query=message.content,
                        capabilities=caps,
                        provider=provider_prefs,
                        cache_control=cache_control,
                        session_id=session_id,
                    )
                    provider_name = full_data.get("provider") if isinstance(full_data, dict) else None
                    routed_model = full_data.get("model") if isinstance(full_data, dict) else None
                    timing_mark(f"chat provider={provider_name or '?'} model={routed_model or current_model}")
                except Exception as e:
                    timing_mark("chat_error")
                    err_msg = str(e).lower()
                    if "no endpoints found for" in err_msg:
                        print(f"⚠️ [Routing Error] {current_model}: endpoint 없음, 모델 목록 재동기화")
                        await bot.or_client.fetch_models_metadata()
                        await message.channel.send(
                            f"❌ `{current_model}` 모델은 현재 OpenRouter에서 실제 endpoint를 찾지 못했습니다.\n"
                            f"목록과 실제 라우팅 상태가 어긋났거나 provider privacy/direct 설정(`{format_provider_mode(config)}`)으로 endpoint가 걸러졌을 수 있습니다.\n"
                            f"`/model refresh` 후 `/model set`으로 다시 선택하거나 `/model provider provider:auto privacy:allow`로 완화해보세요."
                        )
                        return

                    has_any_media = has_media or any(
                        isinstance(m.get("content"), list)
                        and any(part.get("type") in ("image_url", "video_url", "input_audio") for part in m["content"])
                        for m in messages
                    )
                    media_error_markers = (
                        "image_url",
                        "video_url",
                        "input_audio",
                        "unsupported content",
                        "unsupported input",
                        "invalid image",
                        "invalid video",
                        "invalid audio",
                        "does not support images",
                        "does not support image",
                        "does not support videos",
                        "does not support video",
                        "does not support audio",
                        "multimodal",
                        "modalit",
                        "base64",
                        "'choices'",
                    )
                    is_media_error = has_any_media and any(marker in err_msg for marker in media_error_markers)
                    if not is_media_error:
                        raise e

                    print(f"⚠️ [Media Error] {current_model}: 텍스트 모드로 전환.")
                    await message.channel.send(f"⚠️ `{current_model}` 모델이 미디어를 처리하지 못했습니다. 텍스트 모드로 재시도합니다.")
                    for payload in messages:
                        if isinstance(payload.get("content"), list):
                            payload["content"] = [part for part in payload["content"] if part.get("type") == "text"]
                    if has_media and isinstance(messages[-1].get("content"), list):
                        for part in messages[-1]["content"]:
                            if part.get("type") == "text":
                                part["text"] += "\n(System: 미디어 첨부가 있었으나 이 모델이 처리하지 못했습니다.)"
                                break
                    response, usage, images, full_data = await run_chat(
                        current_model,
                        messages,
                        search_mode=search_mode,
                        search_query=message.content,
                        provider=provider_prefs,
                        cache_control=cache_control,
                        session_id=session_id,
                    )

                bot_display_name = get_bot_display_name()
                if response.startswith(f"{bot_display_name}:"):
                    response = response.replace(f"{bot_display_name}:", "", 1).strip()
                response = clean_assistant_response_text(response)

                total_tokens = usage.get("total_tokens", "N/A")
                prompt_tokens = usage.get("prompt_tokens") or 0
                prompt_details = usage.get("prompt_tokens_details", {}) if isinstance(usage, dict) else {}
                cached_tokens = (prompt_details.get("cached_tokens") or 0) if isinstance(prompt_details, dict) else 0
                cache_rate = f" ({cached_tokens * 100 // prompt_tokens}%)" if prompt_tokens else ""
                log_time = datetime.now().strftime("%H:%M:%S")

                print(f"[{log_time}] [{message.guild.name}] {current_display_name}: {message.content}")
                print(
                    f"봇({current_model}): {response} "
                    f"(Tokens: {total_tokens}, In: {prompt_tokens}, Cached: {cached_tokens}{cache_rate})"
                )
                if images:
                    print(f"  [이미지 생성]: {len(images)}장")
                if isinstance(full_data, dict) and full_data.get("provider"):
                    print(f"  [provider] {full_data.get('provider')}")
                print("-------")

                if response:
                    await send_long_message(message.channel, response)

                await send_generated_images(message.channel.send, images)

            except Exception as e:
                print(f"❌ 오류: {str(e)}")
                await message.channel.send(f"❌ 오류: {str(e)}")

# ================= 7. 명령어 =================

async def model_autocomplete(interaction: discord.Interaction, current: str) -> List[app_commands.Choice[str]]:
    return await build_model_choices(current)

async def image_generation_model_autocomplete(interaction: discord.Interaction, current: str) -> List[app_commands.Choice[str]]:
    return await build_model_choices(
        current,
        predicate=lambda m: "image" in bot.or_client.analyze_model_capabilities(m["id"], model_info=m)["output"],
    )

async def image_prompt_model_autocomplete(interaction: discord.Interaction, current: str) -> List[app_commands.Choice[str]]:
    current_lower = current.lower()
    special_choices = [
        app_commands.Choice(name="current - 현재 채팅 모델 사용", value="current"),
    ]
    choices = [
        choice for choice in special_choices
        if current_lower in choice.name.lower() or current_lower in choice.value.lower()
    ]
    # 이미지 prompt 다듬는 텍스트 모델 — vision 입력 지원 모델이면 reference 같이 분석 가능 (하이브리드)
    model_choices = await build_model_choices(
        current,
        predicate=lambda m: "text" in bot.or_client.analyze_model_capabilities(m["id"], model_info=m)["output"],
    )
    return (choices + model_choices)[:25]


async def provider_privacy_autocomplete(interaction: discord.Interaction, current: str) -> List[app_commands.Choice[str]]:
    current_lower = (current or "").lower()
    labels = {
        "allow": "allow - 학습/저장 가능 provider도 허용",
        "deny": "deny - data_collection 없는 provider만 허용",
    }
    return [
        app_commands.Choice(name=label, value=value)
        for value, label in labels.items()
        if not current_lower or current_lower in value or current_lower in label.lower()
    ][:25]


async def provider_autocomplete(interaction: discord.Interaction, current: str) -> List[app_commands.Choice[str]]:
    current_lower = (current or "").lower()
    choices = [
        app_commands.Choice(name="auto - 직접 provider 선택 해제", value="auto"),
    ]
    config = bot.get_config(interaction.guild_id)
    endpoints = await bot.or_client.fetch_model_endpoints(config.get("model", DEFAULT_MODEL)) or []
    for endpoint in endpoints:
        tag = endpoint_tag(endpoint)
        if not tag:
            continue
        latency = format_latency(endpoint)
        throughput = format_throughput(endpoint)
        metrics = [
            f"up {format_uptime(endpoint.get('uptime_last_30m'))}",
            f"lat {latency}" if latency else None,
            throughput,
            format_endpoint_price(endpoint),
        ]
        metrics = [item for item in metrics if item]
        label = (
            f"{tag} | {endpoint_display_name(endpoint)} | "
            + " | ".join(metrics)
        )
        if current_lower and current_lower not in label.lower() and current_lower not in tag.lower():
            continue
        choices.append(app_commands.Choice(name=label[:100], value=tag[:100]))
        if len(choices) >= 25:
            break
    return choices[:25]

model_group = app_commands.Group(name="model", description="모델 관련 설정")


@model_group.command(name="set", description="기본 채팅 모델 변경")
@app_commands.autocomplete(model_id=model_autocomplete)
@app_commands.describe(model_id="모델 이름 검색")
async def set_model(interaction: discord.Interaction, model_id: str):
    await interaction.response.defer()
    clean_id = normalize_model_id(model_id)
    model_info, endpoints, _, error = await validate_model_for_use(clean_id, require_endpoint_check=True)
    if error:
        await interaction.followup.send(
            f"❌ {error}\n`/model refresh` 후 다시 시도해보세요.",
        )
        return

    config = bot.get_config(interaction.guild_id)
    config["model"] = model_info['id']
    config["provider_order"] = []
    bot.save_config()
    
    type_str = bot.or_client.get_type_string(model_info['id'], model_info=model_info)
    created = format_date(model_info.get('created'))
    endpoint_suffix = f" | 엔드포인트: {len(endpoints)}개" if endpoints is not None else ""
    actor_name = get_server_display_name(interaction.guild, interaction.user)

    await interaction.followup.send(
        f"✅ {actor_name}님이 기본 모델을 변경했습니다.\n"
        f"**{model_info['name']}** (`{model_info['id']}`)\n"
        f"타입: **{type_str}** | 출시일: {created}{endpoint_suffix}",
    )

@model_group.command(name="image", description="채팅/고급 이미지 생성용 폴백 모델 설정")
@app_commands.autocomplete(model_id=image_generation_model_autocomplete)
@app_commands.describe(model_id="이미지 생성 폴백 모델 ID 또는 clear")
async def set_fallback_model(interaction: discord.Interaction, model_id: str):
    await interaction.response.defer()
    clean_id = normalize_model_id(model_id)
    config = bot.get_config(interaction.guild_id)

    if clean_id.lower() in ("clear", "none", "off"):
        config["fallback_model"] = None
        bot.save_config()
        actor_name = get_server_display_name(interaction.guild, interaction.user)
        await interaction.followup.send(f"✅ {actor_name}님이 이미지 폴백 모델 설정을 해제했습니다.")
        return

    model_info, endpoints, _, error = await validate_model_for_use(
        clean_id,
        require_output="image",
        require_endpoint_check=True,
    )
    if error:
        await interaction.followup.send(f"❌ {error}")
        return

    config["fallback_model"] = model_info["id"]
    bot.save_config()
    endpoint_suffix = f" | 엔드포인트: {len(endpoints)}개" if endpoints is not None else ""
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    await interaction.followup.send(
        f"✅ {actor_name}님이 이미지 폴백 모델을 설정했습니다.\n"
        f"**{model_info['name']}** (`{model_info['id']}`){endpoint_suffix}",
    )


@model_group.command(name="image_prompt", description="/image 슬래시 커맨드에서 prompt 다듬는 텍스트 모델 설정")
@app_commands.autocomplete(model_id=image_prompt_model_autocomplete)
@app_commands.describe(model_id="텍스트 모델 ID 또는 current (현재 채팅 모델 사용)")
async def set_image_prompt_model(interaction: discord.Interaction, model_id: str | None = None):
    config = bot.get_config(interaction.guild_id)
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    current_setting = config.get("image_prompt_model") or "(현재 채팅 모델)"

    if model_id is None:
        await interaction.response.send_message(
            f"현재 /image prompt 다듬는 모델: `{current_setting}`\n"
            f"채팅 모델: `{config.get('model', DEFAULT_MODEL)}`"
        )
        return

    await interaction.response.defer()
    clean_id = normalize_model_id(model_id)
    if clean_id.lower() in ("current", "auto", "default"):
        config["image_prompt_model"] = None
        bot.save_config()
        await interaction.followup.send(
            f"✅ {actor_name}님이 /image prompt 다듬는 모델을 현재 채팅 모델로 설정했습니다."
        )
        return

    model_info, endpoints, _, error = await validate_model_for_use(
        clean_id,
        require_output="text",
        require_endpoint_check=True,
    )
    if error:
        await interaction.followup.send(f"❌ {error}")
        return

    config["image_prompt_model"] = model_info["id"]
    bot.save_config()
    endpoint_suffix = f" | 엔드포인트: {len(endpoints)}개" if endpoints is not None else ""
    await interaction.followup.send(
        f"✅ {actor_name}님이 /image prompt 다듬는 모델을 설정했습니다.\n"
        f"**{model_info['name']}** (`{model_info['id']}`){endpoint_suffix}"
    )


@bot.tree.command(name="reset", description="히스토리 초기화")
async def reset_history(interaction: discord.Interaction):
    config = bot.get_config(interaction.guild_id)
    config["reset_time"] = datetime.now(timezone.utc).isoformat()
    bot.save_config()
    await interaction.response.send_message("🧹 히스토리 초기화")


@model_group.command(name="provider", description="프로바이더 자동/직접 라우팅 확인 및 설정")
@app_commands.autocomplete(provider=provider_autocomplete, privacy=provider_privacy_autocomplete)
@app_commands.describe(
    provider="직접 사용할 provider tag. auto/off/clear는 직접 선택 해제",
    privacy="allow=OpenRouter 기본 허용, deny=data_collection 없는 provider만 허용",
)
async def set_provider_routing(
    interaction: discord.Interaction,
    provider: str | None = None,
    privacy: str | None = None,
):
    await interaction.response.defer()
    config = bot.get_config(interaction.guild_id)
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    changed = []

    if privacy:
        clean_privacy = normalize_provider_privacy(privacy)
        if privacy.strip().lower() not in PROVIDER_PRIVACY_MODES:
            await interaction.followup.send(
                f"❌ privacy는 `{', '.join(PROVIDER_PRIVACY_MODES)}` 중 하나여야 합니다."
            )
            return
        config["provider_privacy"] = clean_privacy
        changed.append(f"privacy `{clean_privacy}`")

    if provider:
        clean_provider = provider.strip().lower()
        if clean_provider in ("auto", "clear", "none", "off", "disable", "disabled"):
            config["provider_order"] = []
            changed.append("직접 provider 해제")
        else:
            endpoints = await bot.or_client.fetch_model_endpoints(config.get("model", DEFAULT_MODEL))
            available_tags = {endpoint_tag(endpoint) for endpoint in endpoints or [] if endpoint_tag(endpoint)}
            if available_tags and clean_provider not in available_tags:
                await interaction.followup.send(
                    f"❌ `{clean_provider}` provider는 현재 모델 endpoint 목록에 없습니다.\n"
                    f"`/model provider`로 가능한 tag를 확인해보세요."
                )
                return
            config["provider_order"] = [clean_provider]
            changed.append(f"직접 provider `{clean_provider}`")

    if changed:
        bot.save_config()

    endpoints = await bot.or_client.fetch_model_endpoints(config.get("model", DEFAULT_MODEL))
    selected_tags = normalize_provider_order(config.get("provider_order"))
    embed = discord.Embed(title="🛰️ Provider Routing", color=0x4caf50)
    embed.add_field(name="Current", value=format_provider_mode(config), inline=False)
    embed.add_field(
        name=f"Endpoints for `{config.get('model', DEFAULT_MODEL)}`",
        value=format_endpoint_lines(endpoints, selected_tags=selected_tags, limit=12)[:1024],
        inline=False,
    )
    footer = "변경 없음. provider 또는 privacy 옵션을 넣으면 설정됩니다."
    if changed:
        footer = f"{actor_name}: " + ", ".join(changed)
    embed.set_footer(text=footer)
    await interaction.followup.send(embed=embed)


@model_group.command(name="history", description="히스토리 개수·블록 크기·미디어 상한 설정")
@app_commands.describe(
    count=f"항상 남기는 최소 히스토리 메시지 개수 ({MIN_HISTORY_LIMIT}~{MAX_HISTORY_LIMIT}, 사용자+봇 합산)",
    block=f"블록 크기: 히스토리가 count+block개가 되면 오래된 block개를 한 번에 버림 ({MIN_HISTORY_BLOCK}~{MAX_HISTORY_BLOCK})",
    image_count=f"히스토리 속 미디어(이미지·영상·오디오) 상한, 넘으면 블록을 일찍 교체 ({MIN_IMAGE_HISTORY_LIMIT}~{MAX_IMAGE_HISTORY_LIMIT}, 0=끄기)",
)
async def set_history_limit(
    interaction: discord.Interaction,
    count: int | None = None,
    block: int | None = None,
    image_count: int | None = None,
):
    config = bot.get_config(interaction.guild_id)
    if count is None and block is None and image_count is None:
        base = get_history_limit(config)
        block_size = get_history_block(config)
        await interaction.response.send_message(
            f"현재 히스토리: 최소 **{base}개**, 블록 **{block_size}개** → {base}~{base + block_size - 1}개 사이를 오감 "
            f"(사용자+봇 합산)\n"
            f"현재 히스토리 미디어 상한: **{get_image_history_limit(config)}개** "
            f"(범위 {MIN_IMAGE_HISTORY_LIMIT}~{MAX_IMAGE_HISTORY_LIMIT}, 0=끄기)"
        )
        return
    changed = []
    if count is not None:
        limit = clamp_history_limit(count)
        config["history_limit"] = limit
        suffix = "" if limit == count else f" (요청값 {count}에서 허용 범위로 보정)"
        changed.append(f"최소 히스토리 **{limit}개**{suffix}")
    if block is not None:
        block_size = clamp_history_block(block)
        config["history_block"] = block_size
        suffix = "" if block_size == block else f" (요청값 {block}에서 허용 범위로 보정)"
        changed.append(f"블록 **{block_size}개**{suffix}")
    if image_count is not None:
        image_limit = clamp_image_history_limit(image_count)
        config["image_history_limit"] = image_limit
        suffix = "" if image_limit == image_count else f" (요청값 {image_count}에서 허용 범위로 보정)"
        changed.append(f"히스토리 미디어 상한 **{image_limit}개**{suffix}")
    bot.save_config()
    actor_name = get_server_display_name(interaction.guild, interaction.user)
    await interaction.response.send_message(
        f"✅ {actor_name}님이 " + ", ".join(changed) + "로 설정했습니다."
    )

@bot.tree.command(name="dashboard", description="현재 모델 스펙 확인")
async def dashboard(interaction: discord.Interaction):
    await interaction.response.defer()
    config = bot.get_config(interaction.guild_id)
    curr_model = config['model']
    
    if not bot.or_client.models_list:
        await bot.or_client.fetch_models_metadata()
        
    model_meta = bot.or_client.get_model_info(curr_model) or {}
    if not model_meta:
        model_meta = await bot.or_client.fetch_model_details(curr_model) or {}
    type_str = bot.or_client.get_type_string(curr_model, model_info=model_meta)
    caps = bot.or_client.analyze_model_capabilities(curr_model, model_info=model_meta)
    is_multimodal = len(caps["input"]) > 1 or len(caps["output"]) > 1
    endpoints = await bot.or_client.fetch_model_endpoints(curr_model)

    color = 0x00bcd4 if is_multimodal else 0x4caf50
    embed = discord.Embed(title="🎛️ OpenRouter Dashboard", color=color)

    pricing = model_meta.get('pricing', {})
    try:
        in_p = float(pricing.get('prompt', 0)) * 1_000_000
        out_p = float(pricing.get('completion', 0)) * 1_000_000
        def _fmt_price(v):
            if v == 0: return "Free"
            if v < 0.001: return f"${v:.6f}"
            if v < 0.01: return f"${v:.4f}"
            if v < 1: return f"${v:.3f}"
            return f"${v:.2f}"
        price_str = f"In: {_fmt_price(in_p)} / Out: {_fmt_price(out_p)}"
    except: price_str = "N/A"

    context = format_tokens(model_meta.get('context_length'))
    max_out = format_tokens(model_meta.get('top_provider', {}).get('max_completion_tokens'))
    created = format_date(model_meta.get('created'))

    embed.add_field(name=f"🧠 {model_meta.get('name', curr_model)}", value=f"ID: `{curr_model}`\nType: **{type_str}**\nReleased: **{created}**", inline=False)
    embed.add_field(name="⚙️ Specs", value=f"• **Context:** {context}\n• **Max Out:** {max_out}\n• **Price(1M):** {price_str}", inline=True)

    # Modality 표시
    input_str = ", ".join(caps["input"]) if caps["input"] else "text"
    output_str = ", ".join(caps["output"]) if caps["output"] else "text"
    embed.add_field(name="🔄 Modalities", value=f"• **Input:** {input_str}\n• **Output:** {output_str}", inline=True)
    if endpoints is not None:
        embed.add_field(name="🛰️ Routing", value=f"• **Available Endpoints:** {len(endpoints)}", inline=True)

    image_prompt_model_status = config.get("image_prompt_model") or "(current chat model)"
    image_model_status = config.get("fallback_model") or DEFAULT_IMAGE_MODEL or "(none)"

    helpers = [
        f"• **Search:** {describe_search_mode(config, model_meta)}",
        f"• **/image — image model:** `{image_model_status}`",
        f"• **/image — prompt refiner:** `{image_prompt_model_status}`",
        f"• **Image Models discovered:** {count_image_generation_models()}",
        f"• **Provider:** {format_provider_mode(config)}",
        f"• **History:** {get_history_limit(config)}~{get_history_limit(config) + get_history_block(config) - 1} messages "
        f"(block {get_history_block(config)}) / media cap {get_image_history_limit(config)}",
    ]
    embed.add_field(name="🧩 Helpers", value="\n".join(helpers), inline=False)
    if endpoints is not None:
        endpoint_lines = format_endpoint_lines(
            endpoints,
            selected_tags=normalize_provider_order(config.get("provider_order")),
            limit=8,
        )
        embed.add_field(name="🛰️ Provider Endpoints", value=endpoint_lines[:1024], inline=False)

    desc = model_meta.get('description', '설명 없음')
    if len(desc) > 300: desc = desc[:297] + "..."
    embed.add_field(name="📝 Description", value=desc, inline=False)
    embed.set_footer(text=f"Time: {datetime.now(KST).strftime('%H:%M:%S')}")

    await interaction.followup.send(embed=embed)

@bot.tree.command(name="set", description="응답 채널 설정")
async def set_channel(interaction: discord.Interaction):
    config = bot.get_config(interaction.guild_id)
    config["channel_id"] = interaction.channel_id
    bot.save_config()
    await interaction.response.send_message(f"✅ 활성 채널: {interaction.channel.name}")

@bot.tree.command(name="search", description="웹 검색 기능 ON/OFF")
async def toggle_search(interaction: discord.Interaction):
    config = bot.get_config(interaction.guild_id)
    config["search"] = not config.get("search", False)
    bot.save_config()
    state = "ON 🔍" if config["search"] else "OFF"
    lines = [f"✅ 웹 검색: **{state}**"]
    if config["search"]:
        lines.append(f"방식: {describe_search_mode(config, bot.or_client.get_model_info(config['model']))}")
        if not TAVILY_API_KEY:
            lines.append("⚠️ 서버에 `TAVILY_API_KEY`가 없어 실제로는 검색하지 않습니다.")
    lines.append(f"기본 모델: `{config['model']}`")
    await interaction.response.send_message("\n".join(lines))


@bot.tree.command(name="image", description="이미지 생성 (prompt → 텍스트 모델로 다듬은 후 이미지 모델 호출)")
@app_commands.describe(
    prompt="생성할 이미지에 대한 설명 (한국어/영어 자유)",
    reference="참조 이미지 파일 (선택, image-to-image 지원 모델에서만 사용)",
    reference_url="참조 이미지 URL (선택)",
)
async def image_command(
    interaction: discord.Interaction,
    prompt: str,
    reference: discord.Attachment | None = None,
    reference_url: str | None = None,
):
    await interaction.response.defer()
    config = bot.get_config(interaction.guild_id) if interaction.guild_id else {}
    actor_name = (
        get_server_display_name(interaction.guild, interaction.user)
        if interaction.guild else interaction.user.display_name
    )

    # 1. 이미지 모델 결정 + capability 확인
    image_model_id = config.get("fallback_model") or DEFAULT_IMAGE_MODEL
    if not image_model_id:
        await interaction.followup.send(
            "❌ 이미지 모델이 설정되지 않았습니다. `/model image <model_id>`로 먼저 설정해주세요."
        )
        return
    image_info, _, image_caps, image_error = await validate_model_for_use(
        image_model_id, require_output="image"
    )
    if image_error:
        await interaction.followup.send(f"❌ 이미지 모델 오류: {image_error}")
        return
    image_supports_input = "image" in image_caps["input"]

    # 2. Prompt 다듬는 모델 결정 + capability 확인
    prompt_model_id = config.get("image_prompt_model") or config.get("model") or DEFAULT_MODEL
    prompt_info, _, prompt_caps, prompt_error = await validate_model_for_use(
        prompt_model_id, require_output="text"
    )
    if prompt_error:
        await interaction.followup.send(f"❌ Prompt 다듬는 모델 오류: {prompt_error}")
        return
    prompt_model_supports_vision = "image" in prompt_caps["input"]

    # 3. Reference 이미지 다운로드 (있으면)
    reference_b64 = None
    reference_mime = None
    if reference is not None:
        if not is_image_attachment(reference):
            await interaction.followup.send("❌ 첨부 파일이 이미지 형식이 아닙니다.")
            return
        raw, mime = await download_binary(reference.url)
        if not raw:
            await interaction.followup.send("❌ 첨부 이미지를 다운로드하지 못했습니다.")
            return
        reference_mime = infer_attachment_mime(reference, mime)
        reference_b64 = base64.b64encode(raw).decode("utf-8")
    elif reference_url:
        if not reference_url.lower().startswith(("http://", "https://")):
            await interaction.followup.send("❌ reference_url은 http(s) 형식이어야 합니다.")
            return
        b64, mime = await encode_image(reference_url)
        if not b64:
            await interaction.followup.send("❌ URL에서 이미지를 가져오지 못했습니다.")
            return
        reference_b64 = b64
        reference_mime = mime or "image/png"

    has_reference = reference_b64 is not None

    # 4. 사용자 입력 표시 + status 메시지
    user_display = f"**{actor_name}:** {prompt}"
    if has_reference:
        user_display += " *(reference 이미지 포함)*"
    await interaction.followup.send(user_display)

    notes = []
    if has_reference and not image_supports_input:
        notes.append(f"⚠️ 이미지 모델 `{image_model_id}`은 reference 미지원 — 새로 생성됨")
    if has_reference and not prompt_model_supports_vision:
        notes.append(f"⚠️ Prompt 다듬는 모델 `{prompt_model_id}`은 vision 미지원 — reference 분석 없이 텍스트만으로 다듬음")

    status_msg = await interaction.followup.send(
        f"⏳ Prompt 다듬는 중... (`{prompt_model_id}`)"
        + ("\n" + "\n".join(notes) if notes else "")
    )

    # 5. Prompt 다듬기 호출 (옵션 Y: 항상 prompt 다듬는 모델 거침)
    refine_system = (
        "You are a prompt refiner for an AI image generation model. "
        "Rewrite the user's request into a detailed, vivid English prompt suitable for diffusion or image generation models. "
        "Be specific about subject, style, lighting, composition, mood, and key visual details. "
        "Keep it under ~500 words. "
        "Output ONLY the refined prompt — no preamble, no quotes, no explanation."
    )
    refine_user_parts = [{"type": "text", "text": prompt}]
    if has_reference and prompt_model_supports_vision:
        refine_user_parts.append({
            "type": "image_url",
            "image_url": {"url": f"data:{reference_mime};base64,{reference_b64}"},
        })
    refine_messages = [
        {"role": "system", "content": refine_system},
        {"role": "user", "content": refine_user_parts},
    ]
    try:
        refined_text, _, _, _ = await bot.or_client.chat_completion(
            prompt_model_id,
            refine_messages,
            provider=build_provider_preferences(config),
        )
        refined_text = (refined_text or "").strip() or prompt
        if len(refined_text) > 2000:
            refined_text = refined_text[:2000]
    except Exception as e:
        print(f"⚠️ /image prompt 다듬기 실패, 원본 사용: {e}")
        refined_text = prompt

    # 6. 이미지 생성 호출
    await status_msg.edit(content=(
        f"🎨 이미지 생성 중... (`{image_model_id}`)\n"
        f"**Refined Prompt:** ```{refined_text[:1500]}```"
        + ("\n" + "\n".join(notes) if notes else "")
    ))

    if has_reference and image_supports_input:
        image_messages = [{
            "role": "user",
            "content": [
                {"type": "text", "text": refined_text},
                {"type": "image_url", "image_url": {"url": f"data:{reference_mime};base64,{reference_b64}"}},
            ],
        }]
    else:
        image_messages = [{"role": "user", "content": refined_text}]

    try:
        _, _, images, _ = await bot.or_client.chat_completion(
            image_model_id,
            image_messages,
            capabilities=image_caps,
            provider=build_provider_preferences(config),
        )
        if not images:
            await status_msg.edit(content=f"❌ 이미지 생성 실패 — 응답에 이미지 데이터 없음 (`{image_model_id}`)")
            return
        await send_generated_images(interaction.followup.send, images)
        ref_suffix = ""
        if has_reference:
            if image_supports_input:
                ref_suffix += " | reference 사용"
            else:
                ref_suffix += " | reference 무시 (모델 미지원)"
            if prompt_model_supports_vision:
                ref_suffix += " | reference 분석됨"
        await status_msg.edit(content=(
            f"✅ 생성 완료 — 이미지 모델: `{image_model_id}` / prompt: `{prompt_model_id}`{ref_suffix}"
        ))
        print(f"  [/image] image_model={image_model_id} prompt_model={prompt_model_id} "
              f"reference={has_reference} images={len(images)} refined='{refined_text[:80]}'")
    except Exception as e:
        print(f"❌ /image 생성 실패: {e}")
        await status_msg.edit(content=f"❌ 이미지 생성 실패: {str(e)[:300]}")

@model_group.command(name="refresh", description="OpenRouter 모델 리스트 즉시 갱신")
async def refresh_models_cmd(interaction: discord.Interaction):
    await interaction.response.defer()
    before = len(bot.or_client.models_list)
    await bot.or_client.fetch_models_metadata()
    after = len(bot.or_client.models_list)
    image_count = count_image_generation_models()
    if after == 0:
        await interaction.followup.send("[ERROR] 모델 리스트 갱신 실패 — OpenRouter API 응답 확인 필요")
    elif after == before:
        await interaction.followup.send(
            f"✅ 갱신 완료 (변동 없음, {after}개 모델 / 이미지 출력 {image_count}개)",
        )
    else:
        delta = after - before
        sign = "+" if delta > 0 else ""
        await interaction.followup.send(
            f"✅ 갱신 완료: {before}개 → **{after}개** ({sign}{delta}) / 이미지 출력 {image_count}개",
        )

bot.tree.add_command(model_group)

if __name__ == "__main__":
    bot.run(DISCORD_TOKEN)
